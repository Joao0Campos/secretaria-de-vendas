"""Leitura de saidas/ para o painel - SOMENTE LEITURA.

As planilhas são lidas com as MESMAS funções que a auditoria usa
(ler_dados_cadastrais, ler_lista_documentos, ler_resultados_anteriores) e a
categoria do apontamento vem de testar_amostra.categorizar(). Regra nova lá vale
aqui sem cópia: o painel não decide nada, só mostra o que a planilha já diz.

Cache por mtime: cada planilha é relida só quando muda. São ~60 pastas, e
relê-las a cada requisição deixaria o painel lento à toa.
"""
from __future__ import annotations

import hashlib
import hmac
import re
import secrets
import threading
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from analisar_documentos import (
    ABA_CONTEXTO,
    ABA_FALHAS_IA,
    CAMPOS_CONTEXTO,
    LEITURAS_SEM_IA,
    ROTULO_DOCUMENTO_NAO_ANEXADO,
    ROTULO_MODALIDADE,
    ROTULOS_DE_SENTINELA,
    classificar_modalidade,
    extrair_idprecadastro_da_pasta,
    ler_dados_cadastrais,
    ler_lista_documentos,
    ler_resultados_anteriores,
    numero_da_pasta_e_idreserva,
)
from testar_amostra import categorizar

from .banco import CATEGORIAS_RENOMEADAS, STATUS_ABERTO, Banco

TIPO_RESERVA = "reserva"
# Decisão do João (20/10/2026): o projeto não trata mais pré-cadastro. O painel
# lê só saidas/reserva/; saidas/pre_cadastro/ e os apontamentos/eventos de
# pré-cadastro no painel.db ficam onde estão, fora das telas, métricas e
# exportações (rotas /reserva/pre_cadastro/... e /arquivo/pre_cadastro/... = 404).
TIPOS = (TIPO_RESERVA,)

NIVEL_DIVERGENCIA = "divergencia"
NIVEL_AVISO = "aviso"

STATUS_DIVERGENCIA = "divergencia"
STATUS_AVISO = "aviso"
STATUS_OK = "ok"
STATUS_SEM_ANALISE = "sem_analise"
# Documento que a auditoria nem mandou para a IA (fora do escopo, dispensado
# pela modalidade...): não é "ok" nem "sem análise".
STATUS_FORA_DO_ESCOPO = "fora_do_escopo"
# Documento que a IA não conseguiu ler por falha de SISTEMA (sem saldo, limite,
# rede...; aba ABA_FALHAS_IA, desde 21/10/2026): não foi conferido. Não é
# apontamento nem categoria - é o selo "N documento(s) aguardando leitura".
STATUS_AGUARDANDO_LEITURA = "aguardando_leitura"

ABA_DIVERGENCIAS = "Divergencias (IA)"
ABAS_EXTRAS = ("Associados", "Sugestoes (NAO aplicadas)")
# A aba ABA_CONTEXTO (gravar_aba_contexto(), desde 13/10/2026) vence as fontes
# provisórias abaixo, campo a campo; campo vazio nela não apaga a fonte.

# Fontes PROVISÓRIAS de modalidade/empreendimento, por idreserva: as planilhas
# agregadas que testar_amostra.py e varrer_arquivos_repetidos.py gravam na raiz
# de saidas/. Aba "Por reserva"; (padrão do nome, coluna do id, coluna da
# modalidade, coluna do empreendimento) - colunas a partir de 0.
FONTES_PROVISORIAS = (
    ("amostra-*.xlsx", 0, 2, None),
    ("varredura-arquivo-repetido-*.xlsx", 0, 3, 2),
)

# Mesmo separador com que gravar_aba_divergencias() junta os itens de uma célula.
SEPARADOR = " | "


# --------------------------------------------------------------------------
# utilidades
# --------------------------------------------------------------------------

def normalizar(texto) -> str:
    """Sem acento, minúsculo, espaços colapsados."""
    t = unicodedata.normalize("NFKD", str(texto or ""))
    t = "".join(c for c in t if not unicodedata.combining(c)).lower()
    return " ".join(t.split())


def rotulo_modalidade(texto) -> str | None:
    if not texto:
        return None
    return ROTULO_MODALIDADE.get(classificar_modalidade(str(texto)))


def mascarar_documento(documento) -> str:
    """CPF '***.456.789-**', CNPJ '**.345.678/0001-**'. O painel nunca mostra
    o documento inteiro - nome já basta para saber de quem se trata."""
    digitos = re.sub(r"\D", "", str(documento or ""))
    if len(digitos) == 11:
        return f"***.{digitos[3:6]}.{digitos[6:9]}-**"
    if len(digitos) == 14:
        return f"**.{digitos[2:5]}.{digitos[5:8]}/{digitos[8:12]}-**"
    return "***" if digitos else ""


_RE_CNPJ = re.compile(r"(?<!\d)\d{2}\.?\d{3}\.?\d{3}/?\d{4}-?\d{2}(?!\d)")
_RE_CPF = re.compile(r"(?<!\d)\d{3}\.?\d{3}\.?\d{3}-?\d{2}(?!\d)")


def mascarar_no_texto(texto: str) -> str:
    """Mascara CPF/CNPJ dentro de um texto livre. As mensagens de divergência
    ("CPF no documento (...) difere do cadastro (...)") e as abas Associados e
    Sugestoes trazem o número inteiro; a máscara mantém os dígitos do meio, o
    bastante para ver que os dois números diferem. Um telefone de 11 dígitos
    também sai mascarado - preço aceitável."""
    texto = _RE_CNPJ.sub(lambda m: mascarar_documento(m.group()), texto)
    return _RE_CPF.sub(lambda m: mascarar_documento(m.group()), texto)


# Telefone digitado à mão: "(11) 98765-4321", "11 98765 4321", "8765-4321".
_RE_TELEFONE = re.compile(r"(?<!\d)(?:\(\d{2}\)\s*|\d{2}[\s.-]+)?9?\d{4}[\s.-]?\d{4}(?!\d)")
TELEFONE_MASCARADO = "(**) ****-****"


def mascarar_texto_livre(texto: str | None) -> str | None:
    """O "por quê" da tratativa é digitado por gente, e gente cola CPF e
    telefone do cliente. Além do que mascarar_no_texto() pega, qualquer
    sequência com cara de telefone (8 ou 9 dígitos, com ou sem DDD) sai
    mascarada - um número de protocolo de 8 dígitos também, preço aceitável."""
    if not texto:
        return texto
    return _RE_TELEFONE.sub(TELEFONE_MASCARADO, mascarar_no_texto(texto))


# Mensagens que trazem o dado pessoal entre parênteses: "Nome no documento
# ('X') difere do cadastro ('Y')", "Endereço no comprovante ('X') não bate com
# o endereço cadastrado ('Y')", "CPF no documento (X) difere do cadastro (Y)".
CATEGORIAS_COM_NOME = ("Nome diferente do cadastro",)
CATEGORIAS_COM_ENDERECO = ("Endereço não bate com o cadastro",)
CATEGORIAS_COM_DOCUMENTO = ("CPF diferente do cadastro", "CNPJ diferente do cadastro")
_RE_ENTRE_PARENTESES = re.compile(r"\('(.*?)'\)|\(([^()]*)\)")


def _plural(n: int, singular: str, plural: str) -> str:
    return f"{n} {singular if n == 1 else plural}"


def _descricao_do_nome(nome: str) -> str:
    letras = sum(1 for c in nome if c.isalpha())
    return f"nome: {_plural(len(nome.split()), 'palavra', 'palavras')}, {_plural(letras, 'letra', 'letras')}"


def _marcador_de_nome(nome: str) -> str:
    return f"({_descricao_do_nome(nome)})"


# Mensagens que caem em "Outros" no categorizar() mas trazem nome de pessoa (ou
# razão social) - texto EXATO de analisar_documentos.py, 13/10/2026. Vão por
# padrão de texto; o papel/relação entre parênteses ("Cônjuge", "2º
# Promissário Comprador") fica visível, porque ajuda a julgar o apontamento.
MENSAGENS_COM_NOME = (
    # venda PJ: razão social e representante
    (re.compile(r"Cliente é pessoa jurídica \('(.*?)'\), e este documento é de '(.*?)'"
                r"(?= - representante da empresa)"),
     lambda m: (f"Cliente é pessoa jurídica {_marcador_de_nome(m[1])}, "
                f"e este documento é de {_marcador_de_nome(m[2])}")),
    # nome do documento é de um associado
    (re.compile(r"Documento é de (.+?) \(([^()]*)\), não de (.+?)(?= - é co-participante legítimo)"),
     lambda m: (f"Documento é de {_marcador_de_nome(m[1])} ({m[2]}), "
                f"não de {_marcador_de_nome(m[3])}")),
    # mesma mensagem na redação antiga, que ainda está em planilhas já geradas
    # ("..., não do titular - é co-participante..."); achada ao medir a máscara
    # nos 192 apontamentos reais (13/10/2026)
    (re.compile(r"Documento é de (.+?) \(([^()]*)\)(?=, não do titular - é co-participante legítimo)"),
     lambda m: f"Documento é de {_marcador_de_nome(m[1])} ({m[2]})"),
    # CPF do documento é de um associado
    (re.compile(r"CPF do documento pertence a (.+?) \(([^()]*)\)(?=, co-participante da operação)"),
     lambda m: f"CPF do documento pertence a {_marcador_de_nome(m[1])} ({m[2]})"),
    # renda numa operação com mais de uma pessoa: a lista dos associados
    (re.compile(r"a operação tem mais de uma pessoa \((.*?)\)(?=, e o valor do cadastro é a soma)"),
     lambda m: ("a operação tem mais de uma pessoa ("
                + "; ".join(_descricao_do_nome(n) for n in m[1].split(", ")) + ")")),
)


def _marcador(conteudo: str, categoria: str) -> str:
    palavras = conteudo.split()
    if conteudo.startswith(("nome: ", "endereço: ")):
        return f"({conteudo})"  # já é marcador
    if categoria in CATEGORIAS_COM_NOME:
        return _marcador_de_nome(conteudo)
    if categoria in CATEGORIAS_COM_ENDERECO:
        return f"(endereço: {_plural(len(palavras), 'palavra', 'palavras')})"
    # CPF/CNPJ: a máscara de sempre. O que já veio mascarado fica como está
    # (os dígitos do meio mostram que os dois números diferem); número com
    # tamanho estranho, que mascarar_no_texto() não reconhece, vira "***".
    if "*" in conteudo or not re.search(r"\d", conteudo):
        return f"({conteudo})"
    return f"({mascarar_documento(conteudo)})"


def mascarar_para_o_time(texto: str | None, categoria: str) -> str | None:
    """Texto de apontamento (ou "por quê") para arquivo que sai do painel - o
    Excel da revisão vai para o time, inclusive agentes de IA. Nas categorias
    de nome, endereço e CPF/CNPJ, o que está entre parênteses vira um marcador
    que ainda ajuda a julgar o falso positivo sem expor o dado: "(nome: 3
    palavras, 18 letras)", "(endereço: 6 palavras)". As mensagens de
    MENSAGENS_COM_NOME (categoria "Outros") têm os nomes trocados pelo mesmo
    marcador. O resto do texto fica igual; CPF/telefone saem com a máscara de
    sempre."""
    if not texto:
        return texto
    texto = mascarar_texto_livre(texto)
    for padrao, trocar in MENSAGENS_COM_NOME:
        texto = padrao.sub(trocar, texto)
    if categoria not in CATEGORIAS_COM_NOME + CATEGORIAS_COM_ENDERECO + CATEGORIAS_COM_DOCUMENTO:
        return texto
    return _RE_ENTRE_PARENTESES.sub(
        lambda m: _marcador(m.group(1) if m.group(1) is not None else m.group(2), categoria), texto)


def url_cv(modelo: str | None, id_) -> str | None:
    """Link "Abrir no CV": o modelo da configuração (já validado em
    config.modelo_cv) com o id da reserva/pré-cadastro como inteiro."""
    return modelo.format(id=int(id_)) if modelo else None


def formatar_momento(momento: datetime | None) -> str | None:
    return momento.strftime("%d/%m/%Y %H:%M") if momento else None


def _texto_celula(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M")
    return mascarar_no_texto(str(valor))


def _itens_da_celula(valor, eh_divergencia: bool) -> list[str]:
    """Separa uma célula como ler_resultados_anteriores() faz."""
    texto = str(valor or "").strip()
    if not texto or (eh_divergencia and texto == "OK"):
        return []
    return [t for t in texto.split(SEPARADOR) if t.strip()]


# Categoria nova -> a antiga cuja identidade ela herda (ver CATEGORIAS_RENOMEADAS).
_IDENTIDADE_HERDADA = {nova: antiga for antiga, nova in CATEGORIAS_RENOMEADAS.items()}


def identidade_do_apontamento(categoria: str, texto: str) -> str:
    """O que identifica o achado na chave. Com categoria conhecida, a categoria;
    em "Outros", o texto normalizado. Categoria renomeada usa o rótulo ANTIGO
    (15/10/2026): o achado é o mesmo, e a chave tem que continuar a mesma para o
    histórico e as tratativas do painel.db não se soltarem dele."""
    if categoria == "Outros":
        return normalizar(texto)
    return _IDENTIDADE_HERDADA.get(categoria, categoria)


def chave_apontamento(tipo: str, id_: int, pasta: str, arquivo: str, nivel: str,
                      categoria: str, texto: str, ordinal: int) -> str:
    """Chave estável entre rodadas. Com categoria conhecida o texto fica de
    fora: "comprovante com 118 dias" e "com 125 dias" são o MESMO apontamento,
    e não podem virar um novo a cada reanálise. Em "Outros" não há categoria
    que identifique o achado, então entra o texto normalizado."""
    identidade = identidade_do_apontamento(categoria, texto)
    bruto = f"{tipo}|{id_}|{pasta}|{arquivo}|{nivel}|{identidade}|{ordinal}"
    return hashlib.sha1(bruto.encode("utf-8")).hexdigest()


def montar_apontamentos(tipo: str, id_: int, pasta: str,
                        itens: list[tuple[str, str, str]]) -> list[dict]:
    """itens: (arquivo, nivel, texto), na ordem da planilha. A categoria sai do
    texto original; o que é guardado e mostrado é o texto com CPF mascarado."""
    ordinais: dict[tuple, int] = {}
    saida = []
    for arquivo, nivel, original in itens:
        categoria = categorizar(original)
        texto = mascarar_no_texto(original)
        identidade = identidade_do_apontamento(categoria, texto)
        grupo = (arquivo, nivel, identidade)
        ordinal = ordinais.get(grupo, 0)
        ordinais[grupo] = ordinal + 1
        saida.append({
            "chave": chave_apontamento(tipo, id_, pasta, arquivo, nivel, categoria, texto, ordinal),
            "arquivo": arquivo, "nivel": nivel, "categoria": categoria, "texto": texto,
        })
    return saida


# --------------------------------------------------------------------------
# uma pasta (uma pessoa)
# --------------------------------------------------------------------------

@dataclass
class Pasta:
    tipo: str
    id: int
    nome_pasta: str
    caminho: Path
    xlsx: Path
    mtime: float
    cadastro: dict
    documentos: list[dict]
    analisada: bool
    resultados: dict[str, dict]
    apontamentos: list[dict]
    contexto: dict
    abas_extras: list[dict]
    # {token: nome do arquivo} - ver token_do_arquivo().
    tokens: dict[str, str] = field(default_factory=dict)
    # Linhas da aba ABA_FALHAS_IA: {arquivo, tipo_cv, classe, mensagem, em}.
    falhas_ia: list[dict] = field(default_factory=list)

    def token_de(self, arquivo: str) -> str | None:
        for token, nome in self.tokens.items():
            if nome == arquivo:
                return token
        return None

    @property
    def papel(self) -> str:
        return str(self.cadastro.get("papel_na_operacao") or "titular")

    @property
    def eh_titular(self) -> bool:
        return self.papel.strip().lower() == "titular"

    @property
    def nome(self) -> str:
        return str(self.cadastro.get("nome") or self.nome_pasta)

    @property
    def analisado_em(self) -> datetime:
        return datetime.fromtimestamp(self.mtime).replace(microsecond=0)

    def contagens(self) -> dict:
        div = sum(1 for a in self.apontamentos if a["nivel"] == NIVEL_DIVERGENCIA)
        av = sum(1 for a in self.apontamentos if a["nivel"] == NIVEL_AVISO)
        aus = sum(1 for a in self.apontamentos
                  if a["nivel"] == NIVEL_DIVERGENCIA and a["arquivo"] == ROTULO_DOCUMENTO_NAO_ANEXADO)
        return {"divergencias": div, "avisos": av, "ausentes": aus}

    @property
    def status(self) -> str:
        if not self.analisada:
            return STATUS_SEM_ANALISE
        c = self.contagens()
        if c["divergencias"]:
            return STATUS_DIVERGENCIA
        return STATUS_AVISO if c["avisos"] else STATUS_OK

    @property
    def aguardando_leitura(self) -> set[str]:
        """Arquivos com falha da IA na última rodada e sem leitura que valha (na
        releitura das Telas, a leitura velha continua valendo)."""
        return {f["arquivo"] for f in self.falhas_ia if f["arquivo"] not in self.resultados}

    @property
    def documentos_lidos(self) -> int:
        return sum(1 for r in self.resultados.values()
                   if str((r.get("doc_ia") or {}).get("tipo_identificado") or "") not in LEITURAS_SEM_IA)


def id_da_pasta(nome_pasta: str, tipo: str) -> int | None:
    """Número da pasta, conferindo que o formato bate com o tipo (armadilha 5:
    idreserva e idprecadastro não se misturam)."""
    numero = extrair_idprecadastro_da_pasta(nome_pasta)
    if numero is None:
        return None
    eh_reserva = numero_da_pasta_e_idreserva(nome_pasta)
    if (tipo == TIPO_RESERVA) != eh_reserva:
        return None
    return numero


def planilha_da_pasta(pasta: Path) -> Path | None:
    """A mesma escolha de processar_pasta_cliente() (primeiro .xlsx), sem o
    arquivo-trava que o Excel cria quando alguém está com a planilha aberta."""
    candidatas = sorted(p for p in pasta.glob("*.xlsx") if not p.name.startswith("~$"))
    return candidatas[0] if candidatas else None


def _ler_abas_proprias(xlsx: Path) -> tuple[bool, list[tuple[str, list[str], list[str]]], dict, list[dict],
                                             list[dict]]:
    """O que as funções da auditoria não expõem: se a aba de divergências
    existe, as linhas-sentinela, o contexto da análise, as abas extras e os
    documentos aguardando leitura por falha da IA."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        analisada = ABA_DIVERGENCIAS in wb.sheetnames
        sentinelas: list[tuple[str, list[str], list[str]]] = []
        if analisada:
            linhas = wb[ABA_DIVERGENCIAS].iter_rows(values_only=True)
            cabecalho = list(next(linhas, ()) or ())
            for linha in linhas:
                d = dict(zip(cabecalho, linha))
                if d.get("arquivo") in ROTULOS_DE_SENTINELA:
                    sentinelas.append((d["arquivo"],
                                       _itens_da_celula(d.get("divergencias"), True),
                                       _itens_da_celula(d.get("avisos"), False)))

        contexto: dict = {}
        if ABA_CONTEXTO in wb.sheetnames:
            for linha in wb[ABA_CONTEXTO].iter_rows(min_row=2, values_only=True):
                if linha and linha[0]:
                    campo = normalizar(linha[0])
                    if campo in CAMPOS_CONTEXTO and len(linha) > 1 and linha[1] not in (None, ""):
                        contexto[campo] = str(linha[1]).strip()

        extras = []
        for titulo in ABAS_EXTRAS:
            if titulo not in wb.sheetnames:
                continue
            linhas_cruas = []
            for linha in wb[titulo].iter_rows(values_only=True):
                celulas = [_texto_celula(v) for v in linha]
                while celulas and not celulas[-1]:
                    celulas.pop()
                if celulas:
                    linhas_cruas.append(celulas)
            extras.append({"titulo": titulo, "linhas": linhas_cruas})

        falhas_ia: list[dict] = []
        if ABA_FALHAS_IA in wb.sheetnames:
            linhas = wb[ABA_FALHAS_IA].iter_rows(values_only=True)
            cabecalho = list(next(linhas, ()) or ())
            for linha in linhas:
                d = dict(zip(cabecalho, linha))
                if d.get("arquivo"):
                    falhas_ia.append({c: _texto_celula(d.get(c)) for c in ("arquivo", "tipo_cv", "classe",
                                                                           "mensagem", "em")})
    finally:
        wb.close()
    return analisada, sentinelas, contexto, extras, falhas_ia


def token_do_arquivo(segredo: bytes, tipo: str, pasta: str, arquivo: str) -> str:
    """Identificador OPACO do arquivo na URL (13/10/2026, teste com saidas/
    real: o nome do arquivo baixado do CV traz o CPF, e a pasta traz o nome do
    cliente - 14 CPFs apareciam crus nos links). HMAC com o segredo do banco:
    não dá para adivinhar nem montar o token de outro arquivo."""
    mensagem = "\x00".join((tipo, pasta, arquivo)).encode("utf-8")
    return hmac.new(segredo, mensagem, hashlib.sha256).hexdigest()[:24]


def ler_pasta(caminho: Path, tipo: str, segredo: bytes = b"") -> Pasta | None:
    id_ = id_da_pasta(caminho.name, tipo)
    xlsx = planilha_da_pasta(caminho)
    if id_ is None or xlsx is None:
        return None
    mtime = xlsx.stat().st_mtime

    cadastro = ler_dados_cadastrais(xlsx)
    documentos = ler_lista_documentos(xlsx)
    analisada, sentinelas, contexto, extras, falhas_ia = _ler_abas_proprias(xlsx)
    resultados = ler_resultados_anteriores(xlsx) if analisada else {}

    itens: list[tuple[str, str, str]] = []
    for arquivo, divergencias, avisos in sentinelas:
        itens += [(arquivo, NIVEL_DIVERGENCIA, t) for t in divergencias]
        itens += [(arquivo, NIVEL_AVISO, t) for t in avisos]
    for arquivo, r in resultados.items():
        itens += [(arquivo, NIVEL_DIVERGENCIA, t) for t in r.get("divergencias") or []]
        itens += [(arquivo, NIVEL_AVISO, t) for t in r.get("avisos") or []]

    # Só arquivo que a planilha conhece ganha token: a rota de arquivo não
    # serve nada que não esteja na aba Documentos ou numa leitura gravada.
    arquivos = {str(d.get("arquivo_local")) for d in documentos if d.get("arquivo_local")}
    arquivos |= {a for a in resultados if a not in ROTULOS_DE_SENTINELA}
    tokens = {token_do_arquivo(segredo, tipo, caminho.name, a): a for a in sorted(arquivos)}

    return Pasta(
        tipo=tipo, id=id_, nome_pasta=caminho.name, caminho=caminho, xlsx=xlsx, mtime=mtime,
        cadastro=cadastro, documentos=documentos, analisada=analisada, resultados=resultados,
        apontamentos=montar_apontamentos(tipo, id_, caminho.name, itens),
        contexto=contexto, abas_extras=extras, tokens=tokens, falhas_ia=falhas_ia,
    )


# --------------------------------------------------------------------------
# uma reserva (todas as pessoas)
# --------------------------------------------------------------------------

@dataclass
class Reserva:
    tipo: str
    id: int
    pastas: list[Pasta]
    contexto: dict = field(default_factory=dict)

    @property
    def titular(self) -> str:
        for p in self.pastas:
            if p.eh_titular:
                return p.nome
        for p in self.pastas:
            if p.cadastro.get("titular_nome"):
                return str(p.cadastro["titular_nome"])
        return self.pastas[0].nome if self.pastas else ""

    @property
    def analisado_em(self) -> datetime | None:
        datas = [p.analisado_em for p in self.pastas if p.analisada]
        return max(datas) if datas else None

    @property
    def status(self) -> str:
        situacoes = [p.status for p in self.pastas]
        if STATUS_DIVERGENCIA in situacoes:
            return STATUS_DIVERGENCIA
        if STATUS_AVISO in situacoes:
            return STATUS_AVISO
        # Alguém da operação sem análise: não dá para dizer que está ok.
        if STATUS_SEM_ANALISE in situacoes:
            return STATUS_SEM_ANALISE
        return STATUS_OK


def _ordem_das_pessoas(p: Pasta) -> tuple:
    return (0 if p.eh_titular else 1, normalizar(p.papel), p.nome_pasta)


class Repositorio:
    """saidas/ em memória, relida por mtime, e sincronizada com o banco."""

    def __init__(self, saidas: Path, banco: Banco | None = None):
        self.saidas = Path(saidas)
        self.banco = banco
        self.segredo = banco.segredo if banco is not None else secrets.token_bytes(32)
        self._pastas: dict[tuple[str, str], Pasta] = {}
        self._fontes: dict[Path, tuple[float, dict[int, dict]]] = {}
        self.erros: dict[tuple[str, str], str] = {}
        self._trava = threading.Lock()

    def atualizar(self) -> list[Reserva]:
        with self._trava:
            vistas: set[tuple[str, str]] = set()
            for tipo in TIPOS:
                base = self.saidas / tipo
                if not base.is_dir():
                    continue
                for caminho in sorted(base.iterdir()):
                    if not caminho.is_dir():
                        continue
                    chave = (tipo, caminho.name)
                    xlsx = planilha_da_pasta(caminho)
                    if xlsx is None:
                        continue
                    vistas.add(chave)
                    atual = self._pastas.get(chave)
                    try:
                        mtime = xlsx.stat().st_mtime
                    except OSError:
                        continue
                    if atual is not None and atual.xlsx == xlsx and atual.mtime == mtime:
                        continue
                    try:
                        pasta = ler_pasta(caminho, tipo, self.segredo)
                    except Exception as exc:  # planilha corrompida ou aberta pela metade
                        self.erros[chave] = f"{type(exc).__name__}: {exc}"[:300]
                        continue
                    self.erros.pop(chave, None)
                    if pasta is None:
                        self._pastas.pop(chave, None)
                        continue
                    self._pastas[chave] = pasta
                    if self.banco is not None and pasta.analisada:
                        # No banco, o nome do arquivo vai mascarado (ele traz CPF).
                        self.banco.sincronizar_pasta(
                            pasta.tipo, pasta.nome_pasta, pasta.id, pasta.mtime,
                            [dict(a, arquivo=mascarar_no_texto(a["arquivo"])) for a in pasta.apontamentos])
            for chave in list(self._pastas):
                if chave not in vistas:
                    del self._pastas[chave]
            fontes = self._ler_fontes_provisorias()
            return self._agrupar(fontes)

    def _ler_fontes_provisorias(self) -> dict[int, dict]:
        """{idreserva: {modalidade, empreendimento}}; o arquivo mais novo vence."""
        from openpyxl import load_workbook

        arquivos = []
        for padrao, *_ in FONTES_PROVISORIAS:
            arquivos += [p for p in self.saidas.glob(padrao) if not p.name.startswith("~$")]
        for caminho in list(self._fontes):
            if caminho not in arquivos:
                del self._fontes[caminho]

        for padrao, col_id, col_mod, col_emp in FONTES_PROVISORIAS:
            for caminho in self.saidas.glob(padrao):
                if caminho.name.startswith("~$"):
                    continue
                try:
                    mtime = caminho.stat().st_mtime
                    if caminho in self._fontes and self._fontes[caminho][0] == mtime:
                        continue
                    por_id: dict[int, dict] = {}
                    wb = load_workbook(caminho, read_only=True, data_only=True)
                    try:
                        if "Por reserva" in wb.sheetnames:
                            for linha in wb["Por reserva"].iter_rows(min_row=2, values_only=True):
                                try:
                                    idr = int(linha[col_id])
                                except (TypeError, ValueError, IndexError):
                                    continue
                                info = {}
                                if col_mod is not None and len(linha) > col_mod and linha[col_mod]:
                                    info["modalidade"] = str(linha[col_mod]).strip()
                                if col_emp is not None and len(linha) > col_emp and linha[col_emp]:
                                    info["empreendimento"] = str(linha[col_emp]).strip()
                                por_id[idr] = info
                    finally:
                        wb.close()
                    self._fontes[caminho] = (mtime, por_id)
                except Exception as exc:
                    self.erros[("fonte", caminho.name)] = f"{type(exc).__name__}: {exc}"[:300]

        combinado: dict[int, dict] = {}
        for mtime, por_id in sorted(self._fontes.values(), key=lambda f: f[0]):
            for idr, info in por_id.items():
                destino = combinado.setdefault(idr, {})
                destino.update({k: v for k, v in info.items() if v})
        return combinado

    def _agrupar(self, fontes: dict[int, dict]) -> list[Reserva]:
        grupos: dict[tuple[str, int], list[Pasta]] = {}
        for pasta in self._pastas.values():
            grupos.setdefault((pasta.tipo, pasta.id), []).append(pasta)
        reservas = []
        for (tipo, id_), pastas in grupos.items():
            pastas.sort(key=_ordem_das_pessoas)
            contexto: dict = {}
            if tipo == TIPO_RESERVA:
                contexto.update(fontes.get(id_, {}))
            for p in reversed(pastas):  # o titular, primeiro da lista, vence
                contexto.update(p.contexto)
            reservas.append(Reserva(tipo=tipo, id=id_, pastas=pastas, contexto=contexto))
        return reservas


# --------------------------------------------------------------------------
# o que vai para a tela
# --------------------------------------------------------------------------

def montar_apont(a: dict, status_atuais: dict[str, dict]) -> dict:
    ev = status_atuais.get(a["chave"])
    em = ev and ev.get("em")
    return {
        "chave": a["chave"], "nivel": a["nivel"], "texto": a["texto"], "categoria": a["categoria"],
        "status": ev["status"] if ev else STATUS_ABERTO,
        "motivo": ev.get("motivo") if ev else None,
        "detalhe": mascarar_texto_livre(ev.get("detalhe")) if ev else None,
        "veredito": ev.get("veredito") if ev else None,
        "responsavel": ev.get("usuario") if ev else None,
        "atualizado_em": formatar_momento(datetime.fromisoformat(em)) if em else None,
    }


def montar_item(r: Reserva, status_atuais: dict[str, dict], modelos_cv: dict | None = None) -> dict:
    """'modelos_cv': {tipo: modelo do link "Abrir no CV"} (ver config.modelo_cv)."""
    contagens = {"divergencias": 0, "avisos": 0, "ausentes": 0}
    abertos = 0
    for p in r.pastas:
        for k, v in p.contagens().items():
            contagens[k] += v
        abertos += sum(1 for a in p.apontamentos
                       if (status_atuais.get(a["chave"]) or {}).get("status", STATUS_ABERTO) == STATUS_ABERTO)
    ctx = r.contexto
    return {
        "tipo": r.tipo, "id": r.id, "titular": r.titular,
        "modalidade": rotulo_modalidade(ctx.get("modalidade")),
        "empreendimento": ctx.get("empreendimento"),
        "imobiliaria": ctx.get("imobiliaria"),
        "corretor": ctx.get("corretor"),
        "situacao": ctx.get("situacao"),
        "pessoas": len(r.pastas),
        **contagens,
        "abertos": abertos,
        "aguardando_leitura": sum(len(p.aguardando_leitura) for p in r.pastas),
        "analisado_em": formatar_momento(r.analisado_em),
        "status": r.status,
        "url_cv": url_cv((modelos_cv or {}).get(r.tipo), r.id),
    }


# Planilha não é documento do cliente: tem o cadastro inteiro (CPF sem máscara)
# e a leitura da IA. Nunca sai pela rota de arquivo.
EXTENSOES_BLOQUEADAS = (".xlsx", ".xlsm", ".xls", ".xltx", ".xltm")


def _nome_simples(nome: str) -> bool:
    """Um nome de arquivo/pasta sem nada que mude de diretório. No Windows,
    ':' abre fluxo alternativo ou unidade, e ponto/espaço no fim some na hora de
    abrir ("x.xlsx." abre "x.xlsx")."""
    if not nome or nome in (".", "..") or len(nome) > 255:
        return False
    if any(c in nome for c in '/\\:*?"<>|') or any(ord(c) < 32 for c in nome):
        return False
    return nome == nome.rstrip(" .")


def caminho_seguro_do_arquivo(saidas: Path, tipo: str, pasta: str, arquivo: str) -> Path | None:
    """O arquivo pedido, se - e só se - for um arquivo comum dentro de
    saidas/<tipo>/<pasta>/ e não for planilha. Qualquer outra coisa: None."""
    if tipo not in TIPOS or not _nome_simples(pasta) or not _nome_simples(arquivo):
        return None
    if arquivo.lower().endswith(EXTENSOES_BLOQUEADAS):
        return None
    try:
        base = (Path(saidas) / tipo).resolve()
        diretorio = (base / pasta).resolve()
        alvo = (diretorio / arquivo).resolve()
    except (OSError, ValueError):
        return None
    if diretorio.parent != base or alvo.parent != diretorio:
        return None
    if alvo.suffix.lower() in EXTENSOES_BLOQUEADAS or not alvo.is_file():
        return None
    return alvo


def url_do_arquivo(pasta: Pasta, arquivo: str | None) -> str | None:
    """'/arquivo/<tipo>/<id>/<token>' - nenhum nome de cliente, CPF ou nome de
    arquivo na URL. None se o arquivo não existe ou não passa nas defesas."""
    if not arquivo or arquivo in ROTULOS_DE_SENTINELA:
        return None
    token = pasta.token_de(str(arquivo))
    if token is None:
        return None
    saidas = pasta.caminho.parent.parent
    if caminho_seguro_do_arquivo(saidas, pasta.tipo, pasta.nome_pasta, str(arquivo)) is None:
        return None
    return f"/arquivo/{pasta.tipo}/{pasta.id}/{token}"


RE_TOKEN = re.compile(r"[0-9a-f]{24}")


def arquivo_pelo_token(reserva: Reserva, token: str) -> tuple[Pasta, str] | None:
    """(pasta, nome do arquivo) do token, dentro desta reserva."""
    if not RE_TOKEN.fullmatch(token or ""):
        return None
    for pasta in reserva.pastas:
        for candidato, arquivo in pasta.tokens.items():
            if hmac.compare_digest(candidato, token):
                return pasta, arquivo
    return None


def _status_do_documento(resultado: dict | None) -> str:
    if resultado is None:
        return STATUS_SEM_ANALISE
    tipo_ia = str((resultado.get("doc_ia") or {}).get("tipo_identificado") or "")
    if tipo_ia in LEITURAS_SEM_IA:
        return STATUS_FORA_DO_ESCOPO
    if resultado.get("divergencias"):
        return STATUS_DIVERGENCIA
    return STATUS_AVISO if resultado.get("avisos") else STATUS_OK


def montar_pessoa(p: Pasta, status_atuais: dict[str, dict]) -> dict:
    por_arquivo: dict[str, list[dict]] = {}
    da_pasta = []
    for a in p.apontamentos:
        apont = montar_apont(a, status_atuais)
        if a["arquivo"] in ROTULOS_DE_SENTINELA:
            da_pasta.append(apont)
        else:
            por_arquivo.setdefault(a["arquivo"], []).append(apont)

    documentos = []
    ja_listados = set()
    aguardando = p.aguardando_leitura
    for d in p.documentos:
        arquivo = d.get("arquivo_local") or ""
        r = p.resultados.get(arquivo) if arquivo else None
        documentos.append({
            # Nome do arquivo baixado do CV costuma trazer o CPF.
            "arquivo": mascarar_no_texto(arquivo or str(d.get("nome") or "")),
            "url": url_do_arquivo(p, arquivo),
            "tipo_cv": d.get("tipo") or (r or {}).get("tipo_cv") or "",
            "tipo_ia": str(((r or {}).get("doc_ia") or {}).get("tipo_identificado") or "") or None,
            "leitura_completa": bool(((r or {}).get("doc_ia") or {}).get("_leitura_completa")),
            "status": (STATUS_AGUARDANDO_LEITURA if arquivo in aguardando else _status_do_documento(r)),
            "apontamentos": por_arquivo.get(arquivo, []) if arquivo else [],
        })
        if arquivo:
            ja_listados.add(arquivo)
    # Leitura gravada de arquivo que não está mais na aba Documentos.
    for arquivo, r in p.resultados.items():
        if arquivo in ja_listados:
            continue
        documentos.append({
            "arquivo": mascarar_no_texto(arquivo), "url": url_do_arquivo(p, arquivo),
            "tipo_cv": r.get("tipo_cv") or "",
            "tipo_ia": str((r.get("doc_ia") or {}).get("tipo_identificado") or "") or None,
            "leitura_completa": bool((r.get("doc_ia") or {}).get("_leitura_completa")),
            "status": _status_do_documento(r), "apontamentos": por_arquivo.get(arquivo, []),
        })

    return {
        "papel": p.papel, "nome": p.nome,
        "documento": mascarar_documento(p.cadastro.get("documento")),
        "pasta": mascarar_no_texto(p.nome_pasta),
        "analisado_em": formatar_momento(p.analisado_em) if p.analisada else None,
        **p.contagens(),
        "status": p.status,
        "aguardando_leitura": len(aguardando),
        "apontamentos_pasta": da_pasta,
        "documentos": documentos,
        "abas_extras": p.abas_extras,
    }


def montar_detalhe(r: Reserva, status_atuais: dict[str, dict], modelos_cv: dict | None = None) -> dict:
    detalhe = montar_item(r, status_atuais, modelos_cv)
    detalhe["pessoas"] = [montar_pessoa(p, status_atuais) for p in r.pastas]
    return detalhe


# --------------------------------------------------------------------------
# filtros da lista
# --------------------------------------------------------------------------

CAMPOS_FILTRO = ("q", "status", "modalidade", "empreendimento", "imobiliaria",
                 "corretor", "situacao", "tratativa")


def filtro_de(params) -> dict:
    filtro = {c: (params.get(c) or "").strip() for c in CAMPOS_FILTRO}
    if filtro["tratativa"] not in ("abertos", "todos"):
        filtro["tratativa"] = "todos"
    return filtro


def _nomes_da_reserva(r: Reserva) -> list[str]:
    return [r.titular] + [p.nome for p in r.pastas]


def filtrar(reservas_e_itens: list[tuple[Reserva, dict]], filtro: dict) -> list[tuple[Reserva, dict]]:
    q = normalizar(filtro.get("q"))
    saida = []
    for r, item in reservas_e_itens:
        if q and q not in str(item["id"]) and not any(q in normalizar(n) for n in _nomes_da_reserva(r)):
            continue
        if filtro.get("status") and item["status"] != filtro["status"]:
            continue
        if any(filtro.get(c) and (item.get(c) or "") != filtro[c]
               for c in ("modalidade", "empreendimento", "imobiliaria", "corretor", "situacao")):
            continue
        if filtro.get("tratativa") == "abertos" and not item["abertos"]:
            continue
        saida.append((r, item))
    saida.sort(key=lambda par: (par[0].analisado_em or datetime.min, par[0].id), reverse=True)
    return saida


def opcoes(itens: list[dict]) -> dict:
    def distintos(campo):
        return sorted({i[campo] for i in itens if i.get(campo)}, key=normalizar)
    return {
        "modalidades": distintos("modalidade"),
        "empreendimentos": distintos("empreendimento"),
        "imobiliarias": distintos("imobiliaria"),
        "corretores": distintos("corretor"),
        "situacoes": distintos("situacao"),
    }


def totais(itens: list[dict]) -> dict:
    return {
        "analisadas": sum(1 for i in itens if i["status"] != STATUS_SEM_ANALISE),
        "com_divergencia": sum(1 for i in itens if i["status"] == STATUS_DIVERGENCIA),
        "so_aviso": sum(1 for i in itens if i["status"] == STATUS_AVISO),
        "ok": sum(1 for i in itens if i["status"] == STATUS_OK),
        "sem_analise": sum(1 for i in itens if i["status"] == STATUS_SEM_ANALISE),
        "apontamentos_abertos": sum(i["abertos"] for i in itens),
    }
