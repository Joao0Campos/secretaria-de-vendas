"""Exportação para Excel: a lista filtrada (abas "Reservas" e "Apontamentos")
e a revisão de regras (aba "Revisao de regras")."""
from __future__ import annotations

import io

from .dados import mascarar_para_o_time

ROTULO_TIPO = {"reserva": "Reserva"}

COLUNAS_RESERVAS = [
    ("tipo", "tipo"), ("id", "id"), ("titular", "titular"), ("modalidade", "modalidade"),
    ("empreendimento", "empreendimento"), ("imobiliaria", "imobiliária"), ("corretor", "corretor"),
    ("situacao", "situação"), ("pessoas", "pessoas"), ("divergencias", "divergências"),
    ("avisos", "avisos"), ("ausentes", "ausentes"), ("abertos", "abertos"),
    ("analisado_em", "analisado em"), ("status", "status"),
]
COLUNAS_APONTAMENTOS = ["reserva", "pessoa", "papel", "documento", "nivel", "categoria", "texto",
                        "status", "motivo", "detalhe", "responsável", "atualizado em"]
COLUNAS_REVISAO = ["categoria", "reserva", "papel", "documento", "IA leu", "nível", "apontamento",
                   "motivo", "detalhe", "quem", "quando", "ainda aparece"]


def _escrever(ws, valores: list) -> None:
    """Uma linha. Texto que começa com '=' é gravado como TEXTO: o openpyxl o
    gravaria como fórmula, e o texto vem de leitura de documento de terceiros."""
    ws.append(valores)
    for cel in ws[ws.max_row]:
        if isinstance(cel.value, str) and cel.value.startswith("="):
            cel.data_type = "s"


def gerar_xlsx(detalhes: list[dict]) -> bytes:
    """'detalhes': saída de dados.montar_detalhe() de cada reserva da lista."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Reservas"
    _escrever(ws, [rotulo for _, rotulo in COLUNAS_RESERVAS])
    for r in detalhes:
        _escrever(ws, [ROTULO_TIPO.get(r["tipo"], r["tipo"]) if campo == "tipo"
                       else (len(r["pessoas"]) if campo == "pessoas" else r.get(campo))
                       for campo, _ in COLUNAS_RESERVAS])

    ws2 = wb.create_sheet("Apontamentos")
    _escrever(ws2, COLUNAS_APONTAMENTOS)
    for r in detalhes:
        reserva = f"{ROTULO_TIPO.get(r['tipo'], r['tipo'])} {r['id']}"
        for p in r["pessoas"]:
            pares = [("", a) for a in p["apontamentos_pasta"]]
            pares += [(d["arquivo"], a) for d in p["documentos"] for a in d["apontamentos"]]
            for documento, a in pares:
                _escrever(ws2, [reserva, p["nome"], p["papel"], documento or "(pasta)", a["nivel"],
                                a["categoria"], a["texto"], a["status"], a["motivo"] or "",
                                a.get("detalhe") or "", a["responsavel"] or "", a["atualizado_em"] or ""])

    return _salvar(wb, (ws, ws2))


def gerar_xlsx_revisao(grupos: list[dict]) -> bytes:
    """'grupos': saída de revisao.montar()["grupos"]. Uma linha por caso, sem
    nome de cliente (só o papel). Este arquivo vai para o time (inclusive
    agentes de IA): texto e "por quê" passam por mascarar_para_o_time()."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Revisao de regras"
    _escrever(ws, COLUNAS_REVISAO)
    for g in grupos:
        for c in g["casos"]:
            _escrever(ws, [g["categoria"], f"{ROTULO_TIPO.get(c['tipo'], c['tipo'])} {c['id']}",
                           c["papel"] or "", c["tipo_cv"] or "(pasta)", c["tipo_ia"] or "", c["nivel"],
                           mascarar_para_o_time(c["texto"], g["categoria"]), c["motivo"] or "",
                           mascarar_para_o_time(c["detalhe"], g["categoria"]) or "", c["responsavel"] or "",
                           c["atualizado_em"] or "", "sim" if c["presente"] else "não"])
    return _salvar(wb, (ws,))


def _salvar(wb, abas) -> bytes:
    from openpyxl.styles import Font

    for aba in abas:
        for cel in aba[1]:
            cel.font = Font(bold=True)
        aba.freeze_panes = "A2"

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()
