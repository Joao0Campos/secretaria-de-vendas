"""Falhas da IA e o disjuntor (21/10/2026, desenho da Turing) - offline.

Rodar (da pasta do projeto ou de qualquer lugar):
    <python do .venv do projeto> -B painel/tests/test_falhas_ia.py -v

O Gemini é um cliente falso que levanta os erros do SDK (code/status/message);
o ia_estado.json, o cache e as planilhas ficam num diretório temporário. Mesma
blindagem do test_painel.
"""
from __future__ import annotations

import sys
from pathlib import Path

_AQUI = Path(__file__).resolve().parent
if str(_AQUI) not in sys.path:
    sys.path.insert(0, str(_AQUI))

import test_cache_ia as tc  # noqa: E402  (blindagem antes do projeto)

import contextlib  # noqa: E402
import dataclasses  # noqa: E402
import hashlib  # noqa: E402
import io  # noqa: E402
import json  # noqa: E402
import os  # noqa: E402
import sqlite3  # noqa: E402
import types  # noqa: E402
import unittest  # noqa: E402
from datetime import datetime, timedelta  # noqa: E402
from unittest import mock  # noqa: E402

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import analisar_documentos as ad  # noqa: E402
import executor_reanalise as ex  # noqa: E402
import mvp_v1_validador as v1  # noqa: E402
import varredura_diaria as vd  # noqa: E402
from painel import __main__ as cli  # noqa: E402
from painel import ia_falhas  # noqa: E402
from painel import sincronia  # noqa: E402
from painel.app import criar_app  # noqa: E402

tp = tc.tp
LEITURA = tc.LEITURA
CHAVE = "chave-de-teste-do-executor-0123456789"
HASH = hashlib.sha256(CHAVE.encode()).hexdigest()


class ErroDaAPI(Exception):
    """Como google.genai.errors.APIError: .code, .status, .message e o str()."""

    def __init__(self, code, status, message):
        self.code, self.status, self.message = code, status, message
        super().__init__(f"{code} {status}. {{'error': {{'code': {code}, 'message': '{message}'}}}}")


SPENDING_CAP = ErroDaAPI(429, "RESOURCE_EXHAUSTED",
                         "Your project has exceeded its monthly spending cap. Please go to AI Studio to raise it.")
POR_MINUTO = ErroDaAPI(429, "RESOURCE_EXHAUSTED",
                       "You exceeded your current quota, please check your plan and billing details. Quota exceeded "
                       "for metric: generate_content_requests, quotaId: GenerateRequestsPerMinutePerProjectPerModel")


class ClienteFalso:
    """generate_content devolve (ou levanta) o próximo da lista; o último se repete."""

    def __init__(self, *respostas):
        self.respostas, self.chamadas = list(respostas), 0
        self.models = types.SimpleNamespace(generate_content=self._gerar)

    def _gerar(self, **k):
        self.chamadas += 1
        r = self.respostas.pop(0) if len(self.respostas) > 1 else self.respostas[0]
        if isinstance(r, BaseException):
            raise r
        return types.SimpleNamespace(text=json.dumps(r), candidates=None, usage_metadata=None)


class Base(tc.BaseCache):
    def setUp(self):
        super().setUp()
        self.estado = self.tmp / "painel_dados" / "ia_estado.json"
        guardados = {n: getattr(ad, n) for n in ("ARQUIVO_ESTADO_IA", "_dormir", "preparar_conteudo_para_ia")}
        self.addCleanup(lambda: [setattr(ad, n, v) for n, v in guardados.items()])
        ad.ARQUIVO_ESTADO_IA = self.estado
        self.esperas = []
        ad._dormir = self.esperas.append
        ad._seguidos_429 = 0
        self.addCleanup(setattr, ad, "_seguidos_429", 0)
        self.addCleanup(ia_falhas._EM_MEMORIA.clear)
        ad.extrair_dados_do_documento = tc.BaseCache_extrair  # a de verdade, com o cliente falso
        ad.preparar_conteudo_para_ia = lambda c: [("application/pdf", b"x")]
        falso = types.SimpleNamespace(
            Part=types.SimpleNamespace(from_bytes=lambda **k: k), GenerateContentConfig=lambda **k: k)
        genai = mock.patch.dict(sys.modules, {"google.genai": types.SimpleNamespace(types=falso),
                                              "google.genai.types": falso})
        genai.start()
        self.addCleanup(genai.stop)

    def com_cliente(self, cliente, paralelismo=1):
        ad.criar_cliente_ia = lambda chave: cliente
        ad.carregar_config = lambda: {"gemini_api_key": "chave-falsa", "paralelismo_ia": paralelismo}
        return cliente

    def pasta_com(self, arquivos, nome="FULANA-FICTICIA-reserva-777", tipo="Comprovante de Endereço"):
        pasta = self.saidas / "reserva" / nome
        pasta.mkdir(parents=True, exist_ok=True)
        v1.gerar_excel({"nome": "FULANA FICTICIA"},
                       [{"tipo": tipo, "nome": a, "arquivo_local": a} for a in arquivos], pasta / f"{nome}.xlsx")
        for a in arquivos:
            (pasta / a).write_bytes(f"%PDF {a}".encode())
        return pasta

    def falhas_da_planilha(self, pasta):
        wb = load_workbook(next(pasta.glob("*.xlsx")))
        if ad.ABA_FALHAS_IA not in wb.sheetnames:
            return []
        return [(l[0], l[2]) for l in wb[ad.ABA_FALHAS_IA].iter_rows(min_row=2, values_only=True)]

    def lidos(self, pasta):
        return ad.ler_resultados_anteriores(next(pasta.glob("*.xlsx")))


# ==========================================================================
# (1) as classes
# ==========================================================================
class Classes(unittest.TestCase):
    def test_mapeamento(self):
        casos = [
            (SPENDING_CAP, "limite_ou_saldo"),
            (ErroDaAPI(429, "RESOURCE_EXHAUSTED", "Prepayment credits are depleted"), "limite_ou_saldo"),
            (ErroDaAPI(429, "RESOURCE_EXHAUSTED", "Quota exceeded, quotaId: GenerateRequestsPerDayPerProject"),
             "limite_ou_saldo"),
            (ErroDaAPI(403, "PERMISSION_DENIED", "Billing account is disabled for this project"), "limite_ou_saldo"),
            (POR_MINUTO, "rate_limit"),  # traz "billing details" e mesmo assim NÃO é saldo
            (ErroDaAPI(429, "RESOURCE_EXHAUSTED", "Resource has been exhausted"), "indeterminado_429"),
            (ErroDaAPI(429, "RESOURCE_EXHAUSTED", "Rate limit exceeded"), "indeterminado_429"),
            # "billing details" sozinho continua NÃO sendo saldo
            (ErroDaAPI(429, "RESOURCE_EXHAUSTED", "You exceeded your current quota, please check your plan and "
                                                  "billing details."), "indeterminado_429"),
            (ErroDaAPI(503, "UNAVAILABLE", "The model is overloaded"), "rede_timeout_5xx"),
            (ErroDaAPI(500, "INTERNAL", "Internal error"), "rede_timeout_5xx"),
            (TimeoutError("read timed out"), "rede_timeout_5xx"),
            (ConnectionError("conexão recusada"), "rede_timeout_5xx"),
            (ErroDaAPI(400, "INVALID_ARGUMENT", "API key not valid. Please pass a valid API key."), "configuracao"),
            (ErroDaAPI(403, "PERMISSION_DENIED", "Permission denied"), "configuracao"),
            (ErroDaAPI(404, "NOT_FOUND", "models/gemini-9 is not found"), "configuracao"),
            (ia_falhas.FalhaIA("resposta_invalida", "x"), "resposta_invalida"),
        ]
        for exc, classe in casos:
            with self.subTest(exc=str(exc)[:60]):
                self.assertEqual(ia_falhas.classificar(exc).classe, classe)

    def test_detalhe_sanitizado_so_para_o_log(self):
        falha = ia_falhas.classificar(ErroDaAPI(400, "INVALID_ARGUMENT",
                                                "API key AIzaSyA1234567890abcdefghijklmno not valid key=segredo"))
        self.assertNotIn("AIzaSy", falha.detalhe)
        self.assertNotIn("segredo", falha.detalhe)
        self.assertNotIn("AIzaSy", falha.mensagem)
        self.assertLessEqual(len(ia_falhas.sanitizar("x" * 5000)), 400)
        with self.assertRaises(ValueError):
            ia_falhas.FalhaIA("inventada")


# ==========================================================================
# (2) e (6): falha nunca vira divergência; não entra no cache; relida depois
# ==========================================================================
class FalhaNaoEDivergencia(Base):
    def test_falha_de_rede_fica_aguardando_e_a_proxima_rodada_le(self):
        pasta = self.pasta_com(["c.pdf"])
        self.com_cliente(ClienteFalso(ErroDaAPI(503, "UNAVAILABLE", "overloaded")))
        saida = self.processar(pasta)
        self.assertIn("aguardando leitura", saida)
        self.assertNotIn("c.pdf", self.lidos(pasta), "nem divergência, nem linha da aba de divergências")
        self.assertEqual(self.falhas_da_planilha(pasta), [("c.pdf", "rede_timeout_5xx")])
        self.assertEqual(list((self.tmp / "cache_ia").rglob("*.json")), [], "o erro não entra no cache")
        self.assertIsNone(ia_falhas.armado(self.estado), "rede não arma o disjuntor")
        cliente = self.com_cliente(ClienteFalso(dict(LEITURA)))
        self.processar(pasta)
        self.assertEqual(cliente.chamadas, 1, "a próxima rodada tenta de novo")
        self.assertIn("c.pdf", self.lidos(pasta))
        self.assertEqual(self.falhas_da_planilha(pasta), [], "a aba some quando ninguém falha")

    def test_falha_antiga_gravada_como_divergencia_e_relida(self):
        pasta = self.pasta_com(["c.pdf"])
        xlsx = next(pasta.glob("*.xlsx"))
        ad.gravar_aba_divergencias(xlsx, [{"arquivo": "c.pdf", "tipo_cv": "Comprovante de Endereço",
                                           "doc_ia": {"erro": "Falha ao chamar a IA: 429 RESOURCE_EXHAUSTED"},
                                           "divergencias": ["Não foi possível ler o documento: Falha..."],
                                           "avisos": []}])
        self.assertEqual(self.lidos(pasta), {})
        cliente = self.com_cliente(ClienteFalso(dict(LEITURA)))
        self.processar(pasta)
        self.assertEqual(cliente.chamadas, 1)

    def test_documento_que_nao_abre_continua_divergencia(self):
        pasta = self.pasta_com(["c.pdf"])
        ad.preparar_conteudo_para_ia = lambda c: None  # tipo não suportado: problema do arquivo
        self.com_cliente(ClienteFalso(dict(LEITURA)))
        self.processar(pasta)
        self.assertIn("não suportado", self.lidos(pasta)["c.pdf"]["divergencias"][0])
        self.assertEqual(self.falhas_da_planilha(pasta), [])

    def test_rate_limit_backoff_limitado_e_nao_arma(self):
        pasta = self.pasta_com(["c.pdf"])
        cliente = self.com_cliente(ClienteFalso(POR_MINUTO, POR_MINUTO, dict(LEITURA)))
        self.processar(pasta)
        self.assertEqual((cliente.chamadas, self.esperas), (3, [5, 20]))
        self.assertIn("c.pdf", self.lidos(pasta))
        cliente = self.com_cliente(ClienteFalso(POR_MINUTO))
        self.esperas.clear()
        self.processar(self.pasta_com(["d.pdf"], nome="OUTRA-FICTICIA-reserva-778"))
        self.assertEqual((cliente.chamadas, len(self.esperas)), (3, 2), "duas novas tentativas, e só")
        self.assertIsNone(ia_falhas.armado(self.estado), "429 por minuto NÃO arma o disjuntor")


# ==========================================================================
# (3) o disjuntor
# ==========================================================================
class Disjuntor(Base):
    def test_saldo_arma_e_corta_as_chamadas_seguintes_da_pasta(self):
        pasta = self.pasta_com(["a.pdf", "b.pdf", "c.pdf"])
        cliente = self.com_cliente(ClienteFalso(SPENDING_CAP, dict(LEITURA)))
        saida = self.processar(pasta)
        self.assertEqual(cliente.chamadas, 1, "submissão limitada: o resto nem é enviado")
        e = ia_falhas.armado(self.estado)
        self.assertEqual((e["classe"], e["motivo"]), ("limite_ou_saldo", ia_falhas.MENSAGENS["limite_ou_saldo"]))
        self.assertNotIn("RESOURCE_EXHAUSTED", json.dumps(e), "o detalhe técnico não vai para o estado")
        self.assertEqual(sorted(self.falhas_da_planilha(pasta)),
                         [("a.pdf", "limite_ou_saldo"), ("b.pdf", "bloqueada"), ("c.pdf", "bloqueada")])
        self.assertEqual(self.lidos(pasta), {})
        self.assertIn("ia-liberar", saida)
        desde = e["desde"]
        self.processar(self.pasta_com(["d.pdf"], nome="OUTRA-FICTICIA-reserva-778"))
        self.assertEqual(cliente.chamadas, 1, "armado: a pasta seguinte não chama a IA")
        self.assertEqual(ia_falhas.armado(self.estado)["desde"], desde, "o 'desde' não muda")

    def test_cache_continua_valendo_com_o_disjuntor_armado(self):
        pasta = self.pasta_com(["c.pdf"])
        self.com_cliente(ClienteFalso(dict(LEITURA)))
        self.processar(pasta)  # lida e guardada no cache
        ia_falhas.armar(self.estado, "limite_ou_saldo")
        cliente = self.com_cliente(ClienteFalso(SPENDING_CAP))
        self.processar(self.pasta_com(["c.pdf"]))  # planilha recriada, mesmo arquivo
        self.assertEqual(cliente.chamadas, 0)
        self.assertEqual(self.lidos(pasta)["c.pdf"]["doc_ia"]["nome_encontrado"], "FULANA FICTICIA")

    def test_corta_varredura_releitura_e_executor(self):
        ia_falhas.armar(self.estado, "limite_ou_saldo")
        cliente = self.com_cliente(ClienteFalso(dict(LEITURA)))
        # varredura (processar_pasta_cliente, como em _analisar_reserva)
        self.processar(self.pasta_com(["v.pdf"], nome="VARREDURA-FICTICIA-reserva-801"))
        # releitura das Telas (--reler-telas): a Tela velha continua valendo
        tela = self.pasta_com(["t.pdf"], nome="TELA-FICTICIA-reserva-802", tipo="Tela de Aprovação")
        ia_falhas.liberar(self.estado)
        self.processar(tela)
        ia_falhas.armar(self.estado, "limite_ou_saldo")
        chamadas = cliente.chamadas
        self.processar(tela, reler_telas=True)
        self.assertEqual(cliente.chamadas, chamadas)
        self.assertIn("t.pdf", self.lidos(tela), "a leitura velha não some quando a releitura falha")
        # executor (reanálise pela API): _analisar -> vd._analisar_reserva -> processar_pasta_cliente
        pasta_ex = self.pasta_com(["e.pdf"], nome="EXECUTOR-FICTICIA-reserva-803")
        with mock.patch.object(vd, "situacao_atual_no_cv", lambda cfg, idr: None), \
                mock.patch.object(vd, "processar_reserva_completa",
                                  lambda idr, devolver_associados=False: ([pasta_ex], [], None)), \
                mock.patch.object(vd, "_ler_listagem", lambda: {"reservas": {}}), \
                mock.patch.object(vd, "carregar_config", lambda: {}), \
                mock.patch.object(ad, "ORIGEM_DO_CONSUMO", ad.ORIGEM_DO_CONSUMO), \
                contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(ex._analisar(803)[0], ex.PEDIDO_CONCLUIDO)
        self.assertEqual(cliente.chamadas, chamadas, "nenhum dos três fluxos chamou a IA")
        self.assertEqual(self.falhas_da_planilha(pasta_ex), [("e.pdf", "bloqueada")])

    def test_recuperacao_pelo_comando_ia_liberar(self):
        pasta = self.pasta_com(["c.pdf"])
        self.com_cliente(ClienteFalso(SPENDING_CAP))
        self.processar(pasta)
        self.assertTrue(ia_falhas.armado(self.estado))
        with mock.patch.object(ia_falhas, "ARQUIVO_NO_PC", self.estado), \
                contextlib.redirect_stdout(io.StringIO()) as saida:
            cli.main(["ia-liberar"])
        self.assertIn("liberado", saida.getvalue())
        self.assertEqual(ia_falhas.estado(self.estado)["estado"], "livre")
        cliente = self.com_cliente(ClienteFalso(dict(LEITURA)))
        self.processar(pasta)
        self.assertEqual(cliente.chamadas, 1)
        self.assertIn("c.pdf", self.lidos(pasta))
        self.assertEqual(self.falhas_da_planilha(pasta), [])

    def test_o_comando_e_o_pc_usam_o_mesmo_arquivo(self):
        self.assertEqual(ia_falhas.ARQUIVO_NO_PC, ad.PASTA_SCRIPT / "painel_dados" / "ia_estado.json")

    def test_estado_atomico_e_ausente_vale_livre_gravado(self):
        self.assertFalse(self.estado.exists())
        self.assertIsNone(ia_falhas.armado(self.estado))
        self.assertEqual(json.loads(self.estado.read_text(encoding="utf-8"))["estado"], "livre",
                         "ausente = livre, explicitamente: o arquivo novo é gravado")
        self.assertTrue(ia_falhas.armar(self.estado, "limite_ou_saldo"))
        self.assertEqual([p.name for p in self.estado.parent.iterdir()], ["ia_estado.json"], "nenhum .parcial")
        with self.assertRaises(ValueError):
            ia_falhas.armar(self.estado, "rede_timeout_5xx")

    def test_estado_ilegivel_e_falha_fechada(self):
        for conteudo in ("{pela metade", '["lista"]', '{"estado": "talvez"}', ""):
            with self.subTest(conteudo=conteudo):
                self.estado.parent.mkdir(parents=True, exist_ok=True)
                self.estado.write_text(conteudo, encoding="utf-8")
                self.assertEqual(ia_falhas.armado(self.estado)["estado"], "ilegivel")
        cliente = self.com_cliente(ClienteFalso(dict(LEITURA)))
        pasta = self.pasta_com(["c.pdf"])
        saida = self.processar(pasta)
        self.assertEqual(cliente.chamadas, 0, "ilegível bloqueia")
        self.assertEqual(self.falhas_da_planilha(pasta), [("c.pdf", "bloqueada")])
        self.assertIn(ia_falhas.MOTIVO_ILEGIVEL, saida, "e loga")
        self.assertEqual(self.estado.read_text(encoding="utf-8"), "", "não é sobrescrito sozinho")
        with mock.patch.object(ia_falhas, "ARQUIVO_NO_PC", self.estado), contextlib.redirect_stdout(io.StringIO()):
            cli.main(["ia-liberar"])
        self.processar(pasta)
        self.assertEqual(cliente.chamadas, 1, "ia-liberar resolve o ilegível")

    def test_armar_sem_gravar_bloqueia_em_memoria_ate_o_fim_do_processo(self):
        pasta = self.pasta_com(["a.pdf", "b.pdf"])
        cliente = self.com_cliente(ClienteFalso(SPENDING_CAP, dict(LEITURA)))
        with mock.patch.object(ia_falhas, "_gravar", lambda caminho, dados: False):
            saida = self.processar(pasta)
            self.assertIn("NÃO foi gravado", saida, "o retorno de armar() não é ignorado")
            self.assertFalse(self.estado.exists())
            self.assertEqual(ia_falhas.armado(self.estado)["estado"], "armado")
            self.processar(self.pasta_com(["d.pdf"], nome="OUTRA-FICTICIA-reserva-778"))
        self.assertEqual(cliente.chamadas, 1, "nada passou depois do limite confirmado")
        self.assertEqual(sorted(self.falhas_da_planilha(pasta)), [("a.pdf", "limite_ou_saldo"), ("b.pdf", "bloqueada")])
        self.processar(self.pasta_com(["e.pdf"], nome="MAIS-FICTICIA-reserva-779"))
        self.assertEqual(cliente.chamadas, 1, "o arquivo ausente não libera o bloqueio em memória")
        ia_falhas.liberar(self.estado)
        self.processar(self.pasta_com(["e.pdf"], nome="MAIS-FICTICIA-reserva-779"))
        self.assertEqual(cliente.chamadas, 2)


class Indeterminado429(Base):
    SEM_MARCADOR = ErroDaAPI(429, "RESOURCE_EXHAUSTED", "Resource has been exhausted (e.g. check quota).")

    def test_sem_retentativa_e_tres_seguidos_armam(self):
        pasta = self.pasta_com(["a.pdf", "b.pdf", "c.pdf", "d.pdf"])
        cliente = self.com_cliente(ClienteFalso(self.SEM_MARCADOR))
        self.processar(pasta)
        self.assertEqual(self.esperas, [], "sem retentativa automática")
        self.assertEqual(cliente.chamadas, 3, "o 4º nem é enviado")
        e = ia_falhas.armado(self.estado)
        self.assertEqual(e["classe"], "indeterminado")
        self.assertIn("confira no AI Studio", e["motivo"])
        self.assertEqual(sorted(self.falhas_da_planilha(pasta)),
                         [("a.pdf", "indeterminado_429"), ("b.pdf", "indeterminado_429"),
                          ("c.pdf", "indeterminado_429"), ("d.pdf", "bloqueada")])
        self.assertEqual(self.lidos(pasta), {})

    def test_sucesso_no_meio_zera_a_contagem(self):
        pasta = self.pasta_com(["a.pdf", "b.pdf", "c.pdf", "d.pdf", "e.pdf"])
        cliente = self.com_cliente(ClienteFalso(self.SEM_MARCADOR, self.SEM_MARCADOR, dict(LEITURA),
                                                self.SEM_MARCADOR, self.SEM_MARCADOR, dict(LEITURA)))
        self.processar(pasta)
        self.assertEqual(cliente.chamadas, 5)
        self.assertIsNone(ia_falhas.armado(self.estado), "não foram 3 seguidos")
        self.assertEqual(len(self.falhas_da_planilha(pasta)), 4)

    def test_billing_details_sozinho_nao_arma(self):
        pasta = self.pasta_com(["a.pdf"])
        self.com_cliente(ClienteFalso(ErroDaAPI(429, "RESOURCE_EXHAUSTED",
                                                "You exceeded your current quota, please check your plan and "
                                                "billing details.")))
        self.processar(pasta)
        self.assertIsNone(ia_falhas.armado(self.estado))
        self.assertEqual(self.falhas_da_planilha(pasta), [("a.pdf", "indeterminado_429")])


# ==========================================================================
# (4) e (5): banner (só admin), estado desatualizado, selo (todos)
# ==========================================================================
class NoSite(tp.Base):
    def setUp(self):
        super().setUp()
        tp.auth.criar_usuario(self.banco, "chefe", "Chefe", "admin", tp.SENHA)
        self.estado = self.cfg.db.parent / "ia_estado.json"

    def cliente_real(self, cfg=None):
        app = criar_app(cfg or self.cfg)  # templates de verdade
        cliente = TestClient(app, follow_redirects=False)
        self.addCleanup(cliente.close)
        return app, cliente

    def entrar(self, cliente, login):
        csrf = cliente.get("/login").text.split('name="csrf" value="')[1].split('"')[0]
        r = cliente.post("/login", data={"login": login, "senha": tp.SENHA, "next": "/", "csrf": csrf})
        self.assertEqual(r.status_code, 303)

    def armar(self, desde="2026-10-21 14:05:00"):
        self.estado.write_text(json.dumps({"estado": "armado", "classe": "limite_ou_saldo", "desde": desde,
                                           "motivo": "IA sem saldo", "atualizado_em": desde}), encoding="utf-8")

    def test_banner_so_para_admin(self):
        self.armar(datetime.now().strftime("%Y-%m-%d 14:05:00"))
        _, admin = self.cliente_real()
        self.entrar(admin, "chefe")
        html = admin.get("/").text
        self.assertIn("IA sem saldo ou limite atingido desde 14:05", html)
        self.assertNotIn("Pode estar desatualizado", html, "no PC o arquivo é lido agora")
        _, gerente = self.cliente_real()
        self.entrar(gerente, "ana")
        self.assertNotIn("data-ia-estado", gerente.get("/").text)
        ia_falhas.liberar(self.estado)
        self.assertNotIn("data-ia-estado", admin.get("/").text, "livre: sem banner")

    def test_banner_indeterminado_e_ilegivel(self):
        _, admin = self.cliente_real()
        self.entrar(admin, "chefe")
        hoje = datetime.now().strftime("%Y-%m-%d 09:10:00")
        self.estado.write_text(json.dumps({"estado": "armado", "classe": "indeterminado", "desde": hoje,
                                           "motivo": "x", "atualizado_em": hoje}), encoding="utf-8")
        html = admin.get("/").text
        self.assertIn("IA bloqueada desde 09:10", html)
        self.assertIn("confira no AI Studio", html)
        self.estado.write_text("{corrompido", encoding="utf-8")
        self.assertIn("Estado da IA ilegível, rode ia-liberar", admin.get("/").text)
        _, gerente = self.cliente_real()
        self.entrar(gerente, "ana")
        self.assertNotIn("data-ia-estado", gerente.get("/").text)

    def test_no_servidor_vem_do_executor_e_avisa_quando_velho(self):
        cfg = dataclasses.replace(self.cfg, executor_chave_hash=HASH)
        app, cliente = self.cliente_real(cfg)
        tp.auth.criar_usuario(app.state.banco, "chefe2", "Chefe", "admin", tp.SENHA)
        armado = {"estado": "armado", "classe": "limite_ou_saldo", "desde": "2026-10-20 09:30:00",
                  "motivo": "IA sem saldo", "atualizado_em": "2026-10-20 09:30:00", "extra": "x" * 999}
        r = cliente.post("/api/executor/estado", json={"modo": None, "ia": armado},
                         headers={"Authorization": f"Bearer {CHAVE}"})
        self.assertEqual(r.status_code, 200)
        guardado = json.loads(app.state.banco.estado_do_executor()["ia"])
        self.assertNotIn("extra", guardado, "o servidor guarda só os campos conhecidos")
        self.entrar(cliente, "chefe2")
        html = cliente.get("/").text
        self.assertIn("IA sem saldo ou limite atingido desde 20/10 09:30", html)
        self.assertNotIn("Pode estar desatualizado", html)
        velho = (datetime.now() - timedelta(minutes=30)).strftime("%Y-%m-%d %H:%M:%S")
        with contextlib.closing(sqlite3.connect(cfg.db)) as con, con:
            con.execute("UPDATE executor_estado SET visto_em = ?", (velho,))
        self.assertIn("Pode estar desatualizado", cliente.get("/").text)
        # executor livre e em dia vence a cópia velha que veio pelo envio
        self.armar("2026-10-19 08:00:00")
        cliente.post("/api/executor/estado", json={"modo": None, "ia": {}}, headers={"Authorization": f"Bearer {CHAVE}"})
        self.assertNotIn("data-ia-estado", cliente.get("/").text)
        # sem sinal do executor: só a cópia do envio, sempre "pode estar desatualizado"
        with contextlib.closing(sqlite3.connect(cfg.db)) as con, con:
            con.execute("DELETE FROM executor_estado")
        html = cliente.get("/").text
        self.assertIn("desde 19/10 08:00", html)
        self.assertIn("Pode estar desatualizado", html)

    def test_executor_manda_o_estado_no_sinal_de_vida(self):
        enviados = []
        fila = ex.FilaRemota(types.SimpleNamespace(post=lambda rota, dados=None: (enviados.append(dados), {})[1]),
                             eco=False)
        with mock.patch.object(ad, "ARQUIVO_ESTADO_IA", self.estado):
            fila.sinal_de_vida("varredura")
            self.armar()
            fila.sinal_de_vida("varredura")
        self.assertEqual((enviados[0]["modo"], enviados[0]["ia"]["estado"]), ("varredura", "livre"))
        self.assertEqual(enviados[1]["ia"]["estado"], "armado")

    def test_vai_pelo_envio_e_e_opcional(self):
        import sincronizar_servidor as ss
        self.assertTrue(sincronia.dado_permitido("ia_estado.json"))
        self.assertIn("ia_estado.json", ss.ARQUIVOS_DE_DADOS)
        self.assertNotIn("ia_estado.json", ss.ARQUIVOS_OBRIGATORIOS)

    def marcar_falha(self, arquivos):
        xlsx = self.planilhas["titular"]
        ad.gravar_aba_falhas_ia(xlsx, [{"arquivo": a, "tipo_cv": "Documentos Diversos",
                                        "falha_ia": {"classe": "limite_ou_saldo", "mensagem": "IA sem saldo",
                                                     "em": "21/10/2026 14:05"}} for a in arquivos])
        os.utime(xlsx, (tp.ONTEM + 7200, tp.ONTEM + 7200))

    def test_selo_e_apontamentos_intocados(self):
        super().entrar()
        antes = self.reserva(id_=1001)
        chaves = sorted(a["chave"] for p in antes["pessoas"] for d in p["documentos"] for a in d["apontamentos"])
        categorias = sorted(a["categoria"] for p in antes["pessoas"] for a in p["apontamentos_pasta"])
        self.marcar_falha(["faltando.pdf", "rg.pdf"])  # rg.pdf tem leitura que vale: não conta
        item = next(i for i in self.itens()["itens"] if i["id"] == 1001)
        self.assertEqual(item["aguardando_leitura"], 1)
        depois = self.reserva(id_=1001)
        self.assertEqual(sorted(a["chave"] for p in depois["pessoas"] for d in p["documentos"]
                                for a in d["apontamentos"]), chaves, "identidade dos apontamentos intocada")
        self.assertEqual((depois["divergencias"], depois["avisos"]), (antes["divergencias"], antes["avisos"]))
        titular = depois["pessoas"][0]
        self.assertEqual(titular["aguardando_leitura"], 1)
        status = {d["arquivo"]: d["status"] for d in titular["documentos"]}
        self.assertEqual(status["faltando.pdf"], "aguardando_leitura")
        self.assertEqual(status["rg.pdf"], "divergencia")
        self.assertEqual(sorted(a["categoria"] for p in depois["pessoas"] for a in p["apontamentos_pasta"]),
                         categorias, "nenhuma categoria nova no painel")
        self.assertEqual(next(i for i in self.itens()["itens"] if i["id"] == 1002)["aguardando_leitura"], 0)

    def test_selo_para_admin_e_nao_admin(self):
        self.marcar_falha(["faltando.pdf"])
        for login in ("ana", "chefe"):
            with self.subTest(login=login):
                _, cliente = self.cliente_real()
                self.entrar(cliente, login)
                self.assertIn("1 documento(s) aguardando leitura", cliente.get("/").text)
                html = cliente.get("/reserva/reserva/1001").text
                self.assertIn("1 documento(s) aguardando leitura", html)
                self.assertIn("Aguardando leitura: a IA falhou", html)
                self.assertNotIn("aguardando leitura", cliente.get("/reserva/reserva/1002").text)


if __name__ == "__main__":
    unittest.main()
