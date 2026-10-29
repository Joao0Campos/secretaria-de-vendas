"""Testes do painel - offline, planilhas SINTÉTICAS em diretório temporário.

Rodar (da pasta do projeto ou de qualquer lugar):
    <python do .venv do projeto> -B painel/tests/test_painel.py -v

As planilhas são geradas com as mesmas funções da auditoria (gerar_excel,
gravar_aba_divergencias, gravar_aba_sugestoes, gravar_aba_associados), para o
teste quebrar se o formato delas mudar. Nomes e documentos são fictícios.
"""
from __future__ import annotations

import sys
from pathlib import Path

_RAIZ = Path(__file__).resolve().parents[2]
if str(_RAIZ) not in sys.path:
    sys.path.insert(0, str(_RAIZ))

import blindagem  # noqa: E402,F401  (primeiro import do projeto)

import _socket  # noqa: E402
import socket  # noqa: E402


def _socketpair_do_proprio_processo(family=socket.AF_INET, type=socket.SOCK_STREAM, proto=0):
    """Exceção ESTREITA à blindagem, só neste arquivo de teste.

    No Windows o asyncio cria o "self-pipe" do event loop com socket.socketpair(),
    que lá é um listener em 127.0.0.1 + um connect nele - e a blindagem bloqueia
    todo connect. Sem event loop o TestClient não sobe. Aqui o connect só é
    feito no listener que esta própria função acabou de abrir, e o par é
    conferido; qualquer outro connect continua bloqueado."""
    ouvinte = socket.socket(family, type, proto)
    try:
        ouvinte.bind(("127.0.0.1", 0))
        ouvinte.listen(1)
        endereco = ouvinte.getsockname()[:2]
        cliente = socket.socket(family, type, proto)
        try:
            _socket.socket.connect(cliente, endereco)
            servidor, _ = ouvinte.accept()
        except BaseException:
            cliente.close()
            raise
    finally:
        ouvinte.close()
    if servidor.getpeername()[:2] != cliente.getsockname()[:2]:
        servidor.close()
        cliente.close()
        raise ConnectionError("par de sockets inesperado")
    return servidor, cliente


socket.socketpair = _socketpair_do_proprio_processo

import contextlib  # noqa: E402
import subprocess  # noqa: E402
from uvicorn.middleware.proxy_headers import ProxyHeadersMiddleware  # noqa: E402
import dataclasses  # noqa: E402
import io  # noqa: E402
import os  # noqa: E402
import re  # noqa: E402
import shutil  # noqa: E402
import tempfile  # noqa: E402
import time  # noqa: E402
import unittest  # noqa: E402
from datetime import datetime, timedelta  # noqa: E402
from unittest import mock  # noqa: E402

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

import analisar_documentos as ad  # noqa: E402
import auditar_reserva as ar  # noqa: E402
import mvp_v1_validador as v1  # noqa: E402
from testar_amostra import categorizar  # noqa: E402

blindagem.neutralizar_config()

from painel import auth, banco as banco_mod, config as config_mod, dados, metricas  # noqa: E402
from painel import __main__ as cli  # noqa: E402
from painel import app as app_mod  # noqa: E402
from painel.app import criar_app  # noqa: E402
from painel.exportar import gerar_xlsx  # noqa: E402

TEMPLATES_FAKE = Path(__file__).resolve().parent / "templates_fake"
SENHA = "senha-de-teste-123"

CPF_TITULAR = "123.456.789-09"      # exemplo de documentação, não é de ninguém
CPF_SEGUNDO = "987.654.321-00"
MSG_NOME = "Nome no documento ('FULANA OUTRA') difere do cadastro ('FULANA FICTICIA') - similaridade 50%"
MSG_PRAZO = "Comprovante de endereço com {} dias (máximo 90) - emitido em 01/01/2026"
MSG_AUSENTE = "Documento obrigatório não encontrado na pasta: 'certidão de estado civil'"
MSG_AVISO_2O = "Nome no documento ('BELTRANO X') difere do cadastro ('BELTRANO FICTICIO') - similaridade 80%"
MSG_CPF = f"CPF no documento ({CPF_SEGUNDO}) difere do cadastro (111.222.333-96)"
ARQUIVO_COM_CPF = f"CPF {CPF_TITULAR}.pdf"
RE_URL_ARQUIVO = r"^/arquivo/reserva/1001/[0-9a-f]{24}$"


def resultado(arquivo, tipo_cv, tipo_ia, divergencias=(), avisos=()):
    return {"arquivo": arquivo, "tipo_cv": tipo_cv,
            "doc_ia": {"tipo_identificado": tipo_ia, "nome_encontrado": "FICTICIO"},
            "divergencias": list(divergencias), "avisos": list(avisos)}


def criar_pasta(saidas: Path, tipo: str, nome_pasta: str, cadastro: dict, documentos: list[dict],
                resultados=None, ausencia=None, avisos_pasta=None, arquivos=(),
                sugestoes=None, associados=None, contexto=None, mtime: float | None = None) -> Path:
    pasta = saidas / tipo / nome_pasta
    pasta.mkdir(parents=True, exist_ok=True)
    xlsx = pasta / f"{nome_pasta}.xlsx"
    v1.gerar_excel(cadastro, documentos, xlsx)
    if resultados is not None:
        ad.gravar_aba_divergencias(xlsx, resultados, ausencia or [], [], avisos_pasta or [])
    if sugestoes is not None:
        ar.gravar_aba_sugestoes(xlsx, sugestoes, 999)
    if associados is not None:
        ar.gravar_aba_associados(xlsx, associados)
    if contexto is not None:
        wb = load_workbook(xlsx)
        ws = wb.create_sheet("Contexto da analise")
        ws.append(["Campo", "Valor"])
        for k, v in contexto.items():
            ws.append([k, v])
        wb.save(xlsx)
    for nome in arquivos:
        (pasta / nome).write_bytes(b"%PDF-1.4 ficticio " + nome.encode())
    if mtime is not None:
        os.utime(xlsx, (mtime, mtime))
    return xlsx


def regravar_divergencias(xlsx: Path, resultados, ausencia=None, mtime=None):
    ad.gravar_aba_divergencias(xlsx, resultados, ausencia or [], [], [])
    mtime = mtime or (xlsx.stat().st_mtime + 3600)
    os.utime(xlsx, (mtime, mtime))


ONTEM = time.time() - 86400

RES_TITULAR = [
    resultado("rg.pdf", "RG, CPF, CNH ou OUTROS", "RG", divergencias=[MSG_NOME]),
    resultado("comprovante.pdf", "Comprovante de Endereço", "Conta de luz",
              divergencias=[MSG_PRAZO.format(118)]),
]


def montar_saidas(saidas: Path) -> dict[str, Path]:
    """Reserva 1001 (titular com divergência + 2º só com aviso), reserva 1004
    com aba de contexto, reserva 1002 ok e reserva 1003 sem análise. E um
    pré-cadastro ANTIGO (2002, com aviso), que o painel não mostra desde 20/10/2026."""
    p = {}
    p["titular"] = criar_pasta(
        saidas, "reserva", "FULANA-FICTICIA-reserva-1001",
        {"nome": "FULANA FICTICIA", "documento": CPF_TITULAR, "papel_na_operacao": "titular",
         "titular_nome": "FULANA FICTICIA", "titular_documento": CPF_TITULAR},
        [{"tipo": "RG, CPF, CNH ou OUTROS", "nome": "rg", "arquivo_local": "rg.pdf"},
         {"tipo": "Comprovante de Endereço", "nome": "comp", "arquivo_local": "comprovante.pdf"},
         {"tipo": "Documentos Diversos", "nome": "sumiu", "arquivo_local": "faltando.pdf"},
         # como o CV entrega: o CPF no nome do arquivo (teste real de 13/10/2026)
         {"tipo": "CPF", "nome": "cpf", "arquivo_local": ARQUIVO_COM_CPF}],
        resultados=RES_TITULAR, ausencia=[MSG_AUSENTE],
        arquivos=("rg.pdf", "comprovante.pdf", ARQUIVO_COM_CPF),
        sugestoes=[{"campo": "documento", "rotulo": "CPF", "valor_atual": CPF_TITULAR,
                    "valor_sugerido": "-", "acao": "conferir", "origem": "rg.pdf", "divergencia": "x"}],
        mtime=ONTEM)
    p["segundo"] = criar_pasta(
        saidas, "reserva", "BELTRANO-FICTICIO-reserva-1001",
        {"nome": "BELTRANO FICTICIO", "documento": CPF_SEGUNDO,
         "papel_na_operacao": "2º Promissário Comprador",
         "titular_nome": "FULANA FICTICIA", "titular_documento": CPF_TITULAR},
        [{"tipo": "RG, CPF, CNH ou OUTROS", "nome": "rg", "arquivo_local": "rg2.png"}],
        resultados=[resultado("rg2.png", "RG, CPF, CNH ou OUTROS", "CNH", avisos=[MSG_AVISO_2O])],
        arquivos=("rg2.png",),
        associados=[{"tipo_associacao": "2º Promissário Comprador", "nome": "BELTRANO FICTICIO",
                     "documento": CPF_SEGUNDO}],
        mtime=ONTEM - 60)
    p["contexto"] = criar_pasta(
        saidas, "reserva", "JOSE-FICTICIO-reserva-1004",
        {"nome": "JOSÉ FICTÍCIO", "documento": "11122233396"},
        [{"tipo": "Comprovante de Endereço", "nome": "c", "arquivo_local": "c.pdf"}],
        resultados=[resultado("c.pdf", "Comprovante de Endereço", "Conta de água")],
        arquivos=("c.pdf",),
        contexto={"modalidade": "Venda Direta", "empreendimento": "Residencial Sintético",
                  "imobiliaria": "Imob Teste", "corretor": "Corretor Teste", "situacao": "Venda Validada"},
        mtime=ONTEM - 120)
    p["ok"] = criar_pasta(
        saidas, "reserva", "CICLANA-FICTICIA-reserva-1002",
        {"nome": "CICLANA FICTICIA", "documento": "00000000000"},
        [{"tipo": "Comprovante de Endereço", "nome": "c", "arquivo_local": "c.pdf"}],
        resultados=[resultado("c.pdf", "Comprovante de Endereço", "Conta de luz")],
        arquivos=("c.pdf",), mtime=ONTEM - 180)
    p["sem_analise"] = criar_pasta(
        saidas, "reserva", "SEM-ANALISE-reserva-1003", {"nome": "SEM ANALISE"},
        [{"tipo": "RG", "nome": "rg", "arquivo_local": "rg.pdf"}], mtime=ONTEM - 240)
    p["pre_cadastro_antigo"] = criar_pasta(
        saidas, "pre_cadastro", "ANTIGA-FICTICIA-2002",
        {"nome": "ANTIGA FICTICIA", "documento": "00000000000"},
        [{"tipo": "Comprovante de Endereço", "nome": "c", "arquivo_local": "c.pdf"}],
        resultados=[resultado("c.pdf", "Comprovante de Endereço", "Conta de luz", avisos=[MSG_AVISO_2O])],
        arquivos=("c.pdf",), mtime=ONTEM - 300)

    # Fontes provisórias de modalidade/empreendimento.
    wb = Workbook()
    ws = wb.active
    ws.title = "Por reserva"
    ws.append(["idreserva", "cliente", "modalidade"])
    ws.append([1001, "FULANA FICTICIA", "Financiamento Bancário"])
    ws.append([1004, "JOSE", "Financiamento Bancário"])  # perde para a aba de contexto
    wb.save(saidas / "amostra-venda-validada.xlsx")
    wb = Workbook()
    ws = wb.active
    ws.title = "Por reserva"
    ws.append(["idreserva", "cliente", "empreendimento", "modalidade"])
    ws.append([1001, "FULANA", "Empreendimento Sintético", "Financiamento Bancário"])
    wb.save(saidas / "varredura-arquivo-repetido-venda-validada.xlsx")

    # Fora de qualquer pasta de pessoa - nunca pode sair pela rota de arquivo.
    (saidas / "reserva" / "segredo.txt").write_text("SEGREDO-RESERVA")
    (saidas.parent / "fora.txt").write_text("SEGREDO-FORA")
    return p


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="painel_teste_"))
        self.saidas = self.tmp / "saidas"
        self.saidas.mkdir()
        self.planilhas = montar_saidas(self.saidas)
        # mostrar_impacto=True: a tela Impacto está oculta por padrão (20/10/2026),
        # mas o código continua e segue testado.
        self.cfg = config_mod.Config(saidas=self.saidas, db=self.tmp / "painel_dados" / "painel.db",
                                     minutos_manuais_por_reserva=30, mostrar_impacto=True)
        self.app = criar_app(self.cfg, templates_dir=TEMPLATES_FAKE)
        self.banco = self.app.state.banco
        auth.criar_usuario(self.banco, "ana", "Ana Teste", "gerente", SENHA)
        self.client = TestClient(self.app, follow_redirects=False)

    def tearDown(self):
        self.client.close()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def tentar_login(self, login="ana", senha=SENHA, next="/", client=None):
        """GET /login (cookie pré-login + csrf no formulário) e POST /login."""
        client = client or self.client
        pagina = client.get("/login")
        csrf = re.search(r'name="csrf" content="([^"]+)"', pagina.text).group(1)
        return client.post("/login", data={"login": login, "senha": senha, "next": next, "csrf": csrf})

    def entrar(self, client=None):
        client = client or self.client
        r = self.tentar_login(client=client)
        self.assertEqual(r.status_code, 303, r.text)
        pagina = client.get("/")
        return re.search(r'name="csrf" content="([^"]+)"', pagina.text).group(1)

    def reserva(self, tipo="reserva", id_=1001):
        r = self.client.get(f"/api/reserva/{tipo}/{id_}")
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def itens(self, **params):
        r = self.client.get("/api/reservas", params=params)
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def apont(self, detalhe, texto_parcial):
        for p in detalhe["pessoas"]:
            for a in p["apontamentos_pasta"] + [a for d in p["documentos"] for a in d["apontamentos"]]:
                if texto_parcial in a["texto"]:
                    return a
        self.fail(f"apontamento com {texto_parcial!r} não encontrado")


# ==========================================================================
# Leitura de saidas/: contagens, status, papéis, máscara, modalidade
# ==========================================================================
class Leitura(Base):
    def setUp(self):
        super().setUp()
        self.entrar()

    def test_contagens_e_status_por_reserva(self):
        por_id = {(i["tipo"], i["id"]): i for i in self.itens()["itens"]}
        r = por_id[("reserva", 1001)]
        self.assertEqual((r["divergencias"], r["avisos"], r["ausentes"]), (3, 1, 1))
        self.assertEqual((r["pessoas"], r["abertos"], r["status"]), (2, 4, "divergencia"))
        self.assertEqual(r["titular"], "FULANA FICTICIA")
        self.assertEqual(por_id[("reserva", 1004)]["status"], "ok")
        self.assertEqual(por_id[("reserva", 1002)]["status"], "ok")
        self.assertEqual(por_id[("reserva", 1003)]["status"], "sem_analise")
        self.assertIsNone(por_id[("reserva", 1003)]["analisado_em"])
        self.assertRegex(r["analisado_em"], r"^\d{2}/\d{2}/\d{4} \d{2}:\d{2}$")

    def test_totais(self):
        t = self.itens()["totais"]
        self.assertEqual(t, {"analisadas": 3, "com_divergencia": 1, "so_aviso": 0, "ok": 2,
                             "sem_analise": 1, "apontamentos_abertos": 4})

    def test_ordem_por_analisado_em_desc(self):
        ids = [i["id"] for i in self.itens()["itens"]]
        self.assertEqual(ids, [1001, 1004, 1002, 1003])

    def test_titular_primeiro_e_cpf_mascarado(self):
        d = self.reserva()
        self.assertEqual([p["papel"] for p in d["pessoas"]], ["titular", "2º Promissário Comprador"])
        self.assertEqual(d["pessoas"][0]["documento"], "***.456.789-**")
        self.assertEqual(d["pessoas"][1]["documento"], "***.654.321-**")
        self.assertNotIn("123.456.789-09", str(d))
        self.assertNotIn("12345678909", str(d))

    def test_cpf_do_nome_do_arquivo_nao_sai_nem_na_url_nem_na_tela(self):
        d = self.reserva()
        docs = {doc["tipo_cv"]: doc for doc in d["pessoas"][0]["documentos"]}
        self.assertEqual(docs["CPF"]["arquivo"], "CPF ***.456.789-**.pdf")
        self.assertRegex(docs["CPF"]["url"], RE_URL_ARQUIVO)
        self.assertNotIn("FULANA", docs["CPF"]["url"])
        self.assertNotIn("CPF", docs["CPF"]["url"])
        for cru in ("123.456.789-09", "12345678909"):
            self.assertNotIn(cru, str(d))
        with self.banco.conexao() as con:
            arquivos = [linha[0] for linha in con.execute("SELECT arquivo FROM apontamentos")]
        self.assertFalse(any("123.456.789-09" in a for a in arquivos))

    def test_cpf_mascarado_nas_abas_extras_e_nos_textos(self):
        d = self.reserva()
        extras = {e["titulo"]: e["linhas"] for p in d["pessoas"] for e in p["abas_extras"]}
        self.assertIn("Sugestoes (NAO aplicadas)", extras)
        self.assertIn("Associados", extras)
        self.assertTrue(extras["Sugestoes (NAO aplicadas)"][0][0].startswith("*** NADA FOI ALTERADO"))
        self.assertNotIn(CPF_SEGUNDO, str(extras))
        self.assertIn("***.654.321-**", str(extras["Associados"]))
        self.assertEqual(dados.mascarar_no_texto(MSG_CPF),
                         "CPF no documento (***.654.321-**) difere do cadastro (***.222.333-**)")
        self.assertEqual(dados.mascarar_documento("12.345.678/0001-95"), "**.345.678/0001-**")

    def test_sentinela_de_documento_nao_anexado(self):
        titular = self.reserva()["pessoas"][0]
        self.assertEqual([a["texto"] for a in titular["apontamentos_pasta"]], [MSG_AUSENTE])
        self.assertEqual(titular["apontamentos_pasta"][0]["categoria"], "Documento obrigatório não anexado")
        self.assertEqual(titular["apontamentos_pasta"][0]["nivel"], "divergencia")

    def test_documentos_url_e_arquivo_inexistente(self):
        docs = {d["arquivo"]: d for d in self.reserva()["pessoas"][0]["documentos"]}
        self.assertRegex(docs["rg.pdf"]["url"], RE_URL_ARQUIVO)
        self.assertNotEqual(docs["rg.pdf"]["url"], docs["comprovante.pdf"]["url"])
        self.assertIsNone(docs["faltando.pdf"]["url"])
        self.assertEqual(docs["faltando.pdf"]["status"], "sem_analise")
        self.assertEqual(docs["rg.pdf"]["tipo_ia"], "RG")
        self.assertTrue(docs["rg.pdf"]["leitura_completa"])
        self.assertEqual(docs["rg.pdf"]["status"], "divergencia")
        self.assertEqual(docs["rg.pdf"]["apontamentos"][0]["categoria"], "Nome diferente do cadastro")

    def test_pasta_sem_aba_de_divergencias(self):
        d = self.reserva("reserva", 1003)
        self.assertEqual(d["status"], "sem_analise")
        self.assertEqual(d["pessoas"][0]["status"], "sem_analise")
        self.assertEqual(d["pessoas"][0]["documentos"][0]["status"], "sem_analise")

    def test_modalidade_e_contexto(self):
        por_id = {i["id"]: i for i in self.itens()["itens"]}
        self.assertEqual(por_id[1001]["modalidade"], "Financiamento")
        self.assertEqual(por_id[1001]["empreendimento"], "Empreendimento Sintético")
        # A aba "Contexto da analise" vence a planilha de amostra.
        self.assertEqual(por_id[1004]["modalidade"], "Venda direta")
        self.assertEqual(por_id[1004]["empreendimento"], "Residencial Sintético")
        self.assertEqual(por_id[1004]["corretor"], "Corretor Teste")
        self.assertIsNone(por_id[1002]["modalidade"])
        self.assertEqual(self.itens()["opcoes"]["modalidades"], ["Financiamento", "Venda direta"])

    def test_404(self):
        self.assertEqual(self.client.get("/reserva/reserva/9999").status_code, 404)
        self.assertEqual(self.client.get("/api/reserva/reserva/9999").status_code, 404)
        self.assertEqual(self.client.get("/reserva/outro/1001").status_code, 404)
        self.assertEqual(self.client.get("/reserva/reserva/abc").status_code, 422)

    def test_paginas_renderizam(self):
        r = self.client.get("/")
        self.assertEqual(r.status_code, 200)
        self.assertIn("usuario=ana", r.text)
        self.assertIn("itens=4", r.text)
        r = self.client.get("/reserva/reserva/1001")
        self.assertEqual(r.status_code, 200)
        self.assertIn("RESERVA 1001 FULANA FICTICIA <titular> <2º Promissário Comprador>", r.text)
        self.assertIn("periodo=30", self.client.get("/impacto").text)
        self.assertIn("periodo=7", self.client.get("/impacto?periodo=7").text)
        self.assertIn("periodo=30", self.client.get("/impacto?periodo=xyz").text)

    def test_cache_nao_rele_planilha_sem_mudanca(self):
        repo = self.app.state.repo
        repo.atualizar()
        with mock.patch.object(dados, "ler_pasta", wraps=dados.ler_pasta) as espiao:
            repo.atualizar()
            self.assertEqual(espiao.call_count, 0)
            regravar_divergencias(self.planilhas["ok"], [])
            repo.atualizar()
            self.assertEqual(espiao.call_count, 1)

    def test_planilha_aberta_no_excel_e_ignorada(self):
        trava = self.planilhas["titular"].parent / "~$FULANA-FICTICIA-reserva-1001.xlsx"
        trava.write_bytes(b"lixo")
        self.assertEqual(self.itens()["totais"]["com_divergencia"], 1)


# ==========================================================================
# Filtros
# ==========================================================================
class Filtros(Base):
    def setUp(self):
        super().setUp()
        self.entrar()

    def ids(self, **params):
        return sorted(i["id"] for i in self.itens(**params)["itens"])

    def test_q_por_numero_e_nome_sem_acento(self):
        self.assertEqual(self.ids(q="1001"), [1001])
        self.assertEqual(self.ids(q="jose ficticio"), [1004])
        self.assertEqual(self.ids(q="JOSÉ"), [1004])
        self.assertEqual(self.ids(q="beltrano"), [1001])  # nome do 2º comprador
        self.assertEqual(self.ids(q="1003"), [1003])
        self.assertEqual(self.ids(q="2002"), [], "pré-cadastro antigo não aparece")

    def test_filtros_de_campo(self):
        self.assertEqual(self.ids(tipo="pre_cadastro"), [1001, 1002, 1003, 1004], "sem filtro de tipo")
        self.assertEqual(self.ids(status="ok"), [1002, 1004])
        self.assertEqual(self.ids(modalidade="Financiamento"), [1001])
        self.assertEqual(self.ids(empreendimento="Residencial Sintético"), [1004])
        self.assertEqual(self.ids(situacao="Venda Validada"), [1004])
        self.assertEqual(self.ids(tratativa="abertos"), [1001])
        self.assertEqual(self.itens(tratativa="xyz")["filtro"]["tratativa"], "todos")


# ==========================================================================
# Login, sessão e CSRF
# ==========================================================================
class Login(Base):
    def test_sem_login_redireciona_com_next(self):
        r = self.client.get("/")
        self.assertEqual((r.status_code, r.headers["location"]), (303, "/login?next=%2F"))
        r = self.client.get("/reserva/reserva/1001?x=1")
        self.assertEqual(r.headers["location"], "/login?next=%2Freserva%2Freserva%2F1001%3Fx%3D1")
        for caminho in ("/impacto", "/exportar.xlsx", "/arquivo/reserva/1001/" + "0" * 24):
            self.assertEqual(self.client.get(caminho).status_code, 303, caminho)

    def test_api_sem_login_e_401(self):
        self.assertEqual(self.client.get("/api/reservas").status_code, 401)
        self.assertEqual(self.client.post("/api/tratativa", json={}).status_code, 401)

    def test_pedido_do_icone_nao_invalida_o_formulario_de_login(self):
        # 13/10/2026, Chrome do João: logo depois da página de login o navegador
        # pede /favicon.ico; protegido, ele ia para /login, que emitia outro
        # cookie pré-login, e todo envio dava "A página de login expirou".
        pagina = self.client.get("/login", headers={"sec-fetch-dest": "document"})
        csrf = re.search(r'name="csrf" content="([^"]+)"', pagina.text).group(1)
        icone = self.client.get("/favicon.ico", headers={
            "sec-fetch-dest": "image", "accept": "image/avif,image/webp,image/*,*/*;q=0.8"})
        self.assertEqual(icone.status_code, 204)
        self.client.get("/login")  # outra aba abrindo o login
        r = self.client.post("/login", data={"login": "ana", "senha": SENHA, "next": "/", "csrf": csrf})
        self.assertEqual(r.status_code, 303, r.text)

    def test_sem_login_so_navegacao_vai_para_o_login(self):
        self.assertEqual(self.client.get("/", headers={"sec-fetch-dest": "document"}).status_code, 303)
        self.assertEqual(self.client.get("/", headers={"sec-fetch-dest": "image"}).status_code, 401)
        self.assertEqual(self.client.get("/", headers={"accept": "image/png"}).status_code, 401)

    def test_pre_login_perto_de_expirar_nao_e_reaproveitado(self):
        segredo = self.banco.segredo
        novo = auth.emitir_pre_login(segredo)
        self.assertTrue(auth.pre_login_ainda_vale(segredo, novo))
        with mock.patch.object(auth.time, "time", return_value=time.time() - 14 * 60):
            velho = auth.emitir_pre_login(segredo)
        self.assertFalse(auth.pre_login_ainda_vale(segredo, velho))
        self.assertTrue(auth.pre_login_valido(segredo, velho, velho))  # ainda aceito se já estava no formulário
        self.assertFalse(auth.pre_login_ainda_vale(segredo, novo[:-1] + ("0" if novo[-1] != "0" else "1")))

    def test_login_e_static_sao_publicos(self):
        r = self.client.get("/login?next=/impacto")
        self.assertEqual(r.status_code, 200)
        self.assertIn("next=/impacto", r.text)
        self.assertNotEqual(self.client.get("/static/x.css").status_code, 303)

    def test_senha_errada(self):
        r = self.tentar_login(senha="errada")
        self.assertEqual(r.status_code, 401)
        self.assertIn("Login ou senha inválidos", r.text)
        self.assertNotIn(auth.COOKIE_SESSAO, r.cookies)
        self.assertIn(auth.COOKIE_PRE_LOGIN, r.cookies)  # a página volta com um pré-login novo
        r = self.tentar_login(login="ninguem")
        self.assertEqual(r.status_code, 401)

    def test_login_ok_cookie_e_destino(self):
        r = self.tentar_login(login="ANA", next="/impacto?periodo=7")
        self.assertEqual((r.status_code, r.headers["location"]), (303, "/impacto?periodo=7"))
        cookies = r.headers.get_list("set-cookie")
        sessao = next(c for c in cookies if c.startswith(auth.COOKIE_SESSAO + "="))
        self.assertIn("HttpOnly", sessao)
        self.assertIn("SameSite=lax", sessao)
        self.assertIn("Max-Age=43200", sessao)
        self.assertNotIn("Secure", sessao)
        # o pré-login é apagado depois do login
        self.assertTrue(any(c.startswith(auth.COOKIE_PRE_LOGIN + "=") and "Max-Age=0" in c for c in cookies))
        self.assertEqual(self.client.get("/impacto").status_code, 200)

    def test_next_externo_vira_raiz(self):
        for ruim in ("//exemplo.com", "https://exemplo.com", "/\\exemplo.com", "exemplo.com", "/login"):
            self.assertEqual(auth.destino_seguro(ruim), "/", ruim)
        self.assertEqual(auth.destino_seguro("/reserva/reserva/1?x=1"), "/reserva/reserva/1?x=1")
        r = self.tentar_login(next="//exemplo.com")
        self.assertEqual(r.headers["location"], "/")

    def test_login_exige_o_par_do_pre_login(self):
        r = self.client.post("/login", data={"login": "ana", "senha": SENHA, "next": "/"})
        self.assertEqual(r.status_code, 403)              # sem cookie e sem campo
        pagina = self.client.get("/login")
        cookie_pre = self.client.cookies.get(auth.COOKIE_PRE_LOGIN)
        self.assertIn(f'content="{cookie_pre}"', pagina.text)  # mesmo valor no formulário
        r = self.client.post("/login", data={"login": "ana", "senha": SENHA, "csrf": "outro"})
        self.assertEqual(r.status_code, 403)              # campo diferente do cookie
        self.client.cookies.delete(auth.COOKIE_PRE_LOGIN)
        r = self.client.post("/login", data={"login": "ana", "senha": SENHA, "csrf": cookie_pre})
        self.assertEqual(r.status_code, 403)              # campo sem o cookie
        self.assertNotIn(auth.COOKIE_SESSAO, self.client.cookies)

    def test_pre_login_vencido_ou_forjado(self):
        segredo = self.banco.segredo
        bom = auth.emitir_pre_login(segredo)
        self.assertTrue(auth.pre_login_valido(segredo, bom, bom))
        self.assertFalse(auth.pre_login_valido(b"outro-segredo", bom, bom))
        emitido, aleatorio, _ = bom.split(".")
        forjado = f"{emitido}.{aleatorio}.{'0' * 32}"
        self.assertFalse(auth.pre_login_valido(segredo, forjado, forjado))
        with mock.patch.object(auth.time, "time", return_value=time.time() + 16 * 60):
            self.assertFalse(auth.pre_login_valido(segredo, bom, bom))
        self.assertFalse(auth.pre_login_valido(segredo, "", ""))
        self.assertFalse(auth.pre_login_valido(segredo, "a.b", "a.b"))

    def test_senha_guardada_com_pbkdf2(self):
        u = self.banco.buscar_usuario("ana")
        self.assertEqual(u["iteracoes"], 390000)
        self.assertNotIn(SENHA, str(u))
        self.assertTrue(auth.senha_confere(SENHA, u["hash"], u["sal"], u["iteracoes"]))

    def test_sessao_guarda_hash_do_token_e_expira(self):
        self.entrar()
        token = self.client.cookies.get(auth.COOKIE_SESSAO)
        with self.banco.conexao() as con:
            guardados = [linha["token"] for linha in con.execute("SELECT token FROM sessoes")]
        self.assertNotIn(token, guardados)
        with self.banco.conexao() as con:
            con.execute("UPDATE sessoes SET expira_em = '2000-01-01 00:00:00'")
        self.assertEqual(self.client.get("/").status_code, 303)

    def test_csrf_tratativa(self):
        csrf = self.entrar()
        chave = self.apont(self.reserva(), "similaridade 50%")["chave"]
        corpo = {"chave": chave, "status": "corrigido"}
        self.assertEqual(self.client.post("/api/tratativa", json=corpo).status_code, 403)
        self.assertEqual(self.client.post("/api/tratativa", json=corpo,
                                          headers={"X-CSRF-Token": "errado"}).status_code, 403)
        self.assertEqual(self.client.post("/api/tratativa", json=corpo,
                                          headers={"X-CSRF-Token": csrf}).status_code, 200)

    def test_logout_exige_csrf_e_apaga_sessao(self):
        csrf = self.entrar()
        self.assertEqual(self.client.post("/logout", data={}).status_code, 403)
        self.assertEqual(self.client.post("/logout", data={"csrf": "errado"}).status_code, 403)
        r = self.client.post("/logout", data={"csrf": csrf})
        self.assertEqual((r.status_code, r.headers["location"]), (303, "/login"))
        with self.banco.conexao() as con:
            self.assertEqual(con.execute("SELECT COUNT(*) FROM sessoes").fetchone()[0], 0)
        self.assertEqual(self.client.get("/").status_code, 303)

    def test_csrf_de_outra_sessao_nao_vale(self):
        csrf_a = self.entrar()
        outro = TestClient(self.app, follow_redirects=False)
        self.entrar(outro)
        chave = self.apont(self.reserva(), "similaridade 50%")["chave"]
        r = outro.post("/api/tratativa", json={"chave": chave, "status": "corrigido"},
                       headers={"X-CSRF-Token": csrf_a})
        outro.close()
        self.assertEqual(r.status_code, 403)

    def test_cabecalhos_de_seguranca(self):
        self.assertIn("frame-ancestors 'self'", self.client.get("/login").headers["content-security-policy"])
        self.entrar()
        for caminho in ("/", "/reserva/reserva/1001", "/impacto"):
            r = self.client.get(caminho)
            self.assertEqual(r.headers["x-content-type-options"], "nosniff")
            self.assertEqual(r.headers["cache-control"], "no-store")
            self.assertEqual(r.headers["content-security-policy"], app_mod.CSP, caminho)
            self.assertIn("script-src 'self';", r.headers["content-security-policy"])
            self.assertNotIn("strict-transport-security", r.headers)


# ==========================================================================
# Limite de tentativas de login
# ==========================================================================
class LimiteDeTentativas(Base):
    def setUp(self):
        super().setUp()
        auth.criar_usuario(self.banco, "bia", "Bia Teste", "registro", SENHA)

    def cliente(self, ip):
        c = TestClient(self.app, follow_redirects=False, client=(ip, 50000))
        self.addCleanup(c.close)
        return c

    def test_5_falhas_do_login_bloqueiam_ate_com_a_senha_certa(self):
        for i in range(5):
            c = self.cliente(f"10.0.0.{i}")   # IPs diferentes: só o contador do login sobe
            self.assertEqual(self.tentar_login(senha="errada", client=c).status_code, 401)
        c = self.cliente("10.0.0.99")
        r = self.tentar_login(client=c)
        self.assertEqual(r.status_code, 401)
        self.assertIn("Login ou senha inválidos", r.text)  # a mesma mensagem da senha errada
        self.assertEqual(self.tentar_login(login="bia", client=c).status_code, 303)  # outro login segue

    def test_5_falhas_do_ip_bloqueiam_qualquer_login(self):
        c = self.cliente("10.0.0.7")
        for login in ("x1", "x2", "x3", "x4", "x5"):
            self.assertEqual(self.tentar_login(login=login, client=c).status_code, 401)
        self.assertEqual(self.tentar_login(login="bia", client=c).status_code, 401)
        self.assertEqual(self.tentar_login(login="bia", client=self.cliente("10.0.0.8")).status_code, 303)

    def test_bloqueio_acaba_depois_de_15_minutos(self):
        c = self.cliente("10.0.0.7")
        for _ in range(5):
            self.tentar_login(senha="errada", client=c)
        self.assertEqual(self.tentar_login(client=c).status_code, 401)
        antes = banco_mod.para_texto(datetime.now() - timedelta(minutes=16))
        with self.banco.conexao() as con:
            con.execute("UPDATE tentativas_login SET em = ?", (antes,))
        self.assertEqual(self.tentar_login(client=c).status_code, 303)

    def test_login_certo_zera_o_contador_do_login(self):
        for i in range(4):
            self.tentar_login(senha="errada", client=self.cliente(f"10.0.1.{i}"))
        self.assertEqual(self.tentar_login(client=self.cliente("10.0.1.9")).status_code, 303)
        self.assertEqual(self.banco.falhas_recentes("ana", "-", datetime.now() - timedelta(hours=1))[0], 0)
        for i in range(4):
            self.tentar_login(senha="errada", client=self.cliente(f"10.0.2.{i}"))
        self.assertEqual(self.tentar_login(client=self.cliente("10.0.2.9")).status_code, 303)

    def test_bloqueado_gasta_o_mesmo_pbkdf2(self):
        for _ in range(5):
            auth.autenticar_com_limite(self.banco, "ana", "errada", "10.9.9.9")
        with mock.patch.object(auth, "gerar_hash", wraps=auth.gerar_hash) as espiao:
            self.assertIsNone(auth.autenticar_com_limite(self.banco, "ana", SENHA, "10.9.9.9"))
        self.assertEqual(espiao.call_count, 1)


# ==========================================================================
# Etapa B: SECRETARIA_COOKIE_SEGURO=1
# ==========================================================================
class CookieSeguro(Base):
    def setUp(self):
        super().setUp()
        cfg = config_mod.Config(saidas=self.saidas, db=self.tmp / "painel_dados" / "painel.db",
                                cookie_seguro=True)
        self.app = criar_app(cfg, templates_dir=TEMPLATES_FAKE)
        self.client.close()
        self.client = TestClient(self.app, base_url="https://testserver", follow_redirects=False)

    def test_secure_e_hsts(self):
        pagina = self.client.get("/login")
        self.assertIn("Secure", pagina.headers["set-cookie"])
        self.assertEqual(pagina.headers["strict-transport-security"], app_mod.HSTS)
        r = self.tentar_login()
        self.assertEqual(r.status_code, 303)
        sessao = next(c for c in r.headers.get_list("set-cookie") if c.startswith(auth.COOKIE_SESSAO + "="))
        self.assertIn("Secure", sessao)

    def test_hsts_em_toda_resposta_e_cookies_apagados_com_secure(self):
        self.assertEqual(self.client.get("/api/reservas").headers["strict-transport-security"], app_mod.HSTS)
        self.assertEqual(self.client.get("/").headers["strict-transport-security"], app_mod.HSTS)
        csrf = self.entrar()
        r = self.client.post("/logout", data={"csrf": csrf})
        apagado = next(c for c in r.headers.get_list("set-cookie") if c.startswith(auth.COOKIE_SESSAO + "="))
        self.assertIn("Secure", apagado)
        self.assertEqual(r.headers["strict-transport-security"], app_mod.HSTS)

    def test_padrao_do_pc_sem_secure_nem_hsts(self):
        self.client.close()
        self.client = TestClient(criar_app(self.cfg_padrao(), templates_dir=TEMPLATES_FAKE), follow_redirects=False)
        pagina = self.client.get("/login")
        self.assertNotIn("Secure", pagina.headers["set-cookie"])
        self.assertNotIn("strict-transport-security", pagina.headers)
        self.assertNotIn("strict-transport-security", self.client.get("/api/reservas").headers)

    def cfg_padrao(self):
        return config_mod.Config(saidas=self.saidas, db=self.tmp / "painel_dados" / "painel.db")

    def test_variavel_de_ambiente(self):
        with mock.patch.dict(os.environ, {"SECRETARIA_COOKIE_SEGURO": "1"}):
            self.assertTrue(config_mod.carregar().cookie_seguro)
        with mock.patch.dict(os.environ, {"SECRETARIA_COOKIE_SEGURO": "0"}):
            self.assertFalse(config_mod.carregar().cookie_seguro)


# ==========================================================================
# Tratativa
# ==========================================================================
class Tratativa(Base):
    def setUp(self):
        super().setUp()
        self.csrf = self.entrar()
        self.chave = self.apont(self.reserva(), "similaridade 50%")["chave"]

    def tratar(self, **corpo):
        return self.client.post("/api/tratativa", json={"chave": self.chave, **corpo},
                                headers={"X-CSRF-Token": self.csrf})

    def test_corrigido(self):
        r = self.tratar(status="corrigido")
        self.assertEqual(r.status_code, 200, r.text)
        a = r.json()
        self.assertEqual((a["status"], a["responsavel"], a["chave"]), ("corrigido", "ana", self.chave))
        self.assertRegex(a["atualizado_em"], r"^\d{2}/\d{2}/\d{4} \d{2}:\d{2}$")
        self.assertEqual(self.apont(self.reserva(), "similaridade 50%")["status"], "corrigido")
        self.assertEqual(self.itens()["totais"]["apontamentos_abertos"], 3)

    def test_ignorado_exige_motivo(self):
        self.assertEqual(self.tratar(status="ignorado").status_code, 400)
        self.assertEqual(self.tratar(status="ignorado", motivo="porque sim").status_code, 400)
        self.assertEqual(self.tratar(status="ignorado", motivo="Outro: ").status_code, 400)
        r = self.tratar(status="ignorado", motivo="Falso positivo da IA")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["motivo"], "Falso positivo da IA")
        r = self.tratar(status="ignorado", motivo="Outro: documento conferido no papel")
        self.assertEqual(r.json()["motivo"], "Outro: documento conferido no papel")

    def test_status_invalidos(self):
        for status in ("resolvido_na_reanalise", "qualquer", ""):
            self.assertEqual(self.tratar(status=status).status_code, 400, status)
        r = self.client.post("/api/tratativa", content=b"nao e json",
                             headers={"X-CSRF-Token": self.csrf, "content-type": "application/json"})
        self.assertEqual(r.status_code, 400)
        r = self.client.post("/api/tratativa", json={"chave": "nao-existe", "status": "corrigido"},
                             headers={"X-CSRF-Token": self.csrf})
        self.assertEqual(r.status_code, 404)

    def test_reabrir(self):
        self.tratar(status="corrigido")
        r = self.tratar(status="aberto")
        self.assertEqual(r.json()["status"], "aberto")
        self.assertIsNone(r.json()["motivo"])


# ==========================================================================
# Sincronização com a reanálise
# ==========================================================================
class Sincronizacao(Base):
    def setUp(self):
        super().setUp()
        self.csrf = self.entrar()

    def test_some_na_reanalise_e_resolvido_e_reaparece_aberto(self):
        chave = self.apont(self.reserva(), "similaridade 50%")["chave"]
        xlsx = self.planilhas["titular"]
        sem_nome = [resultado("rg.pdf", "RG, CPF, CNH ou OUTROS", "RG"), RES_TITULAR[1]]
        regravar_divergencias(xlsx, sem_nome, [MSG_AUSENTE])
        self.reserva()
        a = self.banco.buscar_apontamento(chave)
        self.assertEqual(a["presente"], 0)
        ev = self.banco.status_atuais()[chave]
        self.assertEqual((ev["status"], ev["usuario"]), ("resolvido_na_reanalise", "sistema"))
        # o momento é o da planilha nova, não o da sincronização
        self.assertEqual(banco_mod.de_texto(ev["em"]),
                         datetime.fromtimestamp(xlsx.stat().st_mtime).replace(microsecond=0))

        regravar_divergencias(xlsx, RES_TITULAR, [MSG_AUSENTE])
        detalhe = self.reserva()
        self.assertEqual(self.banco.buscar_apontamento(chave)["presente"], 1)
        self.assertEqual(self.apont(detalhe, "similaridade 50%")["status"], "aberto")

    def test_tratado_nao_vira_resolvido_e_ignorado_nao_reabre(self):
        detalhe = self.reserva()
        corrigido = self.apont(detalhe, "similaridade 50%")["chave"]
        ignorado = self.apont(detalhe, "118 dias")["chave"]
        h = {"X-CSRF-Token": self.csrf}
        self.client.post("/api/tratativa", json={"chave": corrigido, "status": "corrigido"}, headers=h)
        self.client.post("/api/tratativa", json={"chave": ignorado, "status": "ignorado",
                                                 "motivo": "Falso positivo da IA"}, headers=h)
        xlsx = self.planilhas["titular"]
        regravar_divergencias(xlsx, [], [MSG_AUSENTE])
        self.reserva()
        st = self.banco.status_atuais()
        self.assertEqual((st[corrigido]["status"], st[ignorado]["status"]), ("corrigido", "ignorado"))
        regravar_divergencias(xlsx, RES_TITULAR, [MSG_AUSENTE])
        self.reserva()
        st = self.banco.status_atuais()
        # corrigido que volta = correção que não pegou; ignorado continua ignorado
        self.assertEqual((st[corrigido]["status"], st[ignorado]["status"]), ("aberto", "ignorado"))

    def test_texto_com_numero_diferente_mantem_a_chave(self):
        antes = self.apont(self.reserva(), "118 dias")
        res = [RES_TITULAR[0], resultado("comprovante.pdf", "Comprovante de Endereço", "Conta de luz",
                                         divergencias=[MSG_PRAZO.format(125)])]
        regravar_divergencias(self.planilhas["titular"], res, [MSG_AUSENTE])
        depois = self.apont(self.reserva(), "125 dias")
        self.assertEqual(antes["chave"], depois["chave"])
        self.assertEqual(self.banco.buscar_apontamento(antes["chave"])["texto"], MSG_PRAZO.format(125))
        self.assertNotIn(antes["chave"], self.banco.status_atuais())

    def test_chave_estavel_e_ordinal(self):
        a = dados.montar_apontamentos("reserva", 1, "P", [("x.pdf", "aviso", "texto livre 1"),
                                                         ("x.pdf", "aviso", "texto livre 1")])
        self.assertNotEqual(a[0]["chave"], a[1]["chave"])
        b = dados.montar_apontamentos("reserva", 1, "P", [("x.pdf", "aviso", "Texto   LIVRE 1")])
        self.assertEqual(a[0]["chave"], b[0]["chave"])
        c = dados.montar_apontamentos("reserva", 1, "P", [("x.pdf", "aviso", "texto livre 2")])
        self.assertNotEqual(a[0]["chave"], c[0]["chave"])  # "Outros": o texto conta

    def test_primeiro_visto_e_o_mtime_da_planilha(self):
        chave = self.apont(self.reserva(), "similaridade 50%")["chave"]
        a = self.banco.buscar_apontamento(chave)
        self.assertEqual(banco_mod.de_texto(a["primeiro_visto"]),
                         datetime.fromtimestamp(ONTEM).replace(microsecond=0))

    def test_planilha_mais_velha_nao_desfaz(self):
        chave = self.apont(self.reserva(), "similaridade 50%")["chave"]
        self.assertFalse(self.banco.sincronizar_pasta("reserva", "FULANA-FICTICIA-reserva-1001", 1001,
                                                      ONTEM - 999, []))
        self.assertEqual(self.banco.buscar_apontamento(chave)["presente"], 1)


# ==========================================================================
# Métricas
# ==========================================================================
class Metricas(Base):
    AGORA = datetime(2026, 10, 13, 12, 0, 0)

    def test_calculo_sintetico(self):
        reservas = self.app.state.repo.atualizar()
        dia = lambda n: banco_mod.para_texto(self.AGORA - timedelta(days=n))  # noqa: E731
        aponts = [
            {"chave": "a", "nivel": "divergencia", "categoria": "Nome diferente do cadastro", "primeiro_visto": dia(10)},
            {"chave": "b", "nivel": "divergencia", "categoria": "Nome diferente do cadastro", "primeiro_visto": dia(10)},
            {"chave": "c", "nivel": "aviso", "categoria": "Outros", "primeiro_visto": dia(5)},
            {"chave": "d", "nivel": "aviso", "categoria": "Outros", "primeiro_visto": dia(2)},
            {"chave": "e", "nivel": "divergencia", "categoria": "Outros", "primeiro_visto": dia(200)},
        ]
        eventos = [
            {"id": 1, "chave": "a", "status": "corrigido", "motivo": None, "usuario": "ana", "em": dia(9)},
            {"id": 2, "chave": "b", "status": "resolvido_na_reanalise", "motivo": None, "usuario": "sistema", "em": dia(7)},
            {"id": 3, "chave": "c", "status": "ignorado", "motivo": "Falso positivo da IA", "usuario": "ana", "em": dia(4)},
            {"id": 4, "chave": "e", "status": "corrigido", "motivo": None, "usuario": "bia", "em": dia(1)},
        ]
        m = metricas.calcular(reservas, aponts, eventos, "30", 30, agora=self.AGORA)
        self.assertEqual((m["apontamentos"], m["divergencias"], m["avisos"]), (4, 2, 2))
        self.assertEqual(m["por_status"], {"aberto": 1, "corrigido": 1, "ignorado": 1,
                                           "resolvido_na_reanalise": 1})
        self.assertEqual(m["taxa_resolucao"], round(100 * 2 / 3, 1))  # (a, b) / (4 - 1 ignorado)
        # a: 24 h; b: 72 h; e: 199 dias (resolvido no período, visto antes dele)
        self.assertEqual(m["horas_ate_resolver"]["mediana"], 72.0)
        cats = {c["categoria"]: c for c in m["por_categoria"]}
        self.assertEqual(cats["Outros"]["pct_falso_positivo"], 50.0)
        self.assertEqual(cats["Nome diferente do cadastro"]["pct_falso_positivo"], 0.0)
        self.assertEqual(cats["Nome diferente do cadastro"]["corrigidos"], 1)
        self.assertEqual(m["ignorados_por_motivo"], [{"motivo": "Falso positivo da IA", "n": 1}])
        self.assertEqual(m["por_usuario"], [{"usuario": "ana", "tratativas": 2},
                                            {"usuario": "bia", "tratativas": 1}])
        self.assertLessEqual(len(m["por_semana"]), 12)
        self.assertEqual(sum(s["novos"] for s in m["por_semana"]), 4)
        self.assertEqual(sum(s["resolvidos"] for s in m["por_semana"]), 3)

        tudo = metricas.calcular(reservas, aponts, eventos, "tudo", None, agora=self.AGORA)
        self.assertEqual(tudo["apontamentos"], 5)
        self.assertIsNone(tudo["tempo_economizado_horas"])
        self.assertEqual(len(tudo["por_semana"]), 12)

        vazio = metricas.calcular([], [], [], "7", None, agora=self.AGORA)
        self.assertIsNone(vazio["horas_ate_resolver"])
        self.assertIsNone(vazio["taxa_resolucao"])
        self.assertIsNone(vazio["pct_reservas_com_divergencia"])

    def test_api_impacto_sobre_as_planilhas(self):
        self.entrar()
        r = self.client.get("/api/impacto", params={"periodo": "tudo"})
        self.assertEqual(r.status_code, 200)
        corpo = r.json()
        m = corpo["m"]
        self.assertEqual(corpo["periodo"], "tudo")
        self.assertEqual((m["reservas_analisadas"], m["pessoas"], m["documentos_lidos"]), (3, 4, 5))
        self.assertEqual((m["apontamentos"], m["divergencias"], m["avisos"]), (4, 3, 1))
        self.assertEqual(m["pct_reservas_com_divergencia"], round(100 / 3, 1))
        self.assertEqual(m["tempo_economizado_horas"], 1.5)  # 3 reservas x 30 min
        self.assertEqual(self.client.get("/api/impacto?periodo=abc").json()["periodo"], "30")


# ==========================================================================
# Exportação
# ==========================================================================
class Exportacao(Base):
    def test_xlsx_com_as_duas_abas(self):
        self.entrar()
        r = self.client.get("/exportar.xlsx", params={"status": "divergencia"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r.headers["content-disposition"])
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Reservas", "Apontamentos"])
        reservas = list(wb["Reservas"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(sorted(linha[1] for linha in reservas), [1001])
        self.assertEqual({linha[0] for linha in reservas}, {"Reserva"})
        apont = list(wb["Apontamentos"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(apont), 4)
        self.assertIn(MSG_AUSENTE, [linha[6] for linha in apont])
        self.assertTrue(all(linha[7] == "aberto" for linha in apont))

    def test_texto_com_igual_nao_vira_formula(self):
        detalhe = {"tipo": "reserva", "id": 1, "titular": "=HYPERLINK(\"x\")", "pessoas": [{
            "nome": "N", "papel": "titular", "apontamentos_pasta": [{
                "nivel": "aviso", "categoria": "Outros", "texto": "=1+1", "status": "aberto",
                "motivo": None, "responsavel": None, "atualizado_em": None}], "documentos": []}]}
        wb = load_workbook(io.BytesIO(gerar_xlsx([detalhe])))
        self.assertEqual(wb["Reservas"]["C2"].value, "=HYPERLINK(\"x\")")
        self.assertEqual(wb["Reservas"]["C2"].data_type, "s")
        self.assertEqual(wb["Apontamentos"]["G2"].data_type, "s")


# ==========================================================================
# Rota de arquivo
# ==========================================================================
class Arquivo(Base):
    PASTA = "FULANA-FICTICIA-reserva-1001"

    def setUp(self):
        super().setUp()
        self.entrar()

    def url(self, tipo_cv, id_=1001):
        for p in self.reserva("reserva", id_)["pessoas"]:
            for d in p["documentos"]:
                if d["tipo_cv"] == tipo_cv:
                    return d["url"]
        self.fail(tipo_cv)

    def token(self, arquivo, pasta=PASTA, tipo="reserva"):
        return dados.token_do_arquivo(self.banco.segredo, tipo, pasta, arquivo)

    def test_arquivo_valido_abre_inline_com_nome_neutro(self):
        r = self.client.get(self.url("RG, CPF, CNH ou OUTROS"))
        self.assertEqual(r.status_code, 200)
        self.assertTrue(r.content.startswith(b"%PDF"))
        self.assertEqual(r.headers["content-type"], "application/pdf")
        self.assertEqual(r.headers["content-disposition"], 'inline; filename="documento.pdf"')
        r = self.client.get(self.url("CPF"))  # o arquivo com CPF no nome
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers["content-disposition"], 'inline; filename="documento.pdf"')

    def test_tipo_desconhecido_vai_como_download(self):
        criar_pasta(self.saidas, "reserva", "OUTRA-reserva-1005", {"nome": "OUTRA"},
                    [{"tipo": "Documentos Diversos", "nome": "p", "arquivo_local": "pagina.html"}])
        (self.saidas / "reserva" / "OUTRA-reserva-1005" / "pagina.html").write_text("<script>alert(1)</script>")
        r = self.client.get(self.url("Documentos Diversos", 1005))
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers["content-type"], "application/octet-stream")
        self.assertEqual(r.headers["content-disposition"], 'attachment; filename="documento.html"')

    def test_token_e_amarrado_a_reserva_e_ao_segredo(self):
        url = self.url("RG, CPF, CNH ou OUTROS")
        token = url.rsplit("/", 1)[1]
        self.assertEqual(token, self.token("rg.pdf"))
        ruins = [
            f"/arquivo/reserva/1004/{token}",              # token certo, outra reserva
            f"/arquivo/pre_cadastro/1001/{token}",         # outro tipo
            f"/arquivo/reserva/1001/{token.upper()}",
            f"/arquivo/reserva/1001/{token[:-1]}",
            f"/arquivo/reserva/1001/{token}0",
            "/arquivo/reserva/1001/" + "0" * 24,
            "/arquivo/reserva/1001/" + dados.token_do_arquivo(b"outro-segredo", "reserva", self.PASTA, "rg.pdf"),
            f"/arquivo/outro/1001/{token}",
            f"/arquivo/RESERVA/1001/{token}",
            f"/arquivo/reserva/abc/{token}",
            "/arquivo/reserva/1001/%2e%2e",
            f"/arquivo/reserva/{self.PASTA}/rg.pdf",       # formato antigo
            f"/arquivo/reserva/1001/{token}/..%2F..%2Fsegredo.txt",
        ]
        for u in ruins:
            r = self.client.get(u)
            self.assertNotEqual(r.status_code, 200, u)
            self.assertNotIn(b"%PDF", r.content, u)

    def test_arquivo_fora_da_planilha_nao_sai_nem_com_token_certo(self):
        (self.saidas / "reserva" / self.PASTA / "extra.pdf").write_bytes(b"%PDF extra")
        r = self.client.get(f"/arquivo/reserva/1001/{self.token('extra.pdf')}")
        self.assertEqual(r.status_code, 404)

    def test_planilha_com_nome_malicioso_continua_barrada(self):
        """Mesmo que a aba Documentos traga um caminho, as defesas de caminho valem."""
        pasta = "MALICIA-reserva-1006"
        criar_pasta(self.saidas, "reserva", pasta, {"nome": "MALICIA"},
                    [{"tipo": "A", "nome": "a", "arquivo_local": "../segredo.txt"},
                     {"tipo": "B", "nome": "b", "arquivo_local": f"{pasta}.xlsx"},
                     {"tipo": "C", "nome": "c", "arquivo_local": "..\\..\\fora.txt"}])
        d = self.reserva("reserva", 1006)
        self.assertTrue(all(doc["url"] is None for doc in d["pessoas"][0]["documentos"]))
        for arquivo in ("../segredo.txt", f"{pasta}.xlsx", "..\\..\\fora.txt"):
            r = self.client.get(f"/arquivo/reserva/1006/{self.token(arquivo, pasta)}")
            self.assertEqual(r.status_code, 404, arquivo)
            self.assertNotIn(b"SEGREDO", r.content)

    def test_arquivo_apagado_depois_do_link(self):
        url = self.url("RG, CPF, CNH ou OUTROS")
        (self.saidas / "reserva" / self.PASTA / "rg.pdf").unlink()
        self.assertEqual(self.client.get(url).status_code, 404)

    def test_bloqueios_na_funcao(self):
        s = self.saidas
        self.assertIsNotNone(dados.caminho_seguro_do_arquivo(s, "reserva", self.PASTA, "rg.pdf"))
        casos = [
            ("reserva", "..", "segredo.txt"), ("reserva", self.PASTA, ".."),
            ("reserva", self.PASTA, "../segredo.txt"), ("reserva", self.PASTA, "..\\segredo.txt"),
            ("reserva", "..\\..", "fora.txt"), ("reserva", self.PASTA, str(self.tmp / "fora.txt")),
            ("reserva", str(s / "reserva" / self.PASTA), "rg.pdf"), ("reserva", "/" + self.PASTA, "rg.pdf"),
            ("reserva", self.PASTA, f"{self.PASTA}.xlsx"), ("reserva", self.PASTA, "rg.pdf "),
            ("pre", self.PASTA, "rg.pdf"), ("../reserva", self.PASTA, "rg.pdf"),
            ("reserva", ".", "segredo.txt"), ("reserva", self.PASTA, "rg\x00.pdf"),
        ]
        for tipo, pasta, arquivo in casos:
            self.assertIsNone(dados.caminho_seguro_do_arquivo(s, tipo, pasta, arquivo), (tipo, pasta, arquivo))


# ==========================================================================
# "Por quê" da tratativa (detalhe)
# ==========================================================================
TELEFONE = "(11) 98765-4321"  # fictício
ASSOCIADO_BELTRANO = {"nome": "BELTRANO SOUZA", "documento": CPF_SEGUNDO,
                      "tipo_associacao": "2º Promissário Comprador"}


class Detalhe(Base):
    def setUp(self):
        super().setUp()
        self.csrf = self.entrar()
        self.chave = self.apont(self.reserva(), "118 dias")["chave"]

    def tratar(self, **corpo):
        return self.client.post("/api/tratativa", json={"chave": self.chave, **corpo},
                                headers={"X-CSRF-Token": self.csrf})

    def test_migracao_de_banco_sem_a_coluna(self):
        antigo = self.tmp / "antigo" / "painel.db"
        antigo.parent.mkdir()
        esquema_antigo = banco_mod.ESQUEMA.replace(",\n    detalhe TEXT,\n    veredito TEXT", "")
        self.assertNotEqual(esquema_antigo, banco_mod.ESQUEMA)
        import sqlite3
        con = sqlite3.connect(antigo)
        con.executescript(esquema_antigo)
        con.execute("INSERT INTO apontamentos VALUES ('k1', 'reserva', 1, 'P', 'a.pdf', 'aviso', 'Outros', "
                    "'t', '2026-10-01 10:00:00', '2026-10-01 10:00:00', 1)")
        con.execute("INSERT INTO eventos (chave, status, motivo, usuario, em) VALUES "
                    "('k1', 'ignorado', 'Falso positivo da IA', 'ana', '2026-10-02 10:00:00')")
        con.commit()
        con.close()

        for _ in range(2):  # idempotente: abrir de novo não tenta recriar a coluna
            b = banco_mod.Banco(antigo)
        with b.conexao() as con:
            colunas = {linha["name"] for linha in con.execute("PRAGMA table_info(eventos)")}
        self.assertIn("detalhe", colunas)
        self.assertIn("veredito", colunas)
        ev = b.status_atuais()["k1"]
        self.assertEqual((ev["status"], ev["motivo"], ev["detalhe"]), ("ignorado", "Falso positivo da IA", None))
        self.assertIsNone(ev["veredito"], "evento antigo fica sem veredito")
        self.assertEqual(len(b.listar_eventos()), 1)
        b.registrar_tratativa("k1", "ignorado", "Falso positivo da IA", "ana", "  lido certo  ")
        self.assertEqual(b.status_atuais()["k1"]["detalhe"], "lido certo")

    def test_veredito_combinacoes_validas(self):
        validas = [
            dict(status="corrigido", veredito="real"),
            dict(status="corrigido", veredito="real", detalhe="pedi outro comprovante"),
            dict(status="ignorado", motivo="Corrigido fora do CV", veredito="real"),
            dict(status="ignorado", motivo="Falso positivo da IA", veredito="falso", detalhe="leu a data errada"),
            dict(status="ignorado", motivo="Regra não se aplica a este caso", veredito="falso", detalhe="certidão na pasta do cônjuge"),
            dict(status="ignorado", motivo="Outro: documento conferido no papel", veredito="falso",
                 detalhe="documento conferido no papel"),
            dict(status="aberto", veredito="nao_sei", detalhe="foto cortada"),
        ]
        for corpo in validas:
            r = self.tratar(**corpo)
            self.assertEqual(r.status_code, 200, (corpo, r.text))
            self.assertEqual(r.json()["veredito"], corpo["veredito"])
            self.assertEqual(self.banco.status_atuais()[self.chave]["veredito"], corpo["veredito"])
            self.assertEqual(self.apont(self.reserva(), "118 dias")["veredito"], corpo["veredito"])
        # os motivos gravados continuam os de sempre (métricas e revisão leem estes)
        ev = self.banco.status_atuais()[self.chave]
        self.assertEqual((ev["status"], ev["veredito"]), ("aberto", "nao_sei"))
        # Reabrir (sem veredito) limpa o veredito
        r = self.tratar(status="aberto")
        self.assertEqual(r.status_code, 200)
        self.assertIsNone(r.json()["veredito"])

    def test_veredito_combinacoes_invalidas(self):
        invalidas = [
            dict(status="aberto", veredito="real"),
            dict(status="ignorado", motivo="Falso positivo da IA", veredito="real"),
            dict(status="corrigido", veredito="falso", detalhe="x"),
            dict(status="ignorado", motivo="Falso positivo da IA", veredito="falso"),  # falta detalhe
            dict(status="ignorado", motivo="Falso positivo da IA", veredito="falso", detalhe="  "),
            dict(status="ignorado", motivo="Corrigido fora do CV", veredito="falso", detalhe="x"),
            dict(status="aberto", veredito="nao_sei"),  # falta detalhe
            dict(status="corrigido", veredito="nao_sei", detalhe="x"),
            dict(status="ignorado", motivo="Falso positivo da IA", veredito="nao_sei", detalhe="x"),
            dict(status="corrigido", veredito="talvez"),
            dict(status="corrigido", veredito=5),
        ]
        for corpo in invalidas:
            self.assertEqual(self.tratar(**corpo).status_code, 400, corpo)
        self.assertNotIn(self.chave, self.banco.status_atuais(), "nada foi gravado")
        for corpo in invalidas[:3] + invalidas[6:7]:
            kw = {k: v for k, v in corpo.items() if k != "status"}
            with self.assertRaises(ValueError):
                self.banco.registrar_tratativa(self.chave, corpo["status"], kw.get("motivo"), "ana",
                                               kw.get("detalhe"), kw["veredito"])

    def test_sem_veredito_comportamento_de_hoje(self):
        r = self.tratar(status="ignorado", motivo="Falso positivo da IA")
        self.assertEqual(r.status_code, 200)
        self.assertIsNone(r.json()["veredito"])

    def test_detalhe_salvo_e_devolvido(self):
        r = self.tratar(status="ignorado", motivo="Regra não se aplica a este caso",
                        detalhe="  comprovante é de conta do cônjuge  ")
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["detalhe"], "comprovante é de conta do cônjuge")
        self.assertEqual(self.apont(self.reserva(), "118 dias")["detalhe"], "comprovante é de conta do cônjuge")
        # vazio = None; qualquer status aceita
        self.assertIsNone(self.tratar(status="ignorado", motivo="Falso positivo da IA", detalhe="   ").json()["detalhe"])
        self.assertEqual(self.tratar(status="corrigido", detalhe="pedi outro").json()["detalhe"], "pedi outro")
        self.assertIsNone(self.tratar(status="aberto").json()["detalhe"])

    def test_detalhe_limitado_a_500(self):
        ok = self.tratar(status="ignorado", motivo="Falso positivo da IA", detalhe="x" * 500)
        self.assertEqual(ok.status_code, 200)
        self.assertEqual(self.tratar(status="ignorado", motivo="Falso positivo da IA",
                                     detalhe="x" * 501).status_code, 400)
        self.assertEqual(self.tratar(status="ignorado", motivo="Falso positivo da IA",
                                     detalhe=123).status_code, 400)
        self.assertEqual(self.banco.status_atuais()[self.chave]["detalhe"], "x" * 500)  # o 501 não gravou
        with self.assertRaises(ValueError):
            self.banco.registrar_tratativa(self.chave, "corrigido", None, "ana", "y" * 501)

    def test_cpf_e_telefone_mascarados_ao_exibir_e_exportar(self):
        detalhe = f"cliente ligou do {TELEFONE}, CPF {CPF_TITULAR}"
        r = self.tratar(status="ignorado", motivo="Falso positivo da IA", detalhe=detalhe)
        esperado = f"cliente ligou do {dados.TELEFONE_MASCARADO}, CPF ***.456.789-**"
        self.assertEqual(r.json()["detalhe"], esperado)
        self.assertEqual(self.apont(self.reserva(), "118 dias")["detalhe"], esperado)
        # o banco não guarda o cru
        self.assertEqual(self.banco.listar_eventos()[-1]["detalhe"], esperado)
        self.assertEqual(self.banco.status_atuais()[self.chave]["detalhe"], esperado)
        planilha = self.client.get("/exportar.xlsx")
        wb = load_workbook(io.BytesIO(planilha.content))
        ws = wb["Apontamentos"]
        cabecalho = [c.value for c in ws[1]]
        self.assertEqual(cabecalho[8:10], ["motivo", "detalhe"])
        detalhes = [linha[9] for linha in ws.iter_rows(min_row=2, values_only=True) if linha[9]]
        self.assertEqual(detalhes, [esperado])
        for texto in (r.text, self.client.get("/api/reserva/reserva/1001").text):
            self.assertNotIn(CPF_TITULAR, texto)
            self.assertNotIn("98765-4321", texto)

    def test_mascarar_texto_livre(self):
        m = dados.mascarar_texto_livre
        tel = dados.TELEFONE_MASCARADO
        for telefone in ("(11) 98765-4321", "11 98765 4321", "98765-4321", "8765-4321", "(11)3456-7890"):
            self.assertEqual(m(f"tel {telefone} ok"), f"tel {tel} ok", telefone)
        self.assertEqual(m("12345678909"), "***.456.789-**")
        for intacto in ("emitido em 01/01/2026", "comprovante com 118 dias", "reserva 33937", "R$ 4.500,00"):
            self.assertEqual(m(intacto), intacto)
        self.assertIsNone(m(None))
        self.assertEqual(m(""), "")

    def test_mascarar_para_o_time(self):
        m = dados.mascarar_para_o_time
        self.assertEqual(
            m(MSG_NOME, "Nome diferente do cadastro"),
            "Nome no documento (nome: 2 palavras, 11 letras) difere do cadastro "
            "(nome: 2 palavras, 14 letras) - similaridade 50%")
        self.assertEqual(m("Nome no documento ('JOAO D'AVILA') difere do cadastro ('JOAO') - similaridade 60%",
                           "Nome diferente do cadastro"),
                         "Nome no documento (nome: 2 palavras, 10 letras) difere do cadastro "
                         "(nome: 1 palavra, 4 letras) - similaridade 60%")
        endereco = ("Endereço no comprovante ('RUA DAS FLORES 10 (FUNDOS) CENTRO') não bate com o endereço "
                    "cadastrado ('AV BRASIL 200') - similaridade 30% - ATENÇÃO: IA reportou baixa confiança")
        self.assertEqual(m(endereco, "Endereço não bate com o cadastro"),
                         "Endereço no comprovante (endereço: 6 palavras) não bate com o endereço cadastrado "
                         "(endereço: 3 palavras) - similaridade 30% - ATENÇÃO: IA reportou baixa confiança")
        # CPF: máscara de sempre; número de tamanho estranho também some
        self.assertEqual(m(MSG_CPF, "CPF diferente do cadastro"),
                         "CPF no documento (***.654.321-**) difere do cadastro (***.222.333-**)")
        self.assertEqual(m("CPF no documento (1234567890) difere do cadastro (111.222.333-96)",
                           "CPF diferente do cadastro"),
                         "CPF no documento (***) difere do cadastro (***.222.333-**)")
        self.assertEqual(m("CNPJ no documento (12.345.678/0001-95) difere do cadastro (12345678000195)",
                           "CNPJ diferente do cadastro"),
                         "CNPJ no documento (**.345.678/0001-**) difere do cadastro (**.345.678/0001-**)")
        # outras categorias: só a máscara de CPF/telefone
        self.assertEqual(m(MSG_PRAZO.format(118), "Comprovante de endereço vencido"),
                         MSG_PRAZO.format(118))
        self.assertEqual(m(f"liguei {TELEFONE} (cliente)", "Outros"), f"liguei {dados.TELEFONE_MASCARADO} (cliente)")
        # "por quê" numa categoria de nome: o que está entre parênteses também vira marcador
        self.assertEqual(m("no RG está (FULANA O. FICTICIA), abreviado", "Nome diferente do cadastro"),
                         "no RG está (nome: 3 palavras, 15 letras), abreviado")
        self.assertIsNone(m(None, "Nome diferente do cadastro"))

    # As mensagens de "Outros" com nome de pessoa saem da PRÓPRIA auditoria:
    # se o texto mudar lá, o teste quebra em vez de o nome vazar no Excel.
    def _mensagem(self, inicio, doc_ia, cadastro, tipo_cv="RG", associados=None):
        r = ad.comparar_documento_com_cadastro(doc_ia, cadastro, tipo_cv, associados)
        achadas = [t for t in r["divergencias"] + r["avisos"] if t.startswith(inicio)]
        self.assertEqual(len(achadas), 1, r)
        self.assertEqual(categorizar(achadas[0]), "Outros")  # por isso vai por padrão de texto
        return achadas[0]

    def test_mensagem_de_pj(self):
        msg = self._mensagem("Cliente é pessoa jurídica", {"nome_encontrado": "FULANO DE TAL"},
                             {"nome": "EMPRESA FICTICIA LTDA", "documento": "12345678000195", "cliente_tipo": "J"})
        self.assertEqual(
            dados.mascarar_para_o_time(msg, "Outros"),
            "Cliente é pessoa jurídica (nome: 3 palavras, 19 letras), e este documento é de "
            "(nome: 3 palavras, 11 letras) - representante da empresa, conforme a cartilha. Nome e CPF não "
            "foram comparados com o cadastro (comparar pessoa com empresa não faz sentido)")

    def test_mensagem_de_co_participante_pelo_nome(self):
        msg = self._mensagem("Documento é de", {"nome_encontrado": "BELTRANO SOUZA"},
                             {"nome": "FULANA FICTICIA", "documento": CPF_TITULAR},
                             associados=[ASSOCIADO_BELTRANO])
        self.assertEqual(
            dados.mascarar_para_o_time(msg, "Outros"),
            "Documento é de (nome: 2 palavras, 13 letras) (2º Promissário Comprador), não de "
            "(nome: 2 palavras, 14 letras) - é co-participante legítimo da operação, não divergência")

    def test_mensagem_de_co_participante_redacao_antiga(self):
        # Planilhas já geradas ainda têm "não do titular" no lugar do nome do
        # cadastro; a medição nos apontamentos reais (13/10/2026) achou 5.
        msg = ("Documento é de BELTRANO DE SOUZA LIMA (2º Promissário Comprador), não do titular "
               "- é co-participante legítimo da operação, não divergência")
        self.assertEqual(
            dados.mascarar_para_o_time(msg, "Outros"),
            "Documento é de (nome: 4 palavras, 19 letras) (2º Promissário Comprador), não do titular "
            "- é co-participante legítimo da operação, não divergência")

    def test_mensagem_de_co_participante_pelo_cpf(self):
        msg = self._mensagem("CPF do documento pertence a", {"cpf_encontrado": CPF_SEGUNDO},
                             {"nome": "FULANA FICTICIA", "documento": CPF_TITULAR},
                             associados=[ASSOCIADO_BELTRANO])
        self.assertEqual(
            dados.mascarar_para_o_time(msg, "Outros"),
            "CPF do documento pertence a (nome: 2 palavras, 13 letras) (2º Promissário Comprador), "
            "co-participante da operação - não é divergência de cadastro")

    # test_mensagem_de_renda_com_mais_de_uma_pessoa saiu em 19/10/2026: testava a
    # mensagem da "renda do núcleo" (nomes de todos da operação + "o valor do
    # cadastro é a soma"), que deixou de existir com a renda por proponente (gestora comercial,
    # 14/10/2026, conferir_renda_da_pessoa()). A mensagem nova cai em "Renda
    # diferente da cadastrada", não em "Outros".

    def test_marcador_nao_e_mascarado_de_novo(self):
        m = dados.mascarar_para_o_time
        uma_vez = m(MSG_NOME, "Nome diferente do cadastro")
        self.assertEqual(m(uma_vez, "Nome diferente do cadastro"), uma_vez)


# ==========================================================================
# Revisão de regras
# ==========================================================================
FP = "Falso positivo da IA"
REGRA = "Regra não se aplica a este caso"


class Revisao(Base):
    def setUp(self):
        super().setUp()
        self.csrf = self.entrar()
        d = self.reserva()
        self.nome = self.apont(d, "similaridade 50%")["chave"]       # titular, rg.pdf
        self.prazo = self.apont(d, "118 dias")["chave"]              # titular, comprovante.pdf
        self.ausente = self.apont(d, "certidão de estado civil")["chave"]  # sentinela
        self.segundo = self.apont(d, "similaridade 80%")["chave"]    # 2º comprador, rg2.png

    def tratar(self, chave, status="ignorado", motivo=FP, detalhe=None):
        r = self.client.post("/api/tratativa", json={"chave": chave, "status": status, "motivo": motivo,
                                                     "detalhe": detalhe},
                             headers={"X-CSRF-Token": self.csrf})
        self.assertEqual(r.status_code, 200, r.text)

    def datar(self, chave, momento: datetime):
        with self.banco.conexao() as con:
            con.execute("UPDATE eventos SET em = ? WHERE id = (SELECT MAX(id) FROM eventos WHERE chave = ?)",
                        (banco_mod.para_texto(momento), chave))

    def revisao(self, **params):
        r = self.client.get("/api/revisao", params=params)
        self.assertEqual(r.status_code, 200, r.text)
        return r.json()

    def chaves(self, **params):
        return sorted(c["chave"] for g in self.revisao(**params)["grupos"] for c in g["casos"])

    def test_so_ignorados_atuais_e_filtro_de_motivo(self):
        self.tratar(self.nome, motivo=FP, detalhe="nome com abreviação")
        self.tratar(self.prazo, motivo=REGRA)
        self.tratar(self.ausente, motivo="Corrigido fora do CV")
        self.tratar(self.segundo, status="corrigido", motivo=None)
        r = self.revisao()
        self.assertEqual(r["filtro"], {"motivo": "fp_regra", "periodo": "tudo"})
        self.assertEqual(r["totais"], {"casos": 2, "categorias": 2, "sem_detalhe": 1})
        self.assertEqual(self.chaves(), sorted([self.nome, self.prazo]))
        self.assertEqual(self.chaves(motivo="fp"), [self.nome])
        self.assertEqual(self.chaves(motivo="regra"), [self.prazo])
        self.assertEqual(self.chaves(motivo="todos"), sorted([self.nome, self.prazo, self.ausente]))
        self.assertEqual(self.revisao(motivo="xyz", periodo="5")["filtro"], {"motivo": "fp_regra", "periodo": "tudo"})

    def test_outro_entra_so_em_todos_agrupado(self):
        self.tratar(self.nome, motivo="Outro: conferido no original")
        self.assertEqual(self.chaves(), [])
        g = self.revisao(motivo="todos")["grupos"][0]
        self.assertEqual(g["por_motivo"], [{"motivo": "Outro", "n": 1}])
        self.assertEqual(g["casos"][0]["motivo"], "Outro: conferido no original")

    def test_reaberto_sai_da_revisao(self):
        self.tratar(self.nome)
        self.assertEqual(self.chaves(), [self.nome])
        self.tratar(self.nome, status="aberto", motivo=None)
        self.assertEqual(self.chaves(motivo="todos"), [])
        self.tratar(self.nome, status="ignorado", motivo=REGRA)  # ignorado de novo: volta
        self.assertEqual(self.chaves(), [self.nome])

    def test_periodo_pela_data_do_evento_de_ignorar(self):
        self.tratar(self.nome)
        self.tratar(self.prazo, motivo=REGRA)
        self.datar(self.prazo, datetime.now() - timedelta(days=40))
        self.assertEqual(self.chaves(periodo="30"), [self.nome])
        self.assertEqual(self.chaves(periodo="7"), [self.nome])
        self.assertEqual(self.chaves(periodo="90"), sorted([self.nome, self.prazo]))
        self.assertEqual(self.chaves(), sorted([self.nome, self.prazo]))

    def test_agrupamento_ordenacao_e_pct(self):
        self.tratar(self.nome, motivo=FP)
        self.tratar(self.segundo, motivo=REGRA)
        self.tratar(self.prazo, motivo=FP)
        agora = datetime.now()
        self.datar(self.nome, agora - timedelta(days=3))
        self.datar(self.segundo, agora - timedelta(days=1))
        grupos = self.revisao()["grupos"]
        self.assertEqual([(g["categoria"], g["total"]) for g in grupos],
                         [("Nome diferente do cadastro", 2), ("Comprovante de endereço vencido", 1)])
        nome = grupos[0]
        self.assertEqual([c["chave"] for c in nome["casos"]], [self.segundo, self.nome])  # mais recente antes
        self.assertEqual(nome["por_motivo"], [{"motivo": FP, "n": 1}, {"motivo": REGRA, "n": 1}])
        # o mesmo número da página de impacto, no mesmo período
        impacto = {c["categoria"]: c["pct_falso_positivo"]
                   for c in self.client.get("/api/impacto", params={"periodo": "tudo"}).json()["m"]["por_categoria"]}
        self.assertEqual(nome["pct_falso_positivo"], 50.0)
        self.assertEqual(nome["pct_falso_positivo"], impacto["Nome diferente do cadastro"])
        self.assertEqual(grupos[1]["pct_falso_positivo"], 100.0)

    def test_campos_e_links(self):
        self.tratar(self.nome, detalhe="ok")
        self.tratar(self.segundo)
        self.tratar(self.ausente, motivo="Corrigido fora do CV")
        casos = {c["chave"]: c for g in self.revisao(motivo="todos")["grupos"] for c in g["casos"]}
        c = casos[self.nome]
        self.assertEqual(
            {k: c[k] for k in ("tipo", "id", "papel", "tipo_cv", "tipo_ia", "nivel", "motivo", "detalhe",
                               "responsavel", "url_reserva", "presente")},
            {"tipo": "reserva", "id": 1001, "papel": "titular", "tipo_cv": "RG, CPF, CNH ou OUTROS",
             "tipo_ia": "RG", "nivel": "divergencia", "motivo": FP, "detalhe": "ok", "responsavel": "ana",
             "url_reserva": "/reserva/reserva/1001", "presente": True})
        self.assertEqual(c["texto"], MSG_NOME)
        self.assertRegex(c["atualizado_em"], r"^\d{2}/\d{2}/\d{4} \d{2}:\d{2}$")
        self.assertRegex(c["url_arquivo"], RE_URL_ARQUIVO)
        baixado = self.client.get(c["url_arquivo"])
        self.assertEqual(baixado.status_code, 200)
        self.assertIn(b"rg.pdf", baixado.content)
        self.assertEqual(self.client.get(c["url_reserva"]).status_code, 200)

        s = casos[self.segundo]
        self.assertEqual((s["papel"], s["tipo_ia"], s["nivel"]), ("2º Promissário Comprador", "CNH", "aviso"))
        self.assertIn(b"rg2.png", self.client.get(s["url_arquivo"]).content)

        a = casos[self.ausente]
        self.assertEqual((a["tipo_cv"], a["tipo_ia"], a["url_arquivo"]), (None, None, None))

    def test_ainda_aparece_depois_da_reanalise(self):
        self.tratar(self.prazo, motivo=REGRA)
        sem_prazo = [RES_TITULAR[0], resultado("comprovante.pdf", "Comprovante de Endereço", "Conta de luz")]
        regravar_divergencias(self.planilhas["titular"], sem_prazo, [MSG_AUSENTE])
        c = self.revisao()["grupos"][0]["casos"][0]
        self.assertEqual(c["chave"], self.prazo)
        self.assertFalse(c["presente"])
        # sumiu da planilha, mas o arquivo continua lá: tipo e link ainda resolvem
        self.assertEqual((c["tipo_cv"], c["tipo_ia"]), ("Comprovante de Endereço", "Conta de luz"))
        self.assertIn(b"comprovante.pdf", self.client.get(c["url_arquivo"]).content)

    def test_sem_nome_de_cliente_nem_cpf_cru(self):
        self.tratar(self.prazo, motivo=REGRA, detalhe=f"CPF {CPF_TITULAR}, tel {TELEFONE}")
        self.tratar(self.ausente, motivo=FP)
        respostas = [self.client.get("/api/revisao", params={"motivo": "todos"}).text,
                     self.client.get("/revisao", params={"motivo": "todos"}).text]
        wb = load_workbook(io.BytesIO(self.client.get("/revisao.xlsx", params={"motivo": "todos"}).content))
        respostas.append(" ".join(str(v) for linha in wb.active.iter_rows(values_only=True) for v in linha))
        for texto in respostas:
            for proibido in ("FULANA", "BELTRANO", "FICTICIA", CPF_TITULAR, "12345678909", "98765-4321",
                             "rg.pdf", "comprovante.pdf"):
                self.assertNotIn(proibido, texto)
        self.assertIn("***.456.789-**", respostas[0])
        caso = self.revisao(motivo="todos")["grupos"][0]["casos"][0]
        self.assertNotIn("nome", caso)
        self.assertNotIn("pasta", caso)

    def test_xlsx_mascara_nome_no_texto_mas_a_pagina_nao(self):
        self.tratar(self.nome, detalhe="no RG está (FULANA O. FICTICIA)")
        self.tratar(self.segundo)
        wb = load_workbook(io.BytesIO(self.client.get("/revisao.xlsx").content))
        linhas = list(wb.active.iter_rows(min_row=2, values_only=True))
        textos = {linha[6] for linha in linhas}
        self.assertEqual(textos, {
            "Nome no documento (nome: 2 palavras, 11 letras) difere do cadastro (nome: 2 palavras, 14 letras) "
            "- similaridade 50%",
            "Nome no documento (nome: 2 palavras, 9 letras) difere do cadastro (nome: 2 palavras, 16 letras) "
            "- similaridade 80%"})
        self.assertIn("no RG está (nome: 3 palavras, 15 letras)", {linha[8] for linha in linhas})
        tudo = " ".join(str(v) for linha in linhas for v in linha)
        for nome in ("FULANA", "BELTRANO", "OUTRA"):
            self.assertNotIn(nome, tudo)
        # quem tem acesso ao site vê o texto inteiro
        casos = [c for g in self.revisao()["grupos"] for c in g["casos"]]
        self.assertIn(MSG_NOME, [c["texto"] for c in casos])
        self.assertIn("no RG está (FULANA O. FICTICIA)", [c["detalhe"] for c in casos])

    def test_xlsx_mascara_nome_de_co_participante(self):
        msg = ("Documento é de BELTRANO SOUZA (2º Promissário Comprador), não de FULANA FICTICIA - "
               "é co-participante legítimo da operação, não divergência")
        res = RES_TITULAR + [resultado("rg3.pdf", "RG, CPF, CNH ou OUTROS", "RG", avisos=[msg])]
        regravar_divergencias(self.planilhas["titular"], res, [MSG_AUSENTE])
        chave = self.apont(self.reserva(), "co-participante legítimo")["chave"]
        self.tratar(chave, motivo=REGRA)
        wb = load_workbook(io.BytesIO(self.client.get("/revisao.xlsx").content))
        (linha,) = list(wb.active.iter_rows(min_row=2, values_only=True))
        self.assertEqual((linha[0], linha[6]), (
            "Outros", "Documento é de (nome: 2 palavras, 13 letras) (2º Promissário Comprador), não de "
                      "(nome: 2 palavras, 14 letras) - é co-participante legítimo da operação, não divergência"))
        caso = self.revisao()["grupos"][0]["casos"][0]
        self.assertEqual(caso["texto"], msg)  # a página segue com o texto inteiro

    def test_pagina_xlsx_e_login(self):
        self.tratar(self.nome, detalhe="=HYPERLINK(\"x\")")
        self.tratar(self.prazo, motivo=REGRA)
        pagina = self.client.get("/revisao", params={"motivo": "fp"})
        self.assertEqual(pagina.status_code, 200)
        self.assertIn("REVISAO motivo=fp periodo=tudo casos=1 [Nome diferente do cadastro:1]", pagina.text)
        self.assertIn("Content-Security-Policy", pagina.headers)

        r = self.client.get("/revisao.xlsx")
        self.assertEqual(r.status_code, 200)
        self.assertIn("attachment", r.headers["content-disposition"])
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ["Revisao de regras"])
        ws = wb["Revisao de regras"]
        self.assertEqual([c.value for c in ws[1]],
                         ["categoria", "reserva", "papel", "documento", "IA leu", "nível", "apontamento",
                          "motivo", "detalhe", "quem", "quando", "ainda aparece"])
        linhas = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(linhas), 2)
        nome = next(linha for linha in linhas if linha[0] == "Nome diferente do cadastro")
        self.assertEqual(nome[1:6], ("Reserva 1001", "titular", "RG, CPF, CNH ou OUTROS", "RG", "divergencia"))
        self.assertEqual((nome[7], nome[9], nome[11]), (FP, "ana", "sim"))
        celula = next(c for c in ws["I"][1:] if c.value and c.value.startswith("="))
        self.assertEqual(celula.data_type, "s")  # detalhe com "=" não vira fórmula

        anonimo = TestClient(self.app, follow_redirects=False)
        self.assertEqual(anonimo.get("/api/revisao").status_code, 401)
        self.assertEqual(anonimo.get("/revisao").status_code, 303)
        self.assertEqual(anonimo.get("/revisao.xlsx").status_code, 303)
        anonimo.close()


# ==========================================================================
# Pré-cadastro fora do painel (decisão do João, 20/10/2026)
# ==========================================================================
class SemPreCadastro(Base):
    def test_lista_detalhe_e_arquivo_dao_404(self):
        self.entrar()
        self.assertNotIn(2002, [i["id"] for i in self.itens()["itens"]])
        for caminho in ("/reserva/pre_cadastro/2002", "/api/reserva/pre_cadastro/2002"):
            self.assertEqual(self.client.get(caminho).status_code, 404, caminho)
        token = dados.token_do_arquivo(self.banco.segredo, "pre_cadastro", "ANTIGA-FICTICIA-2002", "c.pdf")
        self.assertEqual(self.client.get(f"/arquivo/pre_cadastro/2002/{token}").status_code, 404)
        self.assertTrue((self.saidas / "pre_cadastro" / "ANTIGA-FICTICIA-2002" / "c.pdf").exists(), "nada apagado")

    def test_apontamentos_de_pre_cadastro_ficam_no_banco_fora_das_metricas(self):
        csrf = self.entrar()
        self.banco.sincronizar_pasta("pre_cadastro", "ANTIGA-FICTICIA-2002", 2002, ONTEM - 300, [
            {"chave": "pc-antigo", "arquivo": "c.pdf", "nivel": "aviso",
             "categoria": "Nome diferente do cadastro", "texto": MSG_AVISO_2O}])
        self.banco.registrar_tratativa("pc-antigo", "ignorado", FP, "ana")
        chave = self.apont(self.reserva(), "118 dias")["chave"]
        self.client.post("/api/tratativa", json={"chave": chave, "status": "ignorado", "motivo": FP},
                         headers={"X-CSRF-Token": csrf})
        self.assertIn("pc-antigo", [a["chave"] for a in self.banco.listar_apontamentos()], "fica no banco")
        self.assertEqual([e["chave"] for e in self.banco.listar_eventos("reserva")].count("pc-antigo"), 0)
        m = self.client.get("/api/impacto", params={"periodo": "tudo"}).json()["m"]
        self.assertEqual((m["apontamentos"], m["avisos"]), (4, 1), "só os das reservas")
        casos = [c for g in self.client.get("/api/revisao").json()["grupos"] for c in g["casos"]]
        self.assertEqual({(c["tipo"], c["id"]) for c in casos}, {("reserva", 1001)})
        for caminho in ("/exportar.xlsx", "/revisao.xlsx"):
            wb = load_workbook(io.BytesIO(self.client.get(caminho).content))
            textos = " ".join(str(v) for ws in wb for linha in ws.iter_rows(values_only=True) for v in linha if v)
            self.assertNotIn("2002", textos, caminho)
            self.assertNotIn("Pré-cadastro", textos, caminho)


# ==========================================================================
# Link "Abrir no CV"
# ==========================================================================
CV_RESERVA = "https://suaempresa.cvcrm.com.br/gestor/comercial/reservas/{id}/administrar"
CV_PRE = "https://suaempresa.cvcrm.com.br/gestor/clientes/precadastros/{id}"


class LinkDoCV(Base):
    def test_modelo_valido(self):
        for modelo in (CV_RESERVA, CV_PRE, "https://SUAEMPRESA.cvcrm.com.br/x?id={id}&aba=1",
                       "https://a.b.cvcrm.com.br:443/r/{id}"):
            self.assertEqual(config_mod.modelo_cv(modelo), modelo, modelo)
        self.assertEqual(config_mod.modelo_cv(f"  {CV_RESERVA}  "), CV_RESERVA)

    def test_modelo_invalido_vira_none(self):
        for modelo in (None, "", "   ",
                       "http://suaempresa.cvcrm.com.br/r/{id}",            # sem https
                       "https://suaempresa.cvcrm.com.br/r/123",            # sem {id}
                       "https://cvcrm.com.br/r/{id}",               # domínio sem subdomínio
                       "https://evil.com/r/{id}",
                       "https://suaempresa.cvcrm.com.br.evil.com/r/{id}",
                       "https://evilcvcrm.com.br/r/{id}",
                       "https://suaempresa.cvcrm.com.br@evil.com/r/{id}",  # usuário na URL
                       "https://x@suaempresa.cvcrm.com.br/r/{id}",
                       "https://suaempresa.cvcrm.com.br\\@evil.com/r/{id}",
                       "https://suaempresa.cvcrm.com.br:8443/r/{id}",
                       "https://suaempresa.cvcrm.com.br/r/{id}/{outro}",
                       "https://suaempresa.cvcrm.com.br/r/{id.__class__}",
                       "https://suaempresa.cvcrm.com.br/r/{id} x",
                       "javascript:alert(1)//{id}.cvcrm.com.br"):
            self.assertIsNone(config_mod.modelo_cv(modelo), modelo)

    def test_carregar_do_ambiente(self):
        import contextlib
        with mock.patch.dict(os.environ, {}, clear=True):
            cfg = config_mod.carregar()
        # Sem variável vale o padrão do código: as telas que o João mandou (13/10/2026).
        self.assertEqual(cfg.cv_url_reserva, config_mod.CV_URL_RESERVA_PADRAO)
        self.assertIn("/reservas/{id}/administrar", cfg.cv_url_reserva)
        self.assertFalse(hasattr(cfg, "cv_url_precadastro"), "pré-cadastro saiu do painel (20/10/2026)")
        with mock.patch.dict(os.environ, {}, clear=True), \
                mock.patch.object(config_mod, "CV_URL_RESERVA_PADRAO", None):
            cfg = config_mod.carregar()
        self.assertIsNone(cfg.cv_url_reserva)
        with mock.patch.dict(os.environ, {"SECRETARIA_CV_URL_RESERVA": CV_RESERVA}, clear=True):
            cfg = config_mod.carregar()
        self.assertEqual(cfg.cv_url_reserva, CV_RESERVA)
        aviso = io.StringIO()
        with mock.patch.dict(os.environ, {"SECRETARIA_CV_URL_RESERVA": "https://evil.com/{id}"}, clear=True), \
                contextlib.redirect_stderr(aviso):
            cfg = config_mod.carregar()
        self.assertIsNone(cfg.cv_url_reserva)
        self.assertIn("SECRETARIA_CV_URL_RESERVA ignorada", aviso.getvalue())
        # o padrão do código vale quando a variável não existe
        with mock.patch.dict(os.environ, {}, clear=True), \
                mock.patch.object(config_mod, "CV_URL_RESERVA_PADRAO", CV_RESERVA):
            self.assertEqual(config_mod.carregar().cv_url_reserva, CV_RESERVA)

    def test_id_como_inteiro(self):
        self.assertEqual(dados.url_cv(CV_RESERVA, "1001"),
                         "https://suaempresa.cvcrm.com.br/gestor/comercial/reservas/1001/administrar")
        self.assertEqual(dados.url_cv(CV_PRE, 2002), "https://suaempresa.cvcrm.com.br/gestor/clientes/precadastros/2002")
        self.assertIsNone(dados.url_cv(None, 1001))
        with self.assertRaises(ValueError):
            dados.url_cv(CV_RESERVA, "1001/../x")

    def test_sem_config_tudo_none(self):
        self.entrar()
        self.assertTrue(all(i["url_cv"] is None for i in self.itens()["itens"]))
        self.assertIsNone(self.reserva()["url_cv"])

    def test_lista_detalhe_e_revisao_com_modelos_diferentes(self):
        self.client.close()
        cfg = config_mod.Config(saidas=self.saidas, db=self.cfg.db, cv_url_reserva=CV_RESERVA)
        self.app = criar_app(cfg, templates_dir=TEMPLATES_FAKE)
        self.client = TestClient(self.app, follow_redirects=False)
        csrf = self.entrar()
        urls = {(i["tipo"], i["id"]): i["url_cv"] for i in self.itens()["itens"]}
        self.assertEqual(urls[("reserva", 1001)], CV_RESERVA.format(id=1001))
        self.assertEqual(urls[("reserva", 1004)], CV_RESERVA.format(id=1004))
        self.assertNotIn(("pre_cadastro", 2002), urls)
        self.assertEqual(self.reserva()["url_cv"], CV_RESERVA.format(id=1001))

        chave = self.apont(self.reserva(), "118 dias")["chave"]
        self.client.post("/api/tratativa", json={"chave": chave, "status": "ignorado", "motivo": FP},
                         headers={"X-CSRF-Token": csrf})
        caso = self.client.get("/api/revisao").json()["grupos"][0]["casos"][0]
        self.assertEqual(caso["url_cv"], CV_RESERVA.format(id=1001))


# ==========================================================================
# Servidor: proxy confiável e IP real no limite de login (20/10/2026)
# ==========================================================================
class ProxyConfiavel(Base):
    IP_A, IP_B = "203.0.113.10", "203.0.113.20"

    def cliente(self, com_proxy: bool):
        """Como o uvicorn faz com proxy_headers=True e forwarded_allow_ips: a
        conexão chega do proxy (127.0.0.1) e o IP real vem do X-Forwarded-For."""
        app = self.app
        if com_proxy:
            with mock.patch.dict(os.environ, {"SECRETARIA_PROXY_CONFIAVEL": "127.0.0.1"}, clear=True):
                app = ProxyHeadersMiddleware(app, trusted_hosts=config_mod.proxies_confiaveis())
        cliente = TestClient(app, client=("127.0.0.1", 50000), follow_redirects=False)
        self.addCleanup(cliente.close)
        return cliente

    def login_de(self, cliente, ip, login, senha):
        pagina = cliente.get("/login")
        csrf = re.search(r'name="csrf" content="([^"]+)"', pagina.text).group(1)
        return cliente.post("/login", data={"login": login, "senha": senha, "next": "/", "csrf": csrf},
                            headers={"X-Forwarded-For": ip})

    def errar_cinco_vezes(self, cliente, ip):
        for i in range(auth.MAXIMO_FALHAS):  # logins diferentes: só o limite POR IP conta
            self.assertEqual(self.login_de(cliente, ip, f"ninguem{i}", "errada-123").status_code, 401)

    def test_dois_ips_atras_do_proxy_nao_se_bloqueiam(self):
        cliente = self.cliente(com_proxy=True)
        self.errar_cinco_vezes(cliente, self.IP_A)
        self.assertEqual(self.login_de(cliente, self.IP_A, "ana", SENHA).status_code, 401, "IP A bloqueado")
        self.assertEqual(self.login_de(cliente, self.IP_B, "ana", SENHA).status_code, 303, "IP B não")
        with self.banco.conexao() as con:
            ips = {linha[0] for linha in con.execute("SELECT ip FROM tentativas_login")}
        self.assertEqual(ips, {self.IP_A}, "o IP gravado é o do cliente, não o do proxy")

    def test_sem_proxy_x_forwarded_for_e_ignorado(self):
        cliente = self.cliente(com_proxy=False)
        self.errar_cinco_vezes(cliente, self.IP_A)
        self.assertEqual(self.login_de(cliente, self.IP_B, "ana", SENHA).status_code, 401,
                         "forjar o cabeçalho não escapa do bloqueio")

    def test_x_forwarded_for_de_quem_nao_e_o_proxy_e_ignorado(self):
        with mock.patch.dict(os.environ, {"SECRETARIA_PROXY_CONFIAVEL": "10.9.9.9"}, clear=True):
            app = ProxyHeadersMiddleware(self.app, trusted_hosts=config_mod.proxies_confiaveis())
        cliente = TestClient(app, client=("127.0.0.1", 50000), follow_redirects=False)
        self.addCleanup(cliente.close)
        self.errar_cinco_vezes(cliente, self.IP_A)
        self.assertEqual(self.login_de(cliente, self.IP_B, "ana", SENHA).status_code, 401)


# ==========================================================================
# Impacto oculto (decisão do João, 20/10/2026)
# ==========================================================================
class ImpactoOculto(Base):
    def test_padrao_oculta_tela_e_api(self):
        self.assertFalse(config_mod.Config(saidas=self.saidas, db=self.cfg.db).mostrar_impacto)
        self.client.close()
        self.app = criar_app(dataclasses.replace(self.cfg, mostrar_impacto=False), templates_dir=TEMPLATES_FAKE)
        self.client = TestClient(self.app, follow_redirects=False)
        self.entrar()
        for caminho in ("/impacto", "/api/impacto", "/api/impacto?periodo=tudo"):
            self.assertEqual(self.client.get(caminho).status_code, 404, caminho)
        self.assertEqual(self.client.get("/api/revisao").status_code, 200, "a revisão continua")

    def test_religar_pelo_ambiente(self):
        with mock.patch.dict(os.environ, {}, clear=True):
            self.assertFalse(config_mod.carregar().mostrar_impacto)
        with mock.patch.dict(os.environ, {"SECRETARIA_MOSTRAR_IMPACTO": "1"}, clear=True):
            self.assertTrue(config_mod.carregar().mostrar_impacto)


# ==========================================================================
# Admin de usuários e "trocar minha senha" (20/10/2026)
# ==========================================================================
class AdminDeUsuarios(Base):
    def setUp(self):
        super().setUp()
        auth.criar_usuario(self.banco, "chefe", "Chefe Teste", "admin", SENHA)
        self.csrf = self.entrar_como("chefe")

    def entrar_como(self, login, senha=SENHA, client=None):
        """Login e o csrf da página para onde o login mandou."""
        client = client or self.client
        r = self.tentar_login(login=login, senha=senha, client=client)
        self.assertEqual(r.status_code, 303, r.text)
        pagina = client.get(r.headers["location"])
        return re.search(r'name="csrf" content="([^"]+)"', pagina.text).group(1)

    def outro_cliente(self):
        cliente = TestClient(self.app, follow_redirects=False)
        self.addCleanup(cliente.close)
        return cliente

    def criar(self, login="bia", perfil="registro", senha="provisoria-123", csrf=None, **extra):
        return self.client.post("/admin/usuarios", data={
            "login": login, "nome": extra.get("nome", "Bia Teste"), "perfil": perfil, "senha": senha,
            "csrf": self.csrf if csrf is None else csrf})

    def acao(self, login, acao, csrf=None, client=None, **dados):
        return (client or self.client).post(f"/admin/usuarios/{login}/{acao}",
                                            data={"csrf": self.csrf if csrf is None else csrf, **dados})

    def test_nao_admin_recebe_403(self):
        outro = self.outro_cliente()
        csrf_ana = self.entrar_como("ana", client=outro)
        self.assertEqual(outro.get("/admin/usuarios").status_code, 403)
        r = outro.post("/admin/usuarios", data={"login": "x1", "nome": "X", "perfil": "admin",
                                                "senha": "senha-longa-1", "csrf": csrf_ana})
        self.assertEqual(r.status_code, 403)
        for acao, dados in (("ativo", {"ativo": "0"}), ("perfil", {"perfil": "registro"}),
                            ("senha", {"senha": "outra-senha-1"})):
            self.assertEqual(self.acao("chefe", acao, csrf=csrf_ana, client=outro, **dados).status_code, 403, acao)
        self.assertIsNone(self.banco.buscar_usuario("x1"))
        self.assertEqual(self.banco.buscar_usuario("chefe")["perfil"], "admin")
        self.assertEqual(self.client.get("/admin/usuarios").status_code, 200, "o admin vê")

    def test_criar_com_o_hash_do_terminal_e_registro(self):
        r = self.criar()
        self.assertEqual((r.status_code, r.headers["location"]), (303, "/admin/usuarios?feito=criado"))
        u = self.banco.buscar_usuario("bia")
        self.assertEqual((u["perfil"], u["ativo"], u["trocar_senha"], u["criado_por"]), ("registro", 1, 1, "chefe"))
        self.assertEqual(u["iteracoes"], auth.ITERACOES)
        self.assertTrue(auth.senha_confere("provisoria-123", u["hash"], u["sal"], u["iteracoes"]))
        pagina = self.client.get("/admin/usuarios").text
        self.assertIn("[bia:registro:1:1]", pagina)
        self.assertIn("[bia:criado:chefe]", pagina)
        self.assertNotIn("hash", self.banco.listar_usuarios()[0])

    def test_criar_duplicado_invalido_e_sem_csrf(self):
        self.criar()
        r = self.criar(login="BIA", nome="Outra")
        self.assertEqual(r.status_code, 409)
        self.assertIn("já existe", r.text)
        self.assertEqual(self.banco.buscar_usuario("bia")["nome"], "Bia Teste")
        for dados, status in (({"login": "c1", "senha": "curta"}, 400), ({"login": "c2", "perfil": "chefe"}, 400),
                              ({"login": "com espaço"}, 400), ({"login": "c3", "csrf": "errado"}, 403)):
            with self.subTest(dados=dados):
                self.assertEqual(self.criar(**dados).status_code, status)
        self.assertEqual([u["login"] for u in self.banco.listar_usuarios()], ["ana", "bia", "chefe"])

    def test_troca_obrigatoria_no_primeiro_login(self):
        self.criar()
        bia = self.outro_cliente()
        r = self.tentar_login(login="bia", senha="provisoria-123", next="/revisao", client=bia)
        self.assertEqual((r.status_code, r.headers["location"]), (303, "/trocar-senha"))
        self.assertEqual((bia.get("/").status_code, bia.get("/").headers["location"]), (303, "/trocar-senha"))
        self.assertEqual(bia.get("/api/reservas").status_code, 403)
        pagina = bia.get("/trocar-senha")
        self.assertIn("obrigatoria=True", pagina.text)
        csrf = re.search(r'name="csrf" content="([^"]+)"', pagina.text).group(1)

        def trocar(atual, nova, confirmacao=None, csrf_=csrf):
            return bia.post("/trocar-senha", data={"atual": atual, "nova": nova, "csrf": csrf_,
                                                   "confirmacao": nova if confirmacao is None else confirmacao})
        self.assertEqual(trocar("provisoria-123", "minha-senha-9", csrf_="x").status_code, 403)
        self.assertEqual(trocar("errada-000", "minha-senha-9").status_code, 400)
        self.assertEqual(trocar("provisoria-123", "minha-senha-9", "outra-coisa").status_code, 400)
        self.assertEqual(trocar("provisoria-123", "provisoria-123").status_code, 400)
        self.assertEqual(trocar("provisoria-123", "curta").status_code, 400)
        r = trocar("provisoria-123", "minha-senha-9")
        self.assertEqual((r.status_code, r.headers["location"]), (303, "/trocar-senha?feito=1"))
        self.assertEqual(bia.get("/").status_code, 200, "a sessão desta aba continua")
        self.assertEqual(self.banco.buscar_usuario("bia")["trocar_senha"], 0)
        self.assertIsNone(auth.autenticar(self.banco, "bia", "provisoria-123"))
        self.assertIsNotNone(auth.autenticar(self.banco, "bia", "minha-senha-9"))
        self.assertIn("[bia:senha_trocada:bia]", self.client.get("/admin/usuarios").text)

    def test_trocar_senha_para_todos_derruba_as_outras_sessoes(self):
        outra = self.outro_cliente()
        self.entrar_como("ana", client=outra)
        aqui = self.outro_cliente()
        csrf = self.entrar_como("ana", client=aqui)
        r = aqui.post("/trocar-senha", data={"atual": SENHA, "nova": "nova-senha-ana", "confirmacao": "nova-senha-ana",
                                             "csrf": csrf})
        self.assertEqual(r.status_code, 303)
        self.assertEqual(aqui.get("/api/reservas").status_code, 200)
        self.assertEqual(outra.get("/api/reservas").status_code, 401)

    def test_desativar_derruba_a_sessao_e_reativar(self):
        auth.criar_usuario(self.banco, "bia", "Bia", "registro", "senha-da-bia")
        bia = self.outro_cliente()
        self.entrar_como("bia", "senha-da-bia", client=bia)
        self.assertEqual(bia.get("/api/reservas").status_code, 200)
        r = self.acao("bia", "ativo", ativo="0")
        self.assertEqual((r.status_code, r.headers["location"]), (303, "/admin/usuarios?feito=desativado"))
        self.assertEqual(bia.get("/api/reservas").status_code, 401)
        with self.banco.conexao() as con:
            self.assertEqual(con.execute("SELECT COUNT(*) FROM sessoes WHERE login = 'bia'").fetchone()[0], 0)
        self.assertEqual(self.tentar_login(login="bia", senha="senha-da-bia", client=bia).status_code, 401)
        self.acao("bia", "ativo", ativo="1")
        self.assertEqual(self.tentar_login(login="bia", senha="senha-da-bia", client=bia).status_code, 303)
        u = self.banco.buscar_usuario("bia")
        self.assertEqual((u["ativo"], u["alterado_por"]), (1, "chefe"))
        pagina = self.client.get("/admin/usuarios").text
        self.assertIn("[bia:desativado:chefe]", pagina)
        self.assertIn("[bia:reativado:chefe]", pagina)

    def test_admin_nao_se_desativa_nem_tira_o_proprio_admin(self):
        self.assertEqual(self.acao("chefe", "ativo", ativo="0").status_code, 400)
        self.assertEqual(self.acao("chefe", "perfil", perfil="gerente").status_code, 400)
        self.assertEqual(self.acao("chefe", "senha", senha="nova-senha-123").status_code, 400)
        u = self.banco.buscar_usuario("chefe")
        self.assertEqual((u["ativo"], u["perfil"], u["trocar_senha"]), (1, "admin", 0))
        self.assertEqual(self.acao("ana", "perfil", perfil="admin").status_code, 303)
        self.assertEqual(self.banco.buscar_usuario("ana")["perfil"], "admin")
        self.assertEqual(self.acao("ninguem", "ativo", ativo="0").status_code, 404)

    def test_redefinir_senha_derruba_e_obriga_a_troca(self):
        bia_antes = self.outro_cliente()
        auth.criar_usuario(self.banco, "bia", "Bia", "registro", "senha-da-bia")
        self.entrar_como("bia", "senha-da-bia", client=bia_antes)
        r = self.acao("bia", "senha", senha="provisoria-456")
        self.assertEqual((r.status_code, r.headers["location"]), (303, "/admin/usuarios?feito=senha"))
        self.assertEqual(bia_antes.get("/api/reservas").status_code, 401)
        self.assertEqual(self.acao("bia", "senha", senha="curta").status_code, 400)
        r = self.tentar_login(login="bia", senha="provisoria-456", client=self.outro_cliente())
        self.assertEqual(r.headers["location"], "/trocar-senha")

    def test_ultimo_acesso_e_terminal_sem_troca(self):
        self.assertIsNotNone(self.banco.buscar_usuario("chefe")["ultimo_acesso"])
        self.assertIsNone(self.banco.buscar_usuario("ana")["ultimo_acesso"])
        u = self.banco.buscar_usuario("ana")  # criada como no terminal
        self.assertEqual((u["criado_por"], u["trocar_senha"]), ("terminal", 0))
        r = self.tentar_login(login="ana", next="/revisao", client=self.outro_cliente())
        self.assertEqual(r.headers["location"], "/revisao")


# ==========================================================================
# analisar_documentos --todas / --reler-telas: só reserva (20/10/2026)
# ==========================================================================
class AnaliseSoDeReserva(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="so_reserva_"))
        nomes = ("PASTA_SAIDAS", "PASTA_SAIDAS_RESERVA", "PASTA_SAIDAS_PRECADASTRO",
                 "processar_pasta_cliente", "reler_telas_das_pastas", "verificar_chave_ia")
        guardados = {n: getattr(ad, n) for n in nomes}
        self.addCleanup(lambda: [setattr(ad, n, v) for n, v in guardados.items()])
        self.addCleanup(shutil.rmtree, self.tmp, True)
        ad.PASTA_SAIDAS = self.tmp
        ad.PASTA_SAIDAS_RESERVA = self.tmp / "reserva"
        ad.PASTA_SAIDAS_PRECADASTRO = self.tmp / "pre_cadastro"
        for p in ("reserva/A-reserva-7", "reserva/B-2o Promissario Comprador-reserva-7",
                  "pre_cadastro/C-7", "cache_ia/ab"):
            (self.tmp / p).mkdir(parents=True)
        ad.verificar_chave_ia = lambda: None
        self.vistas = []

    def test_todas_so_reserva(self):
        ad.processar_pasta_cliente = lambda pasta, **k: self.vistas.append(pasta.name)
        with contextlib.redirect_stdout(io.StringIO()):
            ad.processar_todas_as_pastas()
        self.assertEqual(self.vistas, ["A-reserva-7", "B-2o Promissario Comprador-reserva-7"])

    def test_reler_telas_so_reserva(self):
        ad.reler_telas_das_pastas = lambda pastas: self.vistas.append(sorted(p.name for p in pastas))
        for argv in (["--todas", "--reler-telas"], ["7", "--reler-telas"]):
            with mock.patch.object(sys, "argv", ["analisar_documentos.py", *argv]):
                ad.main()
        esperado = ["A-reserva-7", "B-2o Promissario Comprador-reserva-7"]
        self.assertEqual(self.vistas, [esperado, esperado], "pré-cadastro 7 (outra pessoa) fica de fora")


# ==========================================================================
# Linha de comando
# ==========================================================================
class LinhaDeComando(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="painel_cli_"))
        self.env = mock.patch.dict(os.environ, {
            "SECRETARIA_SAIDAS": str(self.tmp / "saidas"),
            "SECRETARIA_PAINEL_DB": str(self.tmp / "db" / "painel.db"),
        })
        self.env.start()

    def tearDown(self):
        self.env.stop()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_servir_so_em_127_0_0_1(self):
        import uvicorn
        with mock.patch.object(uvicorn, "run") as run:
            cli.main([])
            cli.main(["servir"])
        self.assertEqual(run.call_count, 2)
        for chamada in run.call_args_list:
            self.assertEqual(chamada.kwargs["host"], "127.0.0.1")
            self.assertEqual(chamada.kwargs["port"], 8765)
            self.assertIs(chamada.kwargs["access_log"], False)  # q e ids vão na URL
        self.assertEqual(config_mod.HOST, "127.0.0.1")

    def servir(self, env):
        import uvicorn
        env = {**{k: os.environ[k] for k in ("SECRETARIA_SAIDAS", "SECRETARIA_PAINEL_DB")}, **env}
        with mock.patch.dict(os.environ, env, clear=True), mock.patch.object(uvicorn, "run") as run, \
                contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()) as erro:
            try:
                cli.main(["servir"])
                codigo = None
            except SystemExit as exc:
                codigo = exc.code
        return run, codigo, erro.getvalue()

    def test_padrao_e_o_do_pc_local(self):
        run, codigo, _ = self.servir({})
        self.assertIsNone(codigo)
        k = run.call_args.kwargs
        self.assertEqual((k["host"], k["port"], k["proxy_headers"], k["forwarded_allow_ips"]),
                         ("127.0.0.1", 8765, False, None))

    def test_servidor_com_proxy_na_mesma_maquina(self):
        run, codigo, erro = self.servir({"SECRETARIA_PORTA": "9000", "SECRETARIA_PROXY_CONFIAVEL": "127.0.0.1"})
        k = run.call_args.kwargs
        self.assertEqual((k["host"], k["port"], k["proxy_headers"], k["forwarded_allow_ips"], k["access_log"]),
                         ("127.0.0.1", 9000, True, ["127.0.0.1"], False))
        self.assertNotIn("ATENÇÃO", erro)

    def test_proxy_em_outra_maquina_so_com_permissao_e_aviso(self):
        env = {"SECRETARIA_HOST": "0.0.0.0", "SECRETARIA_PROXY_CONFIAVEL": "10.0.0.50"}
        run, codigo, erro = self.servir(env)
        self.assertEqual(codigo, 2)
        self.assertIn("SECRETARIA_PERMITIR_REDE=1", erro)
        run.assert_not_called()
        run, codigo, erro = self.servir({**env, "SECRETARIA_PERMITIR_REDE": "1"})
        self.assertIsNone(codigo)
        self.assertEqual((run.call_args.kwargs["host"], run.call_args.kwargs["forwarded_allow_ips"]),
                         ("0.0.0.0", ["10.0.0.50"]))
        self.assertIn("firewall SÓ para o IP do proxy (10.0.0.50)", erro)

    def test_recusas_explicadas_no_log(self):
        for env, trecho in (({"SECRETARIA_HOST": "0.0.0.0"}, "SECRETARIA_PROXY_CONFIAVEL precisa"),
                            ({"SECRETARIA_HOST": "10.0.0.10"}, "atrás de um proxy TLS"),
                            # P1 da Turing: proxy local + 0.0.0.0 abriria a porta para contornar o Caddy
                            ({"SECRETARIA_HOST": "0.0.0.0", "SECRETARIA_PROXY_CONFIAVEL": "127.0.0.1",
                              "SECRETARIA_PERMITIR_REDE": "1"}, "o proxy confiável é desta máquina"),
                            ({"SECRETARIA_HOST": "10.0.0.48", "SECRETARIA_PROXY_CONFIAVEL": "::1"},
                             "o proxy confiável é desta máquina"),
                            ({"SECRETARIA_PORTA": "abc"}, "SECRETARIA_PORTA"),
                            ({"SECRETARIA_PORTA": "70000"}, "1 a 65535"),
                            ({"SECRETARIA_PROXY_CONFIAVEL": "*"}, "forjar"),
                            ({"SECRETARIA_PROXY_CONFIAVEL": "proxy.local"}, "não é um IP")):
            with self.subTest(env=env):
                run, codigo, erro = self.servir(env)
                self.assertEqual(codigo, 2)
                self.assertIn("Painel NÃO subiu", erro)
                self.assertIn(trecho, erro)
                run.assert_not_called()

    @unittest.skipUnless(sys.platform == "win32", "script do Windows")
    def test_servico_exige_conta_propria(self):
        # P2 da Turing: sem a conta, o NSSM rodaria o painel como LocalSystem.
        script = Path(__file__).resolve().parents[2] / "ferramentas" / "servidor" / "servico-painel.cmd"
        casos = (("", "informe a conta"), ('"LocalSystem"', "conta do sistema"),
                 (r'"NT AUTHORITY\SYSTEM"', "conta do sistema"),
                 # 21/10/2026: conta virtual só com o nome do serviço; local/domínio exige a senha
                 (r'"NT SERVICE\Outro"', "nome do servico"),
                 (r'".\svc_secretaria"', "precisa da SENHA"),
                 (r'"DOMINIO\svc_secretaria"', "precisa da SENHA"))
        for argumentos, motivo in casos:
            with self.subTest(argumentos=argumentos):
                # como digitado no prompt: o caminho do projeto tem espaço
                r = subprocess.run(f'cmd /s /c ""{script}" {argumentos}"', capture_output=True, text=True,
                                   timeout=30)
                self.assertEqual(r.returncode, 1, r.stdout + r.stderr)
                self.assertIn("ERRO", r.stdout)
                self.assertIn(motivo, r.stdout)
                self.assertNotIn("nssm", r.stderr.lower(), "parou antes de chamar o NSSM")

    @unittest.skipUnless(sys.platform == "win32", "script do Windows")
    def test_conferencia_da_conta_independe_do_idioma(self):
        # 21/10/2026: no servidor em português o rótulo SERVICE_START_NAME sai traduzido; a
        # conferência procura o NOME da conta (e não remove um serviço que está certo).
        script = Path(__file__).resolve().parents[2] / "ferramentas" / "servidor" / "conta-do-servico.cmd"
        modelo_pt = ("[SC] QueryServiceConfig ÊXITO\r\n\r\nNOME_DO_SERVIÇO: SecretariaPainel\r\n"
                     "        TIPO                   : 10  WIN32_OWN_PROCESS\r\n"
                     "        TIPO_DE_INÍCIO         : 2   AUTO_START\r\n"
                     "        NOME_DO_CAMINHO_BINÁRIO: C:\\Ferramentas\\nssm.exe\r\n"
                     "        NOME_DE_EXIBIÇÃO       : Secretaria de Vendas - painel\r\n"
                     "        NOME_DE_INICIO_DO_SERVICO : {conta}\r\n")
        modelo_en = ("[SC] QueryServiceConfig SUCCESS\r\n\r\nSERVICE_NAME: SecretariaPainel\r\n"
                     "        TYPE               : 10  WIN32_OWN_PROCESS\r\n"
                     "        SERVICE_START_NAME : {conta}\r\n")
        casos = ((modelo_pt, "NT Service\\SecretariaPainel", 0), (modelo_en, "NT SERVICE\\SecretariaPainel", 0),
                 (modelo_pt, "LocalSystem", 1), (modelo_pt, "NT Service\\Outro", 1), (modelo_en, "LocalSystem", 1))
        for i, (modelo, conta_no_sc, esperado) in enumerate(casos):
            with self.subTest(conta_no_sc=conta_no_sc, idioma="pt" if modelo is modelo_pt else "en"):
                saida = self.tmp / f"qc-{i}.txt"
                saida.write_bytes(modelo.format(conta=conta_no_sc).encode("cp850"))  # console do Windows em pt-BR
                r = subprocess.run(f'cmd /s /c ""{script}" SecretariaPainel "NT SERVICE\\SecretariaPainel" '
                                   f'"{saida}""', capture_output=True, text=True, timeout=60)
                self.assertEqual(r.returncode, esperado, r.stdout + r.stderr)

    def test_criar_usuario_pede_senha_no_terminal(self):
        with mock.patch.object(cli, "getpass", side_effect=[SENHA, SENHA]):
            cli.main(["criar-usuario", "--login", "Bia", "--nome", "Bia Teste", "--perfil", "registro"])
        u = banco_mod.Banco(Path(os.environ["SECRETARIA_PAINEL_DB"])).buscar_usuario("bia")
        self.assertEqual((u["nome"], u["perfil"]), ("Bia Teste", "registro"))
        with mock.patch.object(cli, "getpass", side_effect=[SENHA, "outra"]):
            with self.assertRaises(SystemExit):
                cli.main(["criar-usuario", "--login", "cia", "--nome", "C", "--perfil", "admin"])
        with self.assertRaises(SystemExit):  # argparse: perfil fora da lista
            cli.main(["criar-usuario", "--login", "d", "--nome", "D", "--perfil", "chefe"])

    def test_config_padrao_e_minutos(self):
        # Sem a variável vale a estimativa do João (30 a 50 min, meio = 40).
        with mock.patch.dict(os.environ, {"SECRETARIA_MINUTOS_MANUAIS_POR_RESERVA": ""}):
            self.assertEqual(config_mod.carregar().minutos_manuais_por_reserva,
                             config_mod.MINUTOS_MANUAIS_PADRAO)
        with mock.patch.dict(os.environ, {"SECRETARIA_MINUTOS_MANUAIS_POR_RESERVA": "45"}):
            self.assertEqual(config_mod.carregar().minutos_manuais_por_reserva, 45)
        with mock.patch.dict(os.environ, {}, clear=True):
            cfg = config_mod.carregar()
            self.assertEqual(cfg.saidas, config_mod.RAIZ_PROJETO / "saidas")
            self.assertEqual(cfg.db, config_mod.RAIZ_PROJETO / "painel_dados" / "painel.db")


# ==========================================================================
# Contexto da análise (analisar_documentos.py)
# ==========================================================================
REGISTRO_PLANO = {  # /v1/cvdw/reservas
    "idreserva": 777, "idprecadastro": 5, "tipovenda": "Financiamento Bancário",
    "empreendimento": "Residencial Sintético", "unidade": "101", "bloco": "A", "etapa": "1",
    "imobiliaria": "Imob Fictícia", "corretor": "Corretor Fictício", "situacao": "Venda Validada",
    "cliente": "FULANA FICTICIA"}
REGISTRO_ANINHADO = {"777": {  # /v1/comercial/reservas/{id}, embrulhado no id
    "idprecadastro": 5, "tipovenda": "Venda Direta",
    "situacao": {"situacao": "Em Validação", "grupo": "G"},
    "unidade": {"empreendimento": "Residencial Sintético", "unidade": "202", "bloco": "B", "etapa": "2"},
    "imobiliaria": {"nome": "Imob Fictícia"}, "corretor": {"corretor": "Corretor Fictício", "imobiliaria": "I"}}}


def linhas_do_contexto(xlsx: Path) -> list[tuple]:
    wb = load_workbook(xlsx, read_only=True)
    try:
        return [tuple(linha[:2]) for linha in wb[ad.ABA_CONTEXTO].iter_rows(values_only=True)]
    finally:
        wb.close()


class ContextoDaAnalise(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="painel_contexto_"))
        self.saidas = self.tmp / "saidas"
        # cache persistente da IA (19/10/2026) num diretório do teste, não em saidas/ do projeto
        cache = ad.PASTA_CACHE_IA
        self.addCleanup(setattr, ad, "PASTA_CACHE_IA", cache)
        ad.PASTA_CACHE_IA = self.tmp / "cache_ia"
        consumo_ia = ad.PASTA_CONSUMO_IA  # consumo da IA (21/10/2026), idem
        self.addCleanup(setattr, ad, "PASTA_CONSUMO_IA", consumo_ia)
        ad.PASTA_CONSUMO_IA = self.tmp / "painel_dados"
        self.addCleanup(setattr, ad, "ARQUIVO_ESTADO_IA", ad.ARQUIVO_ESTADO_IA)  # o disjuntor, idem
        ad.ARQUIVO_ESTADO_IA = self.tmp / "painel_dados" / "ia_estado.json"

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def test_formato_plano(self):
        ctx = ad.contexto_da_reserva(REGISTRO_PLANO, ad.MODALIDADE_FINANCIAMENTO)
        self.assertEqual(ctx, {
            "modalidade": "Financiamento", "tipovenda": "Financiamento Bancário",
            "empreendimento": "Residencial Sintético", "unidade": "101", "bloco": "A", "etapa": "1",
            "imobiliaria": "Imob Fictícia", "corretor": "Corretor Fictício", "situacao": "Venda Validada"})

    def test_formato_aninhado_embrulhado_e_solto(self):
        esperado = {
            "modalidade": "Venda direta", "tipovenda": "Venda Direta",
            "empreendimento": "Residencial Sintético", "unidade": "202", "bloco": "B", "etapa": "2",
            "imobiliaria": "Imob Fictícia", "corretor": "Corretor Fictício", "situacao": "Em Validação"}
        self.assertEqual(ad.contexto_da_reserva(REGISTRO_ANINHADO, ad.MODALIDADE_VENDA_DIRETA), esperado)
        self.assertEqual(ad.contexto_da_reserva(REGISTRO_ANINHADO["777"], ad.MODALIDADE_VENDA_DIRETA), esperado)

    def test_registro_none_e_campos_ausentes(self):
        self.assertEqual(ad.contexto_da_reserva(None, None), {"modalidade": ""})
        self.assertEqual(ad.contexto_da_reserva({}, None), {"modalidade": ""})
        ctx = ad.contexto_da_reserva({"tipovenda": "Loteamento", "corretor": {}, "situacao": None},
                                     ad.MODALIDADE_LOTEAMENTO)
        self.assertEqual(ctx["modalidade"], "Loteamento")
        self.assertEqual((ctx["empreendimento"], ctx["corretor"], ctx["situacao"]), ("", "", ""))

    def test_aba_regravada_sem_duplicar(self):
        xlsx = criar_pasta(self.saidas, "reserva", "X-reserva-777", {"nome": "X"}, [])
        ad.gravar_aba_contexto(xlsx, ad.contexto_da_reserva(REGISTRO_PLANO, ad.MODALIDADE_FINANCIAMENTO))
        ad.gravar_aba_contexto(xlsx, ad.contexto_da_reserva(REGISTRO_ANINHADO, ad.MODALIDADE_VENDA_DIRETA))
        linhas = linhas_do_contexto(xlsx)
        self.assertEqual(linhas[0], ("Campo", "Valor"))
        self.assertEqual([c for c, _ in linhas[1:]], list(ad.CAMPOS_CONTEXTO))
        valores = dict(linhas[1:])
        self.assertEqual((valores["modalidade"], valores["unidade"]), ("Venda direta", "202"))
        self.assertRegex(valores["analisado_em"], r"^\d{2}/\d{2}/\d{4} \d{2}:\d{2}$")
        self.assertEqual(load_workbook(xlsx).sheetnames.count(ad.ABA_CONTEXTO), 1)

    def test_valor_com_igual_nao_vira_formula(self):
        xlsx = criar_pasta(self.saidas, "reserva", "X-reserva-777", {"nome": "X"}, [])
        ad.gravar_aba_contexto(xlsx, {"modalidade": "", "corretor": "=HYPERLINK(\"x\")"})
        wb = load_workbook(xlsx)
        celula = wb[ad.ABA_CONTEXTO]["B3"]
        self.assertEqual((celula.value, celula.data_type), ("=HYPERLINK(\"x\")", "s"))

    # ---- processar_pasta_cliente, com CV falso
    def _com_cv_falso(self, registro):
        nomes = ("carregar_config", "buscar_registro_da_reserva", "buscar_precadastro")
        originais = {n: getattr(ad, n) for n in nomes}
        self.addCleanup(lambda: [setattr(ad, n, v) for n, v in originais.items()])
        ad.carregar_config = lambda: {"gemini_api_key": "chave-falsa"}
        ad.buscar_registro_da_reserva = lambda cfg, n, eh_idreserva=False: registro
        ad.buscar_precadastro = lambda cfg, n: {"renda_total": 3000, "renda_cliente_principal": 1500}

    def _processar(self, pasta: Path):
        import contextlib
        with contextlib.redirect_stdout(io.StringIO()):
            ad.processar_pasta_cliente(pasta, associados_consultados=[], falha_associados=None)

    def test_processar_grava_a_aba_sem_documento_novo_e_o_painel_le(self):
        self._com_cv_falso(REGISTRO_PLANO)
        xlsx = criar_pasta(
            self.saidas, "reserva", "FULANA-FICTICIA-reserva-777", {"nome": "FULANA FICTICIA"},
            [{"tipo": "Comprovante de Endereço", "nome": "c", "arquivo_local": "c.pdf"}],
            resultados=[dict(resultado("c.pdf", "Comprovante de Endereço", "Conta de luz"),
                             doc_ia={"tipo_identificado": "Conta de luz", "_leitura_completa": True})],
            arquivos=("c.pdf",))
        self._processar(xlsx.parent)
        self._processar(xlsx.parent)  # de novo: não duplica
        valores = dict(linhas_do_contexto(xlsx)[1:])
        self.assertEqual(len(linhas_do_contexto(xlsx)), 1 + len(ad.CAMPOS_CONTEXTO))
        self.assertEqual((valores["modalidade"], valores["imobiliaria"]), ("Financiamento", "Imob Fictícia"))

        reservas = dados.Repositorio(self.saidas).atualizar()
        item = dados.montar_item(reservas[0], {})
        self.assertEqual((item["modalidade"], item["empreendimento"], item["imobiliaria"],
                          item["corretor"], item["situacao"]),
                         ("Financiamento", "Residencial Sintético", "Imob Fictícia",
                          "Corretor Fictício", "Venda Validada"))

    def test_pre_cadastro_sem_reserva_grava_so_modalidade_vazia(self):
        self._com_cv_falso(None)
        xlsx = criar_pasta(
            self.saidas, "pre_cadastro", "CICLANA-FICTICIA-2002", {"nome": "CICLANA FICTICIA"}, [],
            resultados=[])
        self._processar(xlsx.parent)
        linhas = linhas_do_contexto(xlsx)
        self.assertEqual([c for c, _ in linhas[1:]], ["modalidade", "analisado_em"])
        self.assertIn(linhas[1][1], ("", None))

    def test_processar_grava_a_aba_com_documento_novo(self):
        """Caminho com leitura nova. A IA é FALSA: o cliente Gemini continua
        bloqueado pela blindagem, aqui só se troca o módulo por um objeto vazio
        e a leitura por um dicionário pronto."""
        import types

        import google
        # Já desembrulhado: buscar_registro_da_reserva() só devolve o registro da
        # busca direta quando 'tipovenda' está no primeiro nível.
        self._com_cv_falso(REGISTRO_ANINHADO["777"])
        # O cliente nasce com timeout (criar_cliente_ia, revisão da Forja de
        # 19/10/2026): o falso precisa de types.HttpOptions e de Client(**k).
        falso = types.SimpleNamespace(Client=lambda **k: object(),
                                      types=types.SimpleNamespace(HttpOptions=lambda **k: k))
        leitura = {"tipo_identificado": "Conta de luz", "nome_encontrado": "FULANA FICTICIA"}
        xlsx = criar_pasta(
            self.saidas, "reserva", "FULANA-FICTICIA-reserva-777", {"nome": "FULANA FICTICIA"},
            [{"tipo": "Comprovante de Endereço", "nome": "c", "arquivo_local": "c.pdf"}],
            arquivos=("c.pdf",))
        with mock.patch.object(google, "genai", falso), \
                mock.patch.dict(sys.modules, {"google.genai": falso, "google.genai.types": falso.types}), \
                mock.patch.object(ad, "extrair_dados_do_documento", return_value=leitura):
            self._processar(xlsx.parent)
        valores = dict(linhas_do_contexto(xlsx)[1:])
        self.assertEqual((valores["modalidade"], valores["situacao"], valores["unidade"]),
                         ("Venda direta", "Em Validação", "202"))


if __name__ == "__main__":
    unittest.main()
