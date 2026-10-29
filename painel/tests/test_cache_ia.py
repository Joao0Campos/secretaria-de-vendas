"""Testes do cache persistente da leitura da IA (19/10/2026) - offline.

Rodar (da pasta do projeto ou de qualquer lugar):
    <python do .venv do projeto> -B painel/tests/test_cache_ia.py -v

O cache fica num diretório temporário; a IA é substituída por uma função que
conta as chamadas. Mesma blindagem do test_painel.
"""
from __future__ import annotations

import sys
from pathlib import Path

_AQUI = Path(__file__).resolve().parent
if str(_AQUI) not in sys.path:
    sys.path.insert(0, str(_AQUI))

import test_painel as tp  # noqa: E402  (blindagem antes do projeto)

import contextlib  # noqa: E402
import io  # noqa: E402
import json  # noqa: E402
import os  # noqa: E402
import shutil  # noqa: E402
import tempfile  # noqa: E402
import types  # noqa: E402
import unittest  # noqa: E402
from unittest import mock  # noqa: E402

import analisar_documentos as ad  # noqa: E402
import mvp_v1_validador as v1  # noqa: E402

LEITURA = {"tipo_identificado": "Conta de luz", "nome_encontrado": "FULANA FICTICIA", "cpf_encontrado": None,
           "data_emissao": None, "data_validade": None, "endereco_encontrado": "Rua Ficticia 1",
           "mes_referencia": "09/2026", "legivel": True, "confianca_campos": {"nome_encontrado": "alta"}}


class BaseCache(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="cache_ia_"))
        self.saidas = self.tmp / "saidas"
        guardados = {n: getattr(ad, n) for n in ("PASTA_CACHE_IA", "PASTA_CONSUMO_IA", "ARQUIVO_ESTADO_IA",
                                                  "extrair_dados_do_documento",
                                                  "carregar_config", "buscar_registro_da_reserva",
                                                  "buscar_precadastro", "listar_pastas_de_cliente")}
        self.addCleanup(lambda: [setattr(ad, n, v) for n, v in guardados.items()])
        ad.PASTA_CACHE_IA = self.tmp / "cache_ia"
        ad.PASTA_CONSUMO_IA = self.tmp / "painel_dados"
        ad.ARQUIVO_ESTADO_IA = self.tmp / "painel_dados" / "ia_estado.json"  # o disjuntor (21/10/2026)
        self.chamadas_ia = []

        def ia_falsa(client, modelo, caminho, tipo_cv):
            self.chamadas_ia.append(Path(caminho).name)
            return dict(LEITURA)
        ad.extrair_dados_do_documento = ia_falsa
        ad.carregar_config = lambda: {"gemini_api_key": "chave-falsa"}
        ad.buscar_registro_da_reserva = lambda cfg, n, eh_idreserva=False: tp.REGISTRO_PLANO
        ad.buscar_precadastro = lambda cfg, n: {"renda_total": 3000, "renda_cliente_principal": 1500}
        ad.criar_cliente_ia = lambda chave: object()
        self.addCleanup(setattr, ad, "criar_cliente_ia", ad.criar_cliente_ia)
        self.addCleanup(shutil.rmtree, self.tmp, True)

    def pasta(self, nome="FULANA-FICTICIA-reserva-777", conteudo=b"%PDF conta de luz 1", arquivo="c.pdf"):
        """Como o processar_reserva_completa() faz: gerar_excel() SEMPRE recria a
        planilha (sem a aba de leitura) e o arquivo é baixado de novo."""
        pasta = self.saidas / "reserva" / nome
        pasta.mkdir(parents=True, exist_ok=True)
        v1.gerar_excel({"nome": "FULANA FICTICIA"},
                       [{"tipo": "Comprovante de Endereço", "nome": "c", "arquivo_local": arquivo}],
                       pasta / f"{nome}.xlsx")
        (pasta / arquivo).write_bytes(conteudo)
        return pasta

    def processar(self, pasta, **k):
        with contextlib.redirect_stdout(io.StringIO()) as saida:
            ad.processar_pasta_cliente(pasta, associados_consultados=[], falha_associados=None, **k)
        return saida.getvalue()


class Cache(BaseCache):
    def test_mesmo_conteudo_nunca_e_relido_mesmo_com_planilha_recriada(self):
        self.processar(self.pasta())
        self.assertEqual(self.chamadas_ia, ["c.pdf"])
        saida = self.processar(self.pasta())  # planilha recriada, arquivo baixado de novo
        self.assertEqual(self.chamadas_ia, ["c.pdf"], "0 chamadas na 2ª vez")
        self.assertIn("cache da IA: 1 do cache, 0 para a IA", saida)
        self.assertIn("(do cache)", saida)
        lido = ad.ler_resultados_anteriores(
            self.saidas / "reserva" / "FULANA-FICTICIA-reserva-777" / "FULANA-FICTICIA-reserva-777.xlsx")
        self.assertEqual(lido["c.pdf"]["doc_ia"]["nome_encontrado"], "FULANA FICTICIA",
                         "a leitura do cache vai para a planilha como as outras")

    def test_outro_nome_e_outra_pasta_tambem_vem_do_cache(self):
        self.processar(self.pasta())
        self.processar(self.pasta(nome="BELTRANO-2o Promissario Comprador-reserva-777", arquivo="outro nome.pdf"))
        self.assertEqual(self.chamadas_ia, ["c.pdf"])

    def test_conteudo_diferente_vai_para_a_ia(self):
        self.processar(self.pasta())
        self.processar(self.pasta(conteudo=b"%PDF conta de luz TROCADA"))
        self.assertEqual(self.chamadas_ia, ["c.pdf", "c.pdf"])

    def test_forcar_rele_e_atualiza_o_cache(self):
        self.processar(self.pasta())
        LEITURA_NOVA = dict(LEITURA, nome_encontrado="LIDA DE NOVO")
        ad.extrair_dados_do_documento = lambda *a: (self.chamadas_ia.append("forcada"), dict(LEITURA_NOVA))[1]
        self.processar(self.pasta(), forcar=True)
        self.assertEqual(self.chamadas_ia, ["c.pdf", "forcada"])
        self.assertEqual(ad.leitura_do_cache(self.pasta() / "c.pdf")["nome_encontrado"], "LIDA DE NOVO")

    def test_erro_da_ia_nao_entra_no_cache(self):
        ad.extrair_dados_do_documento = lambda *a: (self.chamadas_ia.append("x"), {"erro": "timeout"})[1]
        self.processar(self.pasta())
        self.processar(self.pasta())
        self.assertEqual(self.chamadas_ia, ["x", "x"], "falha não vira leitura guardada")
        self.assertEqual(list((self.tmp / "cache_ia").rglob("*.json")), [])

    def test_gravacao_atomica_e_formato(self):
        arquivo = self.tmp / "a.pdf"
        arquivo.write_bytes(b"conteudo")
        self.assertTrue(ad.guardar_no_cache(arquivo, dict(LEITURA, _leitura_completa=True), "Comprovante", "ia"))
        guardados = list((self.tmp / "cache_ia").rglob("*"))
        self.assertEqual([p.suffix for p in guardados if p.is_file()], [".json"], "nenhum .tmp esquecido")
        dados = json.loads(next(p for p in guardados if p.is_file()).read_text(encoding="utf-8"))
        self.assertEqual((dados["versao_prompt"], dados["origem"], "_leitura_completa" in dados["leitura"]),
                         (ad.VERSAO_PROMPT, "ia", False))
        self.assertEqual(dados["sha256"], ad.hash_do_arquivo(arquivo))
        self.assertFalse(ad.guardar_no_cache(arquivo, dict(LEITURA, nome_encontrado="X"), "", "planilha",
                                             sobrescrever=False), "carga inicial não sobrescreve")
        self.assertEqual(ad.leitura_do_cache(arquivo)["nome_encontrado"], "FULANA FICTICIA")

    def test_alimentar_com_as_planilhas(self):
        pasta = self.pasta()
        self.processar(pasta)  # planilha com leitura completa
        shutil.rmtree(self.tmp / "cache_ia")
        ad.listar_pastas_de_cliente = lambda: [pasta]
        c = ad.alimentar_cache_com_as_planilhas()
        self.assertEqual((c["gravadas"], c["ja_estavam"]), (1, 0))
        self.assertEqual(ad.alimentar_cache_com_as_planilhas()["ja_estavam"], 1, "idempotente")
        self.processar(self.pasta())
        self.assertEqual(self.chamadas_ia, ["c.pdf"], "depois da carga, 0 chamadas")

    def test_reler_telas_ignora_o_cache_so_das_telas(self):
        pasta = self.saidas / "reserva" / "FULANA-FICTICIA-reserva-778"
        pasta.mkdir(parents=True)
        docs = [{"tipo": "Comprovante de Endereço", "nome": "c", "arquivo_local": "c.pdf"},
                {"tipo": "Tela de Aprovação", "nome": "t", "arquivo_local": "t.pdf"}]
        xlsx = pasta / "FULANA-FICTICIA-reserva-778.xlsx"
        v1.gerar_excel({"nome": "FULANA FICTICIA"}, docs, xlsx)
        (pasta / "c.pdf").write_bytes(b"conta")
        (pasta / "t.pdf").write_bytes(b"tela antiga")
        self.processar(pasta)  # lê as duas; a "Tela" sem tabela_renda
        self.assertEqual(sorted(self.chamadas_ia), ["c.pdf", "t.pdf"])
        self.processar(pasta, reler_telas=True)
        self.assertEqual(sorted(self.chamadas_ia), ["c.pdf", "t.pdf", "t.pdf"], "só a Tela foi relida")


class LeituraValida(BaseCache):
    """P1 da Forja: só leitura válida e completa entra no cache."""

    def test_leitura_parcial_ou_invalida_nao_entra(self):
        sem_campos = {k: v for k, v in LEITURA.items() if k not in ("nome_encontrado", "legivel")}
        invalidas = [sem_campos, dict(LEITURA, tipo_identificado=""), dict(LEITURA, tipo_identificado=None),
                     dict(LEITURA, legivel="sim"), dict(LEITURA, tabela_renda="cortada"),
                     dict(LEITURA, confianca_campos="alta"), dict(LEITURA, erro="timeout"),
                     dict(LEITURA, tipo_identificado="Fora do escopo de validação")]
        self.assertFalse(ad.leitura_valida(["nao", "e", "dict"]))
        for leitura in invalidas:
            with self.subTest(leitura=leitura):
                self.assertFalse(ad.leitura_valida(leitura))
                ad.extrair_dados_do_documento = lambda *a, l=leitura: (self.chamadas_ia.append("x"), l)[1]
                self.chamadas_ia.clear()
                self.processar(self.pasta())
                self.processar(self.pasta())
                self.assertEqual(self.chamadas_ia, ["x", "x"], "leitura inválida não vira cache")
                self.assertEqual(list((self.tmp / "cache_ia").rglob("*.json")), [])
        self.assertTrue(ad.leitura_valida(dict(LEITURA, tabela_renda=[{"cpf": "1", "renda_bruta": "3.500"}])))

    def test_cache_invalido_no_disco_e_ignorado(self):
        arquivo = self.tmp / "a.pdf"
        arquivo.write_bytes(b"conteudo")
        ad.guardar_no_cache(arquivo, LEITURA, "", "ia")
        destino = next((self.tmp / "cache_ia").rglob("*.json"))
        dados = json.loads(destino.read_text(encoding="utf-8"))
        del dados["leitura"]["legivel"]
        destino.write_text(json.dumps(dados), encoding="utf-8")
        self.assertIsNone(ad.leitura_do_cache(arquivo))

    def test_resposta_cortada_pelo_limite_vira_falha_da_ia(self):
        class Resp:
            text = json.dumps(LEITURA)
            candidates = [type("C", (), {"finish_reason": type("F", (), {"name": "MAX_TOKENS"})()})()]

        class Cliente:
            models = type("M", (), {"generate_content": staticmethod(lambda **k: Resp())})()
        arquivo = self.tmp / "a.png"
        arquivo.write_bytes(b"\x89PNG")
        ad.extrair_dados_do_documento = BaseCache_extrair
        guardado = ad.preparar_conteudo_para_ia
        ad.preparar_conteudo_para_ia = lambda c: [("image/png", b"x")]
        self.addCleanup(setattr, ad, "preparar_conteudo_para_ia", guardado)
        falso = types.SimpleNamespace(
            Part=types.SimpleNamespace(from_bytes=lambda **k: k), GenerateContentConfig=lambda **k: k)
        with mock.patch.dict(sys.modules, {"google.genai": types.SimpleNamespace(types=falso),
                                           "google.genai.types": falso}):
            with self.assertRaises(ad.ia_falhas.FalhaIA) as cortada:
                ad.extrair_dados_do_documento(Cliente(), "m", arquivo, "")
            self.assertEqual(cortada.exception.classe, "resposta_invalida")
            Resp.candidates = [type("C", (), {"finish_reason": type("F", (), {"name": "STOP"})()})()]
            self.assertTrue(ad.leitura_valida(ad.extrair_dados_do_documento(Cliente(), "m", arquivo, "")))
            Resp.text = '["lista", "e nao objeto"]'
            with self.assertRaises(ad.ia_falhas.FalhaIA) as lista:
                ad.extrair_dados_do_documento(Cliente(), "m", arquivo, "")
            self.assertNotIn("lista", lista.exception.detalhe, "o texto lido não vai para o log")


BaseCache_extrair = ad.extrair_dados_do_documento


class ProvaDaPlanilha(BaseCache):
    """P1 da Forja: a carga pelas planilhas só importa com prova de que a
    leitura é do arquivo que está na pasta."""

    def preparar(self):
        pasta = self.pasta()
        self.processar(pasta)  # planilha com leitura completa, gravada depois do download
        shutil.rmtree(self.tmp / "cache_ia")
        ad.listar_pastas_de_cliente = lambda: [pasta]
        return pasta, pasta / "FULANA-FICTICIA-reserva-777.xlsx"

    def test_pdf_trocado_com_o_mesmo_nome_nao_importa(self):
        pasta, xlsx = self.preparar()
        (pasta / "c.pdf").write_bytes(b"%PDF OUTRO papel, mesmo nome")  # depois da planilha
        t = xlsx.stat().st_mtime
        os.utime(pasta / "c.pdf", (t + 60, t + 60))
        c = ad.alimentar_cache_com_as_planilhas()
        self.assertEqual((c["gravadas"], c["sem_prova"]), (0, 1))
        self.assertEqual(list((self.tmp / "cache_ia").rglob("*.json")), [])

    def test_arquivo_muito_anterior_a_planilha_nao_importa(self):
        pasta, xlsx = self.preparar()
        t = xlsx.stat().st_mtime - ad.JANELA_DA_PROVA - 60  # planilha regravada sem download
        os.utime(pasta / "c.pdf", (t, t))
        self.assertEqual(ad.alimentar_cache_com_as_planilhas()["sem_prova"], 1)

    def test_revisar_apaga_importadas_sem_prova_e_invalidas(self):
        pasta, xlsx = self.preparar()
        self.assertEqual(ad.alimentar_cache_com_as_planilhas()["gravadas"], 1)
        provada = ad.hash_do_arquivo(pasta / "c.pdf")
        outros = {}
        for nome, conteudo, leitura, origem in (
                ("sem_prova.pdf", b"importado sem prova", LEITURA, "planilha"),
                ("da_ia.pdf", b"lido pela ia", LEITURA, "ia"),
                ("invalida.pdf", b"leitura antiga", LEITURA, "ia")):
            arq = self.tmp / nome
            arq.write_bytes(conteudo)
            self.assertTrue(ad.guardar_no_cache(arq, leitura, "", origem))
            outros[nome] = ad.hash_do_arquivo(arq)
        destino = ad._arquivo_do_cache(outros["invalida.pdf"])
        dados = json.loads(destino.read_text(encoding="utf-8"))
        dados["leitura"] = {"tipo_identificado": "RG", "observacoes": ""}  # as 20 antigas incompletas
        destino.write_text(json.dumps(dados), encoding="utf-8")
        r = ad.revisar_cache()
        self.assertEqual(r, {"ficaram_ia": 1, "ficaram_planilha": 1, "apagadas_invalidas": 1, "apagadas_sem_prova": 1})
        restantes = {p.stem for p in (self.tmp / "cache_ia").rglob("*.json")}
        self.assertEqual(restantes, {provada, outros["da_ia.pdf"]})
        self.assertEqual(ad.revisar_cache()["apagadas_sem_prova"], 0, "idempotente")

    def test_revisar_apaga_importada_com_leitura_diferente_da_provada(self):
        pasta, xlsx = self.preparar()
        ad.guardar_no_cache(pasta / "c.pdf", dict(LEITURA, nome_encontrado="DE OUTRA RODADA"), "", "planilha")
        self.assertEqual(ad.revisar_cache()["apagadas_sem_prova"], 1)


# Formato real da 34660 (19/10/2026): "CERTIDOES NASC FULANA E BELTRANO.pdf" tem
# duas certidões, e a IA devolveu uma lista com uma leitura para cada uma.
CERTIDOES_EM_LISTA = [
    dict(LEITURA, tipo_identificado="Certidão de Nascimento", nome_encontrado="FULANA FICTICIA",
         endereco_encontrado=None, mes_referencia=None, observacoes="certidão 1 de 2"),
    dict(LEITURA, tipo_identificado="Certidão de Nascimento", nome_encontrado="BELTRANO FICTICIO",
         endereco_encontrado=None, mes_referencia=None, observacoes="certidão 2 de 2"),
]


class LeituraEmLista(BaseCache):
    """Erro 1 de 20/10: 'list' object has no attribute 'get' na 34660."""

    def docs(self, pasta):
        docs = [{"tipo": "Certidão de Estado Civil", "nome": "cert", "arquivo_local": "cert.pdf"},
                {"tipo": "Comprovante de Endereço", "nome": "c", "arquivo_local": "c.pdf"}]
        xlsx = pasta / f"{pasta.name}.xlsx"
        v1.gerar_excel({"nome": "FULANA FICTICIA"}, docs, xlsx)
        (pasta / "cert.pdf").write_bytes(b"%PDF duas certidoes")
        (pasta / "c.pdf").write_bytes(b"%PDF conta")
        return xlsx

    def test_lista_da_ia_vira_falha_so_daquele_documento(self):
        pasta = self.saidas / "reserva" / "FULANA-FICTICIA-2o Promissario Comprador-reserva-34660"
        pasta.mkdir(parents=True)
        xlsx = self.docs(pasta)
        texto = json.dumps(CERTIDOES_EM_LISTA, ensure_ascii=False)
        resp = types.SimpleNamespace(text=texto, candidates=None)
        cliente = types.SimpleNamespace(models=types.SimpleNamespace(generate_content=lambda **k: resp))
        ad.criar_cliente_ia = lambda chave: cliente
        ad.extrair_dados_do_documento = BaseCache_extrair
        guardado = ad.preparar_conteudo_para_ia
        ad.preparar_conteudo_para_ia = lambda c: [("application/pdf", b"x")]
        self.addCleanup(setattr, ad, "preparar_conteudo_para_ia", guardado)
        falso = types.SimpleNamespace(
            Part=types.SimpleNamespace(from_bytes=lambda **k: k), GenerateContentConfig=lambda **k: k)
        with mock.patch.dict(sys.modules, {"google.genai": types.SimpleNamespace(types=falso),
                                           "google.genai.types": falso}):
            self.processar(pasta)  # antes: AttributeError e a pasta inteira falhava
        # 21/10/2026: resposta inválida é falha da IA, não divergência - os dois
        # ficam aguardando leitura (aba própria) e a pasta é concluída.
        self.assertEqual(ad.ler_resultados_anteriores(xlsx), {}, "nenhum vira divergência")
        from openpyxl import load_workbook
        falhas = list(load_workbook(xlsx)[ad.ABA_FALHAS_IA].iter_rows(min_row=2, values_only=True))
        self.assertEqual(sorted((f[0], f[2]) for f in falhas),
                         [("c.pdf", "resposta_invalida"), ("cert.pdf", "resposta_invalida")])
        self.assertEqual(list((self.tmp / "cache_ia").rglob("*.json")), [], "falha não vai para o cache")

    def test_lista_de_qualquer_origem_nao_derruba_a_pasta(self):
        pasta = self.saidas / "reserva" / "FULANA-FICTICIA-reserva-34661"
        pasta.mkdir(parents=True)
        xlsx = self.docs(pasta)
        ad.extrair_dados_do_documento = lambda *a: (self.chamadas_ia.append(Path(a[2]).name),
                                                    list(CERTIDOES_EM_LISTA) if "cert" in Path(a[2]).name
                                                    else dict(LEITURA))[1]
        self.processar(pasta)
        lido = ad.ler_resultados_anteriores(xlsx)
        self.assertNotIn("cert.pdf", lido, "aguardando leitura, não divergência")
        self.assertEqual(lido["c.pdf"]["doc_ia"]["nome_encontrado"], "FULANA FICTICIA")


if __name__ == "__main__":
    unittest.main()
