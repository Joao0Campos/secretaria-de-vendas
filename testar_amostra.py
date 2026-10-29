"""
Teste de amostra: auditar N reservas do mesmo status e consolidar o resultado
=============================================================================

PARA QUE SERVE: transformar anedota em número.

O João observou (22/09/2026) que TODA reserva já validada que ele testava dava
algum apontamento. Depois da recalibração, a reserva 33937 caiu de 4
divergências para 1 — o comprovante de endereço fora do prazo de 90 dias.

Mas 1 caso é anedota. Este script roda a auditoria em VÁRIAS reservas do mesmo
status e responde, com número:

  - de N vendas já validadas, quantas têm apontamento de verdade?
  - qual apontamento se repete? é sempre o mesmo, ou cada uma é diferente?
  - o comprovante de endereço fora do prazo é padrão ou exceção?

Se quase todas tiverem o mesmo problema e a regra for legítima (o João
confirmou em 22/09 que a dos 90 dias é), então o achado é da OPERAÇÃO, não da
ferramenta - e isso é uma informação forte pra levar a uma diretoria.

⚠️ SOMENTE LEITURA. Não altera nada no CV.

⚠️ CUSTO: cada reserva tem ~9 documentos, e cada documento é uma chamada de
IA. Cinco reservas ≈ 45 chamadas. Documentos já lidos numa rodada anterior são
reaproveitados (não use --forcar sem motivo).

Como rodar:
    python testar_amostra.py                       -> pergunta no terminal
    python testar_amostra.py "Venda Validada"      -> 5 reservas desse status
    python testar_amostra.py "Venda Validada" 10   -> 10 reservas
    python testar_amostra.py "Venda Validada" 5 --forcar

Sai um resumo no terminal e uma planilha `saidas/amostra-{status}.xlsx` com
uma linha por reserva e a contagem por tipo de achado.
"""

from __future__ import annotations

import re
import sys
from collections import Counter
from pathlib import Path

from mvp_v1_validador import carregar_config, PASTA_SAIDAS
from mvp_v2_reserva import (
    processar_reserva_completa,
    listar_todas_reservas,
    situacoes_disponiveis,
    casar_situacao,
)
from analisar_documentos import processar_pasta_cliente


# Agrupa apontamentos parecidos num rótulo só, pra dar pra contar. Sem isso, a
# variação de redação da IA (o mesmo comprovante já saiu como "com 168 dias",
# "validade vencida" e "não consegui ler a data") espalharia um problema só em
# três categorias e esconderia o padrão.
# MANUTENÇÃO: toda divergência que o analisar_documentos.py sabe emitir
# precisa de um padrão aqui. Quando a redação de uma mensagem muda lá e o
# padrão não é atualizado, o achado não some - ele passa a ser contado como
# "Outros" e o resumo deixa de nomear o problema. Foi o que aconteceu com
# "Comprovante de endereço sem nenhuma data..." (a mensagem foi reescrita em
# 23/09 junto com a regra da data de referência) e com a checagem de arquivo
# repetido, que nasceu sem padrão. Por isso o resumo agora IMPRIME o texto
# cru de tudo que caiu em "Outros": se aparecer algo ali, é padrão faltando.
CATEGORIAS = [
    # Primeiro da lista de propósito: é o achado mais grave que a ferramenta
    # produz. O mesmo arquivo preenchendo vários campos não é erro de leitura
    # nem documento vencido - é o processo sendo destravado no braço.
    ("Mesmo arquivo anexado em vários tipos de documento",
     [r"o mesmo arquivo .* foi anexado em \d+ tipos"]),
    # Era "fora do prazo de 90 dias" até 14/10/2026, quando a regra virou "mês M
    # vale até o fim de M+3" (gestora comercial). O rótulo mudou junto; o painel usa o rótulo
    # na chave do apontamento, então ele herda a identidade do antigo
    # (painel.banco.CATEGORIAS_RENOMEADAS) e o histórico continua ligado. Ao
    # renomear outra categoria, registre lá também. O padrão antigo fica para
    # planilhas não refeitas.
    ("Comprovante de endereço vencido",
     [r"comprovante de endereço vencido",
      r"comprovante de endereço com \d+ dias"]),
    ("Comprovante de endereço sem data legível",
     [r"comprovante de endereço sem nenhuma data",
      # redação antiga, mantida pra planilhas de rodadas anteriores
      r"não foi possível identificar a data de emissão do comprovante"]),
    ("Data do comprovante em formato não reconhecido",
     [r"data do comprovante em formato inesperado"]),
    ("Endereço não bate com o cadastro",
     [r"endereço no comprovante .* não bate"]),
    ("Endereço não identificado no comprovante",
     [r"não foi possível identificar o endereço no comprovante"]),
    ("Documento obrigatório não anexado",
     [r"documento obrigatório não encontrado"]),
    ("Documento em branco ou ilegível",
     [r"ilegível", r"em branco"]),
    ("Arquivo não pôde ser aberto ou lido",
     [r"não foi possível ler o documento"]),
    ("Renda diferente da cadastrada",
     [r"renda do documento .* difere",
      # por proponente, desde 14/10/2026 (conferir_renda_da_pessoa)
      r"renda d[oa] .* na tela de avaliação de risco .* difere"]),
    ("Renda não conferida por proponente",
     [r"sem a tabela renda completa",
      r"não traz renda legível",
      r"não aparece na tabela renda",
      r"não consegui a renda .* no cv",
      r"o cv não tem renda para esta pessoa"]),
    ("Certidão de casamento do cônjuge vale para os dois",
     [r"certidão de casamento lida na pasta do"]),
    ("Documento de identificação vencido",
     [r"validade vencida"]),
    ("CPF diferente do cadastro",
     [r"cpf no documento .* difere"]),
    # 20/10/2026: na CNH, o Nº de registro ou o RG lido no lugar do CPF.
    ("Número da CNH lido como CPF",
     [r"número lido na cnh .* não é um cpf válido"]),
    ("CNPJ diferente do cadastro",
     [r"cnpj no documento .* difere"]),
    ("Nome diferente do cadastro",
     [r"nome no documento .* difere"]),
]


# Categorias que tratam do MESMO problema de fundo, só que lidas de formas
# diferentes pela IA. Sem este agrupamento a contagem se dividiria e esconderia
# o padrão - o comprovante de endereço da 33937, por exemplo, já foi lido como
# "com 168 dias", como "validade vencida" e como "não consegui ler a data", em
# três execuções do MESMO arquivo. São a mesma reserva com o mesmo problema.
FAMILIAS = {
    "Comprovante de endereço vencido": "Comprovante de endereço com problema de data",
    "Comprovante de endereço sem data legível": "Comprovante de endereço com problema de data",
    "Data do comprovante em formato não reconhecido": "Comprovante de endereço com problema de data",

    "Endereço não bate com o cadastro": "Endereço do comprovante não confere com o cadastro",
    "Endereço não identificado no comprovante": "Endereço do comprovante não confere com o cadastro",

    "CPF diferente do cadastro": "Nome ou documento diferente do cadastro",
    "CNPJ diferente do cadastro": "Nome ou documento diferente do cadastro",
    "Nome diferente do cadastro": "Nome ou documento diferente do cadastro",

    "Documento em branco ou ilegível": "Documento que não dá pra conferir",
    "Arquivo não pôde ser aberto ou lido": "Documento que não dá pra conferir",
}
# "Documento de identificação vencido" NÃO entra na família do comprovante de
# endereço. Entrava até 28/09/2026, de quando "Validade vencida" ainda podia
# vir de um comprovante (o vencimento do boleto lido como validade). Desde a
# correção de 22/09 o comprovante de endereço sai da função ANTES da regra
# genérica de validade, então "Validade vencida" só pode vir de RG, CNH e
# afins - contá-la como problema de comprovante de endereço inflaria o número
# que vai pra diretoria com um caso que é de outra natureza.


def categorizar(texto: str) -> str:
    t = texto.lower()
    for rotulo, padroes in CATEGORIAS:
        if any(re.search(p, t) for p in padroes):
            return rotulo
    return "Outros"


def familia_de(categoria: str) -> str:
    """Agrupa categorias que são o mesmo problema de fundo. Categoria sem
    família declarada é família de si mesma."""
    return FAMILIAS.get(categoria, categoria)


def ler_achados(pasta: Path) -> tuple[list[str], list[str], str]:
    """Lê a aba de divergências da pasta já processada. Devolve
    (divergencias, avisos, nome_do_cliente)."""
    from openpyxl import load_workbook

    caminhos = list(pasta.glob("*.xlsx"))
    if not caminhos:
        return [], [], pasta.name
    wb = load_workbook(caminhos[0], data_only=True)

    nome = pasta.name
    if "Dados Cadastrais" in wb.sheetnames:
        for linha in wb["Dados Cadastrais"].iter_rows(min_row=2, values_only=True):
            if linha and linha[0] == "nome":
                nome = linha[1] or pasta.name
                break

    if "Divergencias (IA)" not in wb.sheetnames:
        return [], [], nome

    ws = wb["Divergencias (IA)"]
    cabecalho = [c.value for c in ws[1]]
    divergencias: list[str] = []
    avisos: list[str] = []
    for linha in ws.iter_rows(min_row=2, values_only=True):
        if not any(linha):
            continue
        d = dict(zip(cabecalho, linha))
        texto_div = str(d.get("divergencias") or "")
        if texto_div and texto_div != "OK":
            divergencias.extend(x.strip() for x in texto_div.split(" | ") if x.strip())
        texto_avi = str(d.get("avisos") or "")
        if texto_avi:
            avisos.extend(x.strip() for x in texto_avi.split(" | ") if x.strip())
    return divergencias, avisos, nome


def gravar_planilha(
    caminho: Path, linhas: list[dict], contagem: Counter,
    contagem_familia: Counter, total: int,
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    ws = wb.active
    ws.title = "Por reserva"
    colunas = ["idreserva", "cliente", "modalidade", "pessoas na operação",
               "quem tem divergência", "divergencias", "avisos", "o_que_foi_apontado",
               # no fim, para não deslocar as colunas antigas (09/10/2026)
               "incompleta (associados não consultados)"]
    ws.append(colunas)
    for cel in ws[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="0E6070")
    for r in linhas:
        ws.append([
            r["idreserva"], r["cliente"], r["modalidade"],
            r.get("n_pessoas", 1),
            ", ".join(r.get("pessoas_com_achado") or []) or "-",
            r["n_divergencias"], r["n_avisos"],
            " | ".join(r["categorias"]) or "(nada)",
            r.get("incompleta") or "",
        ])
        if r["n_divergencias"]:
            for cel in ws[ws.max_row]:
                cel.fill = PatternFill("solid", fgColor="F7E3E5")
    for idx, largura in enumerate([12, 34, 26, 18, 34, 14, 12, 60, 40], start=1):
        ws.column_dimensions[get_column_letter(idx)].width = largura

    ws2 = wb.create_sheet("Padrões")
    ws2.append([f"Amostra de {total} reserva(s) já validadas"])
    ws2["A1"].font = Font(bold=True, size=12)
    ws2.append([])
    ws2.append(["PROBLEMA DE FUNDO", "em quantas reservas", "% da amostra"])
    for cel in ws2[3]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="0E6070")
    for rotulo, n in contagem_familia.most_common():
        pct = (n / total * 100) if total else 0
        ws2.append([rotulo, n, f"{pct:.0f}%"])
        if pct >= 80:
            for cel in ws2[ws2.max_row]:
                cel.fill = PatternFill("solid", fgColor="F7E3E5")

    ws2.append([])
    ws2.append(["detalhe - como cada um foi lido", "em quantas reservas", "% da amostra"])
    for cel in ws2[ws2.max_row]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="6B6B63")
    for rotulo, n in contagem.most_common():
        pct = (n / total * 100) if total else 0
        ws2.append([rotulo, n, f"{pct:.0f}%"])
    for idx, largura in enumerate([56, 22, 14], start=1):
        ws2.column_dimensions[get_column_letter(idx)].width = largura
    for linha in ws2.iter_rows():
        for cel in linha:
            cel.alignment = Alignment(vertical="top", wrap_text=True)

    pendentes = [
        (r["idreserva"], t) for r in linhas for t in r.get("nao_classificados", [])
    ]
    if pendentes:
        ws3 = wb.create_sheet("Nao classificados")
        ws3.append(["Apontamentos que caíram em 'Outros' - falta padrão em CATEGORIAS"])
        ws3["A1"].font = Font(bold=True, size=12)
        ws3.append([])
        ws3.append(["idreserva", "texto do apontamento"])
        for cel in ws3[3]:
            cel.font = Font(bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", fgColor="6B6B63")
        for idr, texto in pendentes:
            ws3.append([idr, texto])
        ws3.column_dimensions["A"].width = 12
        ws3.column_dimensions["B"].width = 110
        for linha in ws3.iter_rows():
            for cel in linha:
                cel.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(caminho)


def testar(situacao: str, quantidade: int, forcar: bool = False) -> None:
    cfg = carregar_config()

    print("=" * 72)
    print(f"AMOSTRA: {quantidade} reserva(s) com situação '{situacao}'")
    print("Somente leitura - nada é alterado no CV.")
    print("=" * 72)
    print()
    print("Buscando a lista de reservas (demora um pouco, é a lista completa)...")
    todas = listar_todas_reservas(cfg)

    exata, opcoes = casar_situacao(situacao, todas)
    if exata is None:
        print(f"\nNão consegui identificar a situação '{situacao}'.")
        print("Você quis dizer alguma destas?" if len(opcoes) < len(situacoes_disponiveis(todas))
              else "Situações que existem no seu CV agora:")
        for s in opcoes:
            print(f"  - {s}")
        return
    if exata.strip().lower() != situacao.strip().lower():
        print(f"\nEntendi '{situacao}' como \"{exata}\".")
    situacao = exata
    selecionadas = [r for r in todas if (r.get("situacao") or "").strip() == exata]

    # As mais recentes primeiro: refletem melhor como o processo está hoje.
    selecionadas.sort(key=lambda r: r.get("idreserva", 0), reverse=True)
    amostra = selecionadas[:quantidade]

    print(f"\n{len(selecionadas)} reserva(s) nesse status. Vou auditar as {len(amostra)} mais recentes:")
    for r in amostra:
        print(f"  - {r.get('idreserva')}  {r.get('empreendimento', '')}")
    print()

    linhas: list[dict] = []
    falhas: list[tuple[int, str]] = []

    for i, item in enumerate(amostra, start=1):
        idreserva = item["idreserva"]
        print()
        print("#" * 72)
        print(f"# [{i}/{len(amostra)}] RESERVA {idreserva}")
        print("#" * 72)
        try:
            # Uma pasta POR PESSOA desde 05/10/2026 (titular, 2º proponente,
            # avalista). Os achados das três entram na mesma linha da reserva:
            # a pergunta da amostra é "esta VENDA tem apontamento?", não "esta
            # pessoa tem" - uma venda com problema no 2º comprador tem
            # problema, ponto.
            #
            # Associados consultados UMA vez por reserva e repassados às pastas
            # (09/10/2026). Antes cada pasta consultava de novo, e uma falha
            # nessa consulta (429) não aparecia no resumo: a reserva entrava na
            # contagem como se tivesse sido auditada por inteiro.
            pastas, associados, falha_associados = processar_reserva_completa(
                idreserva, devolver_associados=True
            )
            for pasta in pastas:
                processar_pasta_cliente(pasta, forcar=forcar,
                                        associados_consultados=associados,
                                        falha_associados=falha_associados)
        except SystemExit as e:
            print(f"  FALHOU: {e}")
            falhas.append((idreserva, str(e)))
            continue
        except Exception as e:  # uma reserva quebrada não pode derrubar a amostra
            print(f"  FALHOU: {e}")
            falhas.append((idreserva, str(e)))
            continue

        divergencias: list[str] = []
        avisos: list[str] = []
        cliente = ""
        pessoas_com_achado: list[str] = []
        for pasta in pastas:
            d, a_, nome = ler_achados(pasta)
            if not cliente:
                cliente = nome          # a primeira é a do titular
            if d:
                pessoas_com_achado.append(nome)
            divergencias.extend(d)
            avisos.extend(a_)
        if len(pastas) > 1:
            print(f"  (operação com {len(pastas)} pessoas; "
                  f"{len(pessoas_com_achado)} com divergência)")
        categorias = sorted({categorizar(d) for d in divergencias})
        # Guarda o texto cru do que caiu em "Outros". Sem isso, um padrão
        # desatualizado em CATEGORIAS vira uma linha muda de "Outros 60%" e o
        # achado mais importante da rodada some do resumo sem ninguém notar.
        nao_classificados = [d for d in divergencias if categorizar(d) == "Outros"]
        linhas.append({
            "idreserva": idreserva,
            "cliente": cliente,
            "modalidade": item.get("tipovenda", ""),
            "n_divergencias": len(divergencias),
            "n_avisos": len(avisos),
            "categorias": categorias,
            "nao_classificados": nao_classificados,
            "n_pessoas": len(pastas),
            "pessoas_com_achado": pessoas_com_achado,
            "incompleta": falha_associados,
        })

    if not linhas:
        print("\nNenhuma reserva pôde ser auditada.")
        return

    # Conta em QUANTAS reservas cada apontamento apareceu (não quantas vezes
    # no total) - é isso que responde "é padrão ou exceção?".
    contagem: Counter = Counter()
    for r in linhas:
        for c in r["categorias"]:
            contagem[c] += 1

    total = len(linhas)
    limpas = sum(1 for r in linhas if r["n_divergencias"] == 0)

    print()
    print("=" * 72)
    print(f"RESULTADO DA AMOSTRA - {total} reserva(s) já validadas")
    print("=" * 72)
    print()
    print(f"  Sem nenhuma divergência : {limpas} de {total}")
    print(f"  Com alguma divergência  : {total - limpas} de {total}")
    com_2o = sum(1 for r in linhas if r.get("n_pessoas", 1) > 1)
    if com_2o:
        so_no_2o = sum(
            1 for r in linhas
            if r.get("n_pessoas", 1) > 1 and r["n_divergencias"]
            and r["cliente"] not in (r.get("pessoas_com_achado") or [])
        )
        print(f"  Com 2º proponente       : {com_2o} de {total}")
        print(f"     ... cuja divergência está SÓ em quem não é o titular: {so_no_2o}"
              f"  (invisível antes de 05/10/2026)")
    print()
    print("  Por reserva:")
    for r in linhas:
        marca = "OK " if r["n_divergencias"] == 0 else "!! "
        incompleta = (f"  [INCOMPLETA: associados não consultados - {r['incompleta']}]"
                      if r.get("incompleta") else "")
        print(f"    {marca}{r['idreserva']}  {r['cliente'][:32]:32} "
              f"{r['n_divergencias']} div, {r['n_avisos']} aviso(s){incompleta}")
    incompletas = sum(1 for r in linhas if r.get("incompleta"))
    if incompletas:
        print(f"\n  ATENÇÃO: {incompletas} reserva(s) INCOMPLETA(S) - a consulta de associados "
              f"falhou; nelas o 2º comprador não foi validado e pode haver divergência "
              f"falsa de nome/CPF. Rodar de novo.")

    # Agrupado por problema de fundo: é esta contagem que responde "é padrão
    # ou exceção?", porque não se perde na variação de redação da IA.
    contagem_familia: Counter = Counter()
    for r in linhas:
        for f in {familia_de(c) for c in r["categorias"]}:
            contagem_familia[f] += 1

    print()
    print("  PADRÃO (agrupado por problema de fundo):")
    for rotulo, n in contagem_familia.most_common():
        pct = n / total * 100
        destaque = "  <-- em quase todas" if pct >= 80 else ""
        print(f"    {n}/{total} ({pct:3.0f}%)  {rotulo}{destaque}")

    print()
    print("  Detalhe (como cada um foi lido):")
    for rotulo, n in contagem.most_common():
        pct = n / total * 100
        print(f"    {n}/{total} ({pct:3.0f}%)  {rotulo}")

    pendentes = [
        (r["idreserva"], t) for r in linhas for t in r["nao_classificados"]
    ]
    if pendentes:
        print()
        print("  ATENÇÃO - apontamentos que caíram em 'Outros' (falta padrão em")
        print("  CATEGORIAS; enquanto isso eles não aparecem nomeados acima):")
        for idr, texto in pendentes:
            print(f"    - [{idr}] {texto[:150]}")

    if falhas:
        print()
        print(f"  {len(falhas)} reserva(s) falharam:")
        for idr, erro in falhas:
            print(f"    - {idr}: {erro[:100]}")

    nome_arquivo = re.sub(r"[^\w\-]+", "-", situacao.strip().lower()).strip("-")
    caminho = PASTA_SAIDAS / f"amostra-{nome_arquivo}.xlsx"
    gravar_planilha(caminho, linhas, contagem, contagem_familia, total)
    print()
    print(f"Planilha da amostra: {caminho}")
    print("Abas: 'Por reserva' e 'Padrões'")


def main() -> None:
    argv = sys.argv[1:]
    forcar = "--forcar" in argv
    argv = [a for a in argv if a != "--forcar"]

    if argv:
        situacao = argv[0]
        quantidade = int(argv[1]) if len(argv) > 1 and argv[1].isdigit() else 5
    else:
        situacao = input("Situação da reserva (ex: Venda Validada): ").strip()
        if not situacao:
            print("Precisa informar a situação.")
            sys.exit(1)
        entrada = input("Quantas reservas auditar? (padrão 5): ").strip()
        quantidade = int(entrada) if entrada.isdigit() else 5

    testar(situacao, quantidade, forcar=forcar)


if __name__ == "__main__":
    main()
