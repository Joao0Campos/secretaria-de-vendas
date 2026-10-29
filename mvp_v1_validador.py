"""
MVP v1 - Secretaria de Vendas
=============================

Fluxo (conforme combinado):
  1) GET /v1/comercial/precadastro/{idPrecadastro}/documentos
     -> pega os links dos documentos anexados + o idpessoa_cv do titular
  2) GET /v1/cadastros/clientes?idcliente={idpessoa_cv}
     -> pega o que está de fato cadastrado no CV (dados do titular)
  3) Salva tudo em uma pasta "{nome do cliente}-{idprecadastro}":
     - os documentos baixados
     - um Excel com os dados cadastrais + a checklist de documentos

O que este script NÃO faz ainda (fica pra próxima etapa):
  - Comparar dado a dado (isso é o "encontrar erros" - v1.1)
  - Processar cônjuge / 2º comprador (só o titular por enquanto,
    porque o endpoint de documentos só devolve o idpessoa_cv do titular)

Como rodar:
    python mvp_v1_validador.py                  -> pergunta no terminal
    python mvp_v1_validador.py 34135             -> processa só esse pré-cadastro
    python mvp_v1_validador.py --situacao Novo   -> processa TODOS com essa situação

O modo "por situação" usa GET /v1/cvdw/precadastros (paginado) pra listar
todos os pré-cadastros do CV, filtra pelo texto da situação (ex: "Novo",
"Aprovado" - exatamente como aparece no CV, sem diferenciar maiúsculas de
minúsculas) e roda o mesmo fluxo de cima em cada um, um por vez. Se um
pré-cadastro falhar, os outros continuam - no final sai um resumo com
quantos deram certo e quais falharam.

Configuração (uma das duas formas):
  a) copie config.example.json para config.json e preencha
  b) ou defina as variáveis de ambiente CVCRM_DOMINIO, CVCRM_EMAIL, CVCRM_TOKEN
"""

from __future__ import annotations

import json
import os
import re
import sys
import unicodedata
from pathlib import Path

import requests

# --------------------------------------------------------------------------
# Configuração
# --------------------------------------------------------------------------

PASTA_SCRIPT = Path(__file__).resolve().parent
PASTA_SAIDAS = PASTA_SCRIPT / "saidas"

# Particionamento (29/08/2026): dentro de saidas/, o mvp_v1_validador.py
# (pré-cadastro) grava em "pre_cadastro/", e o mvp_v2_reserva.py (reserva)
# grava em "reserva/" - evita misturar os dois tipos na mesma pasta.
PASTA_SAIDAS_PRECADASTRO = PASTA_SAIDAS / "pre_cadastro"
PASTA_SAIDAS_RESERVA = PASTA_SAIDAS / "reserva"


def carregar_config() -> dict:
    """Lê config.json se existir; senão cai para variáveis de ambiente."""
    caminho_config = PASTA_SCRIPT / "config.json"
    if caminho_config.exists():
        with open(caminho_config, "r", encoding="utf-8") as f:
            cfg = json.load(f)
    else:
        cfg = {
            "dominio": os.environ.get("CVCRM_DOMINIO", ""),
            "email": os.environ.get("CVCRM_EMAIL", ""),
            "token": os.environ.get("CVCRM_TOKEN", ""),
        }

    faltando = [k for k in ("dominio", "email", "token") if not cfg.get(k)]
    if faltando:
        raise SystemExit(
            f"Faltam campos de configuração: {', '.join(faltando)}. "
            "Copie config.example.json para config.json e preencha, "
            "ou defina CVCRM_DOMINIO / CVCRM_EMAIL / CVCRM_TOKEN."
        )
    return cfg


def base_url(cfg: dict) -> str:
    return f"https://{cfg['dominio']}.cvcrm.com.br/api"


def headers_v1(cfg: dict) -> dict:
    # Esquema de autenticação "email + token" (v1 / CVDW) - o mesmo que
    # o teste manual de conexão usava.
    return {
        "email": cfg["email"],
        "token": cfg["token"],
        "Accept": "application/json",
    }


# --------------------------------------------------------------------------
# Chamadas à API com retry
# --------------------------------------------------------------------------

def requisicao_com_retry(
    url: str, *, headers: dict, params: dict | None = None,
    timeout: int = 30, tentativas: int = 4, contexto: str = "",
):
    """GET que insiste quando o CV responde 429 (limite de requisições) ou 5xx.

    07/10/2026: numa amostra de 10 reservas o CV devolveu **429 Too many
    requests** na listagem de associados de duas delas. O script seguiu só com
    o titular, avisando - mas se aquelas reservas tivessem 2º comprador, ele
    teria sido ignorado em silêncio. Pior: na 34239 o mesmo 429 derrubou a
    busca da reserva e levou a ferramenta a exibir a renda de outra pessoa.

    Não é caso raro: a listagem de associados tem 227 páginas e a de reservas
    tem 55. Varrer tudo em sequência bate no limite com facilidade.

    Espera 2s, 4s, 8s entre as tentativas (respeitando 'Retry-After' quando o
    servidor manda um). Erro que não é de limite nem de servidor - 401, 404 -
    não é repetido: insistir não resolve e só atrasa.
    """
    import time

    espera = 2
    ultima = None
    for tentativa in range(1, tentativas + 1):
        resp = requests.get(url, headers=headers, params=params, timeout=timeout)
        ultima = resp
        if resp.status_code == 200:
            return resp
        if resp.status_code != 429 and resp.status_code < 500:
            return resp          # erro definitivo: devolve pra quem chamou tratar

        if tentativa == tentativas:
            break
        pausa = espera
        cabecalho = resp.headers.get("Retry-After")
        if cabecalho:
            try:
                pausa = max(pausa, int(float(cabecalho)))
            except (TypeError, ValueError):
                pass
        motivo = "limite de requisições" if resp.status_code == 429 else f"erro {resp.status_code}"
        print(f"      ({contexto or url}): {motivo}, tentando de novo em {pausa}s "
              f"[{tentativa}/{tentativas - 1}]")
        time.sleep(pausa)
        espera *= 2

    return ultima


# --------------------------------------------------------------------------
# Utilitários
# --------------------------------------------------------------------------

def sanitizar_nome(texto: str) -> str:
    """Deixa um texto seguro para usar como nome de pasta/arquivo no Windows."""
    if not texto:
        return "SEM_NOME"
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode("ascii")
    texto = re.sub(r'[<>:"/\\|?*]', "", texto)
    texto = re.sub(r"\s+", " ", texto).strip(" .")
    return texto or "SEM_NOME"


def extensao_do_link(link: str) -> str:
    """O link vem com querystring assinada (?X-Amz-...). A extensão real
    está antes do '?'."""
    sem_query = link.split("?", 1)[0]
    _, ext = os.path.splitext(sem_query)
    return ext if ext else ".pdf"  # sem extensão = normalmente é PDF


# --------------------------------------------------------------------------
# Passo 1: documentos do pré-cadastro
# --------------------------------------------------------------------------

def buscar_documentos_precadastro(cfg: dict, idprecadastro: int) -> dict:
    """GET /v1/comercial/precadastro/{idPrecadastro}/documentos

    Retorna o objeto 'dados' inteiro (não só os documentos), porque é lá
    que vêm idpessoa_cv, idempreendimento_cv, idunidade_cv etc.

    IMPORTANTE (verificado no schema oficial + num teste manual):
    'documentos' é um DICIONÁRIO agrupado por papel da pessoa
    (ex.: "titular"), não uma lista simples. Cada item de cada grupo tem:
    idprecadastrosdocumentos, nome, tipo, situacao, motivo_reprovacao,
    descricao_motivo_reprovacao, link_download.
    """
    url = f"{base_url(cfg)}/v1/comercial/precadastro/{idprecadastro}/documentos"
    resp = requests.get(url, headers=headers_v1(cfg), timeout=30)

    if resp.status_code == 204:
        raise SystemExit(f"Pré-cadastro {idprecadastro}: não há documentos cadastrados.")
    if resp.status_code != 200:
        raise SystemExit(
            f"Erro {resp.status_code} ao buscar documentos do pré-cadastro "
            f"{idprecadastro}: {resp.text[:500]}"
        )

    corpo = resp.json()
    dados = corpo.get("dados")
    if not dados:
        raise SystemExit(f"Resposta sem 'dados' para o pré-cadastro {idprecadastro}: {corpo}")
    return dados


def documentos_do_titular(dados_precadastro: dict) -> list[dict]:
    """Extrai a lista de documentos do titular, deduplicada por tipo
    (mantendo o de maior id = mais recente, mesma regra usada no fluxo
    de documentos de reserva)."""
    grupos = dados_precadastro.get("documentos") or {}
    docs_titular = grupos.get("titular") or []

    mais_recente_por_tipo: dict[str, dict] = {}
    for doc in docs_titular:
        tipo = doc.get("tipo") or doc.get("idtipo")
        atual = mais_recente_por_tipo.get(tipo)
        id_doc = doc.get("idprecadastrosdocumentos", 0)
        if atual is None or id_doc > atual.get("idprecadastrosdocumentos", 0):
            mais_recente_por_tipo[tipo] = doc

    return list(mais_recente_por_tipo.values())


def baixar_documentos(cfg: dict, documentos: list[dict], pasta_destino: Path) -> list[dict]:
    """Baixa cada documento para pasta_destino. Retorna a lista de documentos
    com um campo extra 'arquivo_local' (ou 'erro_download' se falhar).

    Baixa TODOS os documentos, independente da situação (Aprovado /
    Pendente / Reprovado) - ao contrário do fluxo de "montar pasta final"
    de outra integração interna, aqui o objetivo é justamente enxergar o que está
    pendente/reprovado para a validação.
    """
    pasta_destino.mkdir(parents=True, exist_ok=True)
    resultado = []

    # Nomes gravados NESTA execução. Serve pra distinguir dois casos que antes
    # eram confundidos (ver o bloco de desambiguação mais abaixo).
    nomes_desta_rodada: set[str] = set()

    for i, doc in enumerate(documentos, start=1):
        nome_original = doc.get("nome") or f"documento_{i}"
        link = doc.get("link_download") or doc.get("link")
        registro = dict(doc)

        if not link:
            registro["erro_download"] = "sem link de download"
            resultado.append(registro)
            print(f"  [{i}/{len(documentos)}] {nome_original}: sem link, pulando")
            continue

        ext = extensao_do_link(link)
        base_nome = sanitizar_nome(Path(nome_original).stem)
        tipo_prefixo = sanitizar_nome(doc.get("tipo") or "documento")
        nome_arquivo = f"{tipo_prefixo} - {base_nome}{ext}"

        # Desambiguação de nome. Há DOIS casos diferentes aqui, e tratá-los
        # igual causou um bug caro (descoberto em 22/09/2026):
        #
        # (a) Dois documentos DIFERENTES com o mesmo nome genérico na MESMA
        #     rodada - o CV faz isso direto (ex: "Branco.pdf" em CTPS, FGTS e
        #     Simulação). Aí sim precisa de sufixo "(2)", senão um sobrescreve
        #     o outro.
        #
        # (b) O MESMO documento, numa nova execução, encontrando o arquivo que
        #     ele próprio baixou antes. Aqui sufixo é errado: é pra sobrescrever.
        #
        # A versão anterior só olhava `caminho.exists()` e tratava tudo como
        # (a). Resultado: cada execução duplicava a pasta inteira - a da reserva 33937
        # chegou a 4 cópias de cada documento (36 arquivos em vez de 9). Pior
        # que o disco: como o nome mudava a cada rodada, o cache de resultados
        # da IA (que casa por nome de arquivo) NUNCA acertava, e toda execução
        # relia tudo com IA mesmo respondendo "não reprocessar" - custo cheio
        # de API em toda rodada, silenciosamente.
        caminho = pasta_destino / nome_arquivo
        sufixo = 2
        while nome_arquivo in nomes_desta_rodada:
            nome_arquivo = f"{tipo_prefixo} - {base_nome} ({sufixo}){ext}"
            caminho = pasta_destino / nome_arquivo
            sufixo += 1
        nomes_desta_rodada.add(nome_arquivo)

        print(f"  [{i}/{len(documentos)}] Baixando: {nome_arquivo} ...", end=" ")
        try:
            # binário: NÃO usar .text, ler .content
            resp = requisicao_com_retry(
                link, headers=headers_v1(cfg), timeout=60,
                contexto=f"download de {nome_arquivo}")
            if resp.status_code == 200:
                caminho.write_bytes(resp.content)
                registro["arquivo_local"] = str(caminho.name)
                print("ok")
            else:
                registro["erro_download"] = f"status {resp.status_code}"
                print(f"falhou ({resp.status_code})")
        except requests.RequestException as exc:
            registro["erro_download"] = str(exc)
            print(f"falhou ({exc})")

        resultado.append(registro)

    return resultado


# --------------------------------------------------------------------------
# Passo 2: dados cadastrais do cliente
# --------------------------------------------------------------------------

def buscar_precadastro(cfg: dict, idprecadastro: int) -> dict | None:
    """GET /v1/comercial/precadastro/{idPrecadastro} - o registro do
    pré-cadastro em si (não os documentos dele, nem o cadastro da pessoa).

    Existe por causa de um erro caro, descoberto em 23/09/2026: a validação de
    renda comparava contra o campo ERRADO. O CV mostra dois valores na tela do
    pré-cadastro, e eles vivem em campos diferentes deste endpoint:

        "Renda mensal bruta"    -> renda_cliente_principal   (só o titular)
        "Renda familiar bruta"  -> renda_total               (o núcleo)

    O cadastro da PESSOA (/v1/cadastros/clientes) tem um campo chamado
    `renda_familiar` que, apesar do nome, traz o valor INDIVIDUAL - o mesmo
    de renda_cliente_principal. Comparar a renda do núcleo, lida na Tela de
    Avaliação de Risco, contra esse campo acusava erro em cadastro correto.

    Medido no pré-cadastro 34626: renda_cliente_principal = 1.636,23,
    renda_total = 5.555,57, e 1.636,23 + 3.919,34 (do 2º comprador)
    = 5.555,57 exatamente. O cadastro estava certo o tempo todo.

    Devolve None em qualquer falha - o chamador cai no comportamento antigo.
    """
    url = f"{base_url(cfg)}/v1/comercial/precadastro/{idprecadastro}"
    try:
        resp = requests.get(url, headers=headers_v1(cfg), timeout=30)
    except requests.RequestException:
        return None
    if resp.status_code != 200:
        return None
    try:
        corpo = resp.json()
    except ValueError:
        return None

    registros = corpo.get("precadastros") if isinstance(corpo, dict) else None
    if isinstance(registros, list) and registros:
        return registros[0]
    dados = corpo.get("dados") if isinstance(corpo, dict) else None
    if isinstance(dados, dict):
        return dados
    return corpo if isinstance(corpo, dict) else None


def buscar_dados_cliente(cfg: dict, idpessoa: str | int) -> dict:
    """GET /v1/cadastros/clientes?idcliente={idpessoa}

    Atenção: o parâmetro da query é 'idcliente', não 'idpessoa' - o valor
    que você recebe como idpessoa_cv no passo 1 entra aqui.
    """
    url = f"{base_url(cfg)}/v1/cadastros/clientes"
    params = {"idcliente": idpessoa}
    resp = requisicao_com_retry(url, headers=headers_v1(cfg), params=params,
                                contexto=f"cadastro do cliente {idpessoa}")

    if resp.status_code != 200:
        raise SystemExit(
            f"Erro {resp.status_code} ao buscar cliente idcliente={idpessoa}: "
            f"{resp.text[:500]}"
        )

    corpo = resp.json()
    pessoas = corpo.get("pessoas") or []
    if not pessoas:
        raise SystemExit(f"Nenhum cliente encontrado para idcliente={idpessoa}: {corpo}")
    return pessoas[0]


# --------------------------------------------------------------------------
# Passo 3: Excel com dados cadastrais + checklist de documentos
# --------------------------------------------------------------------------

def gerar_excel(cliente: dict, documentos: list[dict], caminho_saida: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # --- Aba 1: Dados cadastrais (o que está no CV) ---
    ws1 = wb.active
    ws1.title = "Dados Cadastrais"
    ws1.append(["Campo", "Valor"])
    for cel in ws1[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="0E6070")

    # Campos mais relevantes para a validação primeiro, o resto depois
    campos_prioritarios = [
        "nome", "documento_tipo", "documento", "cliente_tipo",
        "estado_civil", "data_casamento", "rg", "rg_orgao_emissor",
        "rg_data_emissao", "cep", "endereco", "numero", "complemento",
        "bairro", "cidade", "estado", "telefone", "celular", "email",
        "profissao", "renda_familiar",
    ]
    ja_mostrados = set()
    for campo in campos_prioritarios:
        if campo in cliente:
            ws1.append([campo, cliente.get(campo)])
            ja_mostrados.add(campo)
    for campo, valor in cliente.items():
        if campo in ja_mostrados or campo == "tags":
            continue
        if isinstance(valor, (dict, list)):
            valor = json.dumps(valor, ensure_ascii=False)
        ws1.append([campo, valor])

    ws1.column_dimensions["A"].width = 32
    ws1.column_dimensions["B"].width = 60

    # --- Aba 2: Documentos anexados (checklist p/ conferência) ---
    ws2 = wb.create_sheet("Documentos")
    colunas = ["tipo", "nome", "situacao", "motivo_reprovacao",
               "descricao_motivo_reprovacao", "arquivo_local", "erro_download"]
    ws2.append(colunas)
    for cel in ws2[1]:
        cel.font = Font(bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="0E6070")

    for doc in documentos:
        linha = [doc.get(c, "") for c in colunas]
        ws2.append(linha)
        situacao = (doc.get("situacao") or "").lower()
        cor = None
        if "reprov" in situacao:
            cor = "F7E3E5"
        elif "pend" in situacao:
            cor = "F6EBDA"
        elif "aprov" in situacao:
            cor = "E2F0E9"
        if cor:
            for cel in ws2[ws2.max_row]:
                cel.fill = PatternFill("solid", fgColor=cor)

    for idx, largura in enumerate([14, 32, 14, 22, 40, 28, 20], start=1):
        ws2.column_dimensions[get_column_letter(idx)].width = largura

    wb.save(caminho_saida)


# --------------------------------------------------------------------------
# Passo 0 (opcional): listar pré-cadastros por situação
# --------------------------------------------------------------------------

def buscar_pagina_precadastros(cfg: dict, pagina: int, por_pagina: int = 500) -> dict:
    """GET /v1/cvdw/precadastros - lista paginada de pré-cadastros (CVDW,
    mesma autenticação email+token). Cada item já vem com 'idprecadastro',
    'idpessoa' e 'situacao' (texto, ex: 'Novo')."""
    url = f"{base_url(cfg)}/v1/cvdw/precadastros"
    params = {"pagina": pagina, "registros_por_pagina": por_pagina}
    resp = requests.get(url, headers=headers_v1(cfg), params=params, timeout=30)
    if resp.status_code != 200:
        raise SystemExit(
            f"Erro {resp.status_code} ao listar pré-cadastros (página {pagina}): "
            f"{resp.text[:500]}"
        )
    return resp.json()


def listar_todos_precadastros(cfg: dict) -> list[dict]:
    """Percorre todas as páginas de /v1/cvdw/precadastros e junta tudo numa lista só."""
    pagina = 1
    todos: list[dict] = []
    while True:
        corpo = buscar_pagina_precadastros(cfg, pagina)
        itens = corpo.get("dados") or []
        todos.extend(itens)
        total_paginas = corpo.get("total_de_paginas", pagina)
        print(f"      página {pagina}/{total_paginas} ({len(itens)} registros)")
        if pagina >= total_paginas or not itens:
            break
        pagina += 1
    return todos


def situacoes_disponiveis(itens: list[dict]) -> list[str]:
    return sorted({(item.get("situacao") or "").strip() for item in itens if item.get("situacao")})


def processar_por_situacao(situacao_alvo: str) -> None:
    """Busca TODOS os pré-cadastros, filtra pela situação informada (comparação
    sem diferenciar maiúsculas/minúsculas) e processa um por um, reaproveitando
    processar_precadastro(). Se um item falhar, os outros continuam."""
    cfg = carregar_config()

    print("Buscando lista completa de pré-cadastros (pode demorar um pouco)...")
    todos = listar_todos_precadastros(cfg)
    print(f"Total encontrado: {len(todos)} pré-cadastro(s)")

    alvo_norm = situacao_alvo.strip().lower()
    selecionados = [
        item for item in todos
        if (item.get("situacao") or "").strip().lower() == alvo_norm
    ]

    if not selecionados:
        print(f"\nNenhum pré-cadastro encontrado com situação '{situacao_alvo}'.")
        print("Situações que existem no seu CV agora:")
        for s in situacoes_disponiveis(todos):
            print(f"  - {s}")
        return

    print(f"\n{len(selecionados)} pré-cadastro(s) com situação '{situacao_alvo}'. Processando...\n")

    ok: list[int] = []
    falhas: list[tuple[int, str]] = []
    for i, item in enumerate(selecionados, start=1):
        idprecadastro = item["idprecadastro"]
        print(f"\n=== [{i}/{len(selecionados)}] Pré-cadastro {idprecadastro} ===")
        try:
            processar_precadastro(idprecadastro)
            ok.append(idprecadastro)
        except SystemExit as exc:
            print(f"      FALHOU: {exc}")
            falhas.append((idprecadastro, str(exc)))
        except Exception as exc:  # não deixa um erro travar o lote inteiro
            print(f"      FALHOU (erro inesperado): {exc}")
            falhas.append((idprecadastro, str(exc)))

    print("\n=== Resumo ===")
    print(f"Processados com sucesso: {len(ok)}")
    print(f"Falharam: {len(falhas)}")
    for idp, motivo in falhas:
        print(f"  - {idp}: {motivo}")


# --------------------------------------------------------------------------
# Orquestração
# --------------------------------------------------------------------------

def processar_precadastro(idprecadastro: int) -> Path:
    cfg = carregar_config()

    print(f"=== Pré-cadastro {idprecadastro} ===")
    print("[1/4] Buscando documentos do pré-cadastro...")
    dados_pc = buscar_documentos_precadastro(cfg, idprecadastro)
    idpessoa = dados_pc.get("idpessoa_cv")
    if not idpessoa:
        raise SystemExit(
            "Não veio 'idpessoa_cv' na resposta de documentos - "
            "confira a resposta bruta no Postman antes de seguir."
        )
    print(f"      idpessoa_cv do titular: {idpessoa}")

    docs = documentos_do_titular(dados_pc)
    print(f"      {len(docs)} documento(s) (deduplicados por tipo)")

    print("[2/4] Buscando dados cadastrais do cliente...")
    cliente = buscar_dados_cliente(cfg, idpessoa)
    nome_cliente = cliente.get("nome", "SEM_NOME")
    print(f"      Cliente: {nome_cliente}")

    nome_pasta = f"{sanitizar_nome(nome_cliente)}-{idprecadastro}"
    pasta_cliente = PASTA_SAIDAS_PRECADASTRO / nome_pasta
    print(f"[3/4] Baixando documentos em: {pasta_cliente}")
    docs = baixar_documentos(cfg, docs, pasta_cliente)

    print("[4/4] Gerando Excel...")
    caminho_excel = pasta_cliente / f"{nome_pasta}.xlsx"
    gerar_excel(cliente, docs, caminho_excel)
    print(f"      Excel salvo em: {caminho_excel}")

    print("\nConcluído.")
    return pasta_cliente


def pedir_idprecadastro() -> int:
    """Pergunta o id do pré-cadastro no terminal, repetindo até vir um número válido."""
    while True:
        entrada = input("Número do pré-cadastro: ").strip()
        if entrada.isdigit():
            return int(entrada)
        print("Digite apenas números (ex: 34139). Tente de novo.")


def menu_interativo() -> None:
    print("O que você quer fazer?")
    print("  1 - Processar UM pré-cadastro específico (informar o número)")
    print("  2 - Processar TODOS os pré-cadastros de uma situação (ex: Novo, Aprovado...)")
    escolha = input("Escolha 1 ou 2: ").strip()

    if escolha == "2":
        situacao = input("Situação a processar (exatamente como aparece no CV, ex: Novo): ").strip()
        if not situacao:
            print("Situação vazia, cancelando.")
            return
        processar_por_situacao(situacao)
        return

    idprecadastro = pedir_idprecadastro()
    processar_precadastro(idprecadastro)


def main() -> None:
    # python mvp_v1_validador.py 34139           -> processa só esse
    # python mvp_v1_validador.py --situacao Novo  -> processa todos com essa situação
    # python mvp_v1_validador.py (sem nada)       -> pergunta no terminal
    if len(sys.argv) == 2 and sys.argv[1].isdigit():
        processar_precadastro(int(sys.argv[1]))
    elif len(sys.argv) >= 3 and sys.argv[1] == "--situacao":
        processar_por_situacao(" ".join(sys.argv[2:]))
    elif len(sys.argv) == 1:
        menu_interativo()
    else:
        print("Uso:")
        print("  python mvp_v1_validador.py                 (pergunta no terminal)")
        print("  python mvp_v1_validador.py <idPrecadastro>")
        print("  python mvp_v1_validador.py --situacao <nome da situação>")
        sys.exit(1)


if __name__ == "__main__":
    main()
